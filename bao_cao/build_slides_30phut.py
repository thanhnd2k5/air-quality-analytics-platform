#!/usr/bin/env python3
"""Generate Slide_BTL_THUYET_TRINH_30phut.pptx — 49 slides (31 chính + 18 phụ lục A0–A17).

Nội dung bám đúng bao_cao/KE_HOACH_NOI_DUNG_THUYET_TRINH.md (kịch bản nói 30 phút).
Mọi số lấy từ artifacts/ qua report_lib.Facts + vài artifact phụ (onset_6h, cdc_eval, eda).
Speaker notes mỗi slide chứa lời nói + mồi → lá chắn + cửa cắt (không bao giờ hiện trên slide).

    py -3 bao_cao/build_slides_30phut.py
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import pandas as pd  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from report_lib import ART, FROM_SCRATCH_MINUTES, Facts, load_json, vn, vn_int  # noqa: E402

OUT = HERE / "Slide_BTL_THUYET_TRINH_30phut.pptx"
SHOTS = ART / "dashboard_screenshots"
N_TOTAL = 49

# Palette: smog charcoal (dominant) · alert orange (accent) · clean-air teal (support).
INK = RGBColor(0x1F, 0x2A, 0x37)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
MUTED = RGBColor(0x64, 0x74, 0x8B)
CARD = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_ORANGE = RGBColor(0xFF, 0xED, 0xD5)
SOFT_TEAL = RGBColor(0xCC, 0xFB, 0xF1)
SLATE_L = RGBColor(0xCB, 0xD5, 0xE1)  # muted text trên nền INK
AMBER = RGBColor(0xFD, 0xBA, 0x74)    # kicker trên nền INK
HEAD = "Cambria"
BODY = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CW = W - 2 * MARGIN  # content width

SUPERSCRIPT = str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")


def sci(p: float) -> str:
    mant, exp = f"{p:.0e}".split("e")
    return f"p ≈ {mant}×10{str(int(exp)).translate(SUPERSCRIPT)}"


# -------------------------------------------------------------------- deck
class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.title_only = self.prs.slide_layouts[5]
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # --- frames ------------------------------------------------------------
    def content(self, title: str, kicker: str, notes: str, title_size: int = 32):
        s = self.prs.slides.add_slide(self.title_only)
        self.n += 1
        t = s.shapes.title
        t.left, t.top, t.width, t.height = MARGIN, Inches(0.34), CW, Inches(0.92)
        t.text_frame.word_wrap = True
        t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = t.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0] if p.runs else p.add_run()
        run.text = title
        run.font.name, run.font.size, run.font.bold, run.font.color.rgb = HEAD, Pt(title_size), True, INK
        self.text(s, kicker, MARGIN, Inches(0.10), Inches(9), Inches(0.3), 11, MUTED, bold=True)
        self.text(s, f"{self.n} / {N_TOTAL}", W - MARGIN - Inches(1.5), H - Inches(0.44), Inches(1.5),
                  Inches(0.3), 11, MUTED, align=PP_ALIGN.RIGHT)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def dark(self, notes: str):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = INK
        self.text(s, f"{self.n} / {N_TOTAL}", W - MARGIN - Inches(1.5), H - Inches(0.44), Inches(1.5),
                  Inches(0.3), 11, SLATE_L, align=PP_ALIGN.RIGHT)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def shield(self, code: str, title: str, question: str, notes: str):
        """Slide phụ lục: câu hỏi nó đỡ + các khối trả lời dày."""
        s = self.content(title, f"PHỤ LỤC {code} — CHỈ MỞ KHI BỊ HỎI", notes, title_size=28)
        self.text(s, "Trả lời cho: " + question, MARGIN, Inches(1.28), CW, Inches(0.4), 15, MUTED,
                  italic=True)
        return s

    # --- primitives ----------------------------------------------------------
    def text(self, slide, text, x, y, w, h, size=16, color=INK, bold=False, italic=False, font=BODY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if spacing:
                p.space_after = Pt(spacing)
            r = p.add_run()
            r.text = line
            r.font.name, r.font.size, r.font.bold, r.font.italic = font, Pt(size), bold, italic
            r.font.color.rgb = color
        return box

    def rich(self, slide, parts, x, y, w, h, size=15, color=INK, gap=6, lead_color=None):
        """parts = list các đoạn; mỗi đoạn = list[(kind, text)] — 'b' bold, 'o' bold orange."""
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.04)
        for i, para in enumerate(parts):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            pre = p.add_run()
            pre.text = "▸  "
            pre.font.name, pre.font.size, pre.font.bold = BODY, Pt(size), True
            pre.font.color.rgb = lead_color or ORANGE
            for kind, txt in para:
                r = p.add_run()
                r.text = txt
                r.font.name, r.font.size = BODY, Pt(size)
                r.font.bold = "b" in kind or "o" in kind
                r.font.color.rgb = ORANGE if "o" in kind else (TEAL if "t" in kind else color)
        return box

    def card(self, slide, x, y, w, h, fill=CARD, line=None):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shp.adjustments[0] = 0.055
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if line:
            shp.line.color.rgb = line
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def stat(self, slide, value, label, x, y, w, color=ORANGE, fill=CARD, h=Inches(1.55),
             value_size=40, label_size=13):
        # co cỡ chữ giá trị tới khi vừa một dòng trong thẻ (Cambria bold ~0.56em/ký tự)
        while value_size > 15 and len(value) * value_size * 0.0079 > w.inches - 0.44:
            value_size -= 1
        self.card(slide, x, y, w, h, fill)
        if h < Inches(1.3):  # bản compact: value + label ngồi khít
            self.text(slide, value, x + Inches(0.18), y + Inches(0.06), w - Inches(0.36), Inches(0.5),
                      value_size, color, bold=True, font=HEAD)
            self.text(slide, label, x + Inches(0.18), y + Inches(0.56), w - Inches(0.36), h - Inches(0.6),
                      label_size, MUTED)
        else:
            self.text(slide, value, x + Inches(0.2), y + Inches(0.10), w - Inches(0.4), Inches(0.82),
                      value_size, color, bold=True, font=HEAD)
            self.text(slide, label, x + Inches(0.2), y + Inches(0.90), w - Inches(0.4), h - Inches(0.95),
                      label_size, MUTED)

    def chip(self, slide, text, x, y, w, h=Inches(0.42), fill=INK, color=WHITE, size=13, bold=True):
        # chip tắt word-wrap nên phải tự co chữ để không tràn khỏi khung (bold ~0.55em/ký tự)
        w_in = Emu(int(w)).inches
        while size > 9 and len(text) * size * 0.0076 > w_in - 0.12:
            size -= 0.5
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shp.adjustments[0] = 0.5
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = BODY, Pt(size), bold, color
        return shp

    def badge(self, slide, number: str, x, y, color=ORANGE, d=Inches(0.46)):
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
        c.fill.solid()
        c.fill.fore_color.rgb = color
        c.line.fill.background()
        c.shadow.inherit = False
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = number
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = BODY, Pt(15), True, WHITE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    def hline(self, slide, x, y, w, color=SLATE_L):
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(0.75))
        ln.fill.solid()
        ln.fill.fore_color.rgb = color
        ln.line.fill.background()
        ln.shadow.inherit = False
        return ln

    def arrow_right(self, slide, x, y, w=Inches(0.32), color=MUTED):
        a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.22))
        a.fill.solid()
        a.fill.fore_color.rgb = color
        a.line.fill.background()
        a.shadow.inherit = False
        return a

    def picture(self, slide, path: Path, x, y, w=None, h=None):
        return slide.shapes.add_picture(str(path), x, y, width=w, height=h)

    def dtable(self, slide, header, rows, x, y, col_widths, row_h=Inches(0.55), size=14,
               highlight_row=None, header_h=Inches(0.5)):
        """Bảng vẽ bằng shapes: hàng cố định, không tự giãn, không dính table style mặc định."""
        total_w = Inches(sum(col_widths))
        self.card(slide, x, y, total_w, header_h, INK)
        cx = x
        for c, wd in enumerate(col_widths):
            self.text(slide, header[c], cx + Inches(0.14), y, Inches(wd - 0.24), header_h, size - 0.5,
                      WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE,
                      align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
            cx += Inches(wd)
        ry = y + header_h + Inches(0.07)
        for r, row in enumerate(rows):
            fill = SOFT_ORANGE if highlight_row == r else (CARD if r % 2 == 0 else WHITE)
            if fill == WHITE:
                self.card(slide, x, ry, total_w, row_h, WHITE, line=RGBColor(0xE2, 0xE8, 0xF0))
            else:
                self.card(slide, x, ry, total_w, row_h, fill)
            cx = x
            for c, wd in enumerate(col_widths):
                val = str(row[c])
                fs = size - 1
                # co chữ tới khi vừa tối đa 2 dòng trong ô
                while fs > 10.5 and len(val) * fs * 0.0071 > (wd - 0.24) * 2:
                    fs -= 0.5
                self.text(slide, val, cx + Inches(0.14), ry, Inches(wd - 0.24), row_h, fs, INK,
                          bold=(highlight_row == r), anchor=MSO_ANCHOR.MIDDLE,
                          align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
                cx += Inches(wd)
            ry += row_h + Inches(0.07)
        return ry

    def style_chart(self, chart, colors, legend=True):
        chart.font.name = BODY
        chart.font.size = Pt(13)
        chart.font.color.rgb = INK
        chart.has_legend = legend
        if legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(13)
        for series, color in zip(chart.plots[0].series, colors):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(12)
        va.tick_labels.font.color.rgb = MUTED
        ca = chart.category_axis
        ca.tick_labels.font.size = Pt(13)
        ca.tick_labels.font.color.rgb = INK
        ca.format.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    def labels_on(self, chart, fmt="0.000", size=12):
        pl = chart.plots[0]
        pl.has_data_labels = True
        pl.data_labels.number_format = fmt
        pl.data_labels.number_format_is_linked = False
        pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        pl.data_labels.font.size = Pt(size)
        return pl


# --------------------------------------------------------------------------- build
def build(f: Facts) -> Presentation:
    d = Deck()
    loc = f.test(f.locked)
    per = f.test("persistence")

    # ---- S1 · bìa (dark) ------------------------------------------------------
    s = d.dark('10". Chào thầy/cô và các bạn; giới thiệu nhóm và tên đề tài: "Kho dữ liệu chất lượng '
               'không khí Bắc Kinh và cảnh báo sớm PM2.5 vượt ngưỡng."')
    d.text(s, "INT1422 · KHO DỮ LIỆU VÀ KHAI PHÁ DỮ LIỆU · BẢO VỆ 30 PHÚT", MARGIN, Inches(0.95),
           Inches(11), Inches(0.4), 14, AMBER, bold=True)
    d.text(s, ["Kho dữ liệu chất lượng không khí Bắc Kinh", "& cảnh báo sớm PM2.5 vượt ngưỡng"],
           MARGIN, Inches(1.55), Inches(12.1), Inches(2.2), 38, WHITE, bold=True, font=HEAD)
    facts = [(vn_int(420768), "giờ đo"), ("12", "trạm quan trắc"), ("4 năm", "03/2013 – 02/2017"),
             ("4 chiều", "lược đồ hình sao"), ("Giờ tới", "phân loại cảnh báo PM2.5")]
    fw = Emu(int(CW) // 5)
    for i, (fv, flab) in enumerate(facts):
        fx = MARGIN + i * fw
        if i:
            sep = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, fx - Emu(9525 // 2), Inches(4.06),
                                     Pt(0.9), Inches(0.82))
            sep.fill.solid()
            sep.fill.fore_color.rgb = RGBColor(0x3A, 0x48, 0x5C)
            sep.line.fill.background()
            sep.shadow.inherit = False
        d.text(s, fv, fx, Inches(4.02), fw, Inches(0.55), 28, WHITE, bold=True, font=HEAD,
               align=PP_ALIGN.CENTER)
        d.text(s, flab, fx, Inches(4.62), fw, Inches(0.38), 12.5, SLATE_L, align=PP_ALIGN.CENTER)
    d.hline(s, MARGIN, Inches(5.45), Inches(5.2), RGBColor(0x3A, 0x48, 0x5C))
    d.text(s, "NHÓM THỰC HIỆN", MARGIN, Inches(5.62), Inches(5), Inches(0.3), 11, AMBER, bold=True)
    members = [("Phạm Nhật Khánh", "B23DCCC088"), ("Phạm Xuân Kỳ", "B23DCCC092"),
               ("Nguyễn Đức Thành", "B23DCCC152")]
    for i, (name, mssv) in enumerate(members):
        mx = MARGIN + i * Inches(3.6)
        d.text(s, name, mx, Inches(5.98), Inches(3.5), Inches(0.35), 16, WHITE, bold=True)
        d.text(s, mssv, mx, Inches(6.36), Inches(3.5), Inches(0.3), 12, SLATE_L)
    d.text(s, "Giảng viên hướng dẫn: TS. Phan Lý Huỳnh · Học viện Công nghệ Bưu chính Viễn thông",
           MARGIN, Inches(6.78), Inches(10), Inches(0.35), 13, SLATE_L)

    # ---- S2 · câu hỏi nghiệp vụ ------------------------------------------------
    s = d.content("Giờ tới có cần phát cảnh báo không?", "A · VÌ SAO & DỮ LIỆU",
                  "1'. Câu mở: đề tài bắt đầu từ một câu hỏi rất cụ thể của người vận hành. Trung tâm quan trắc "
                  "mỗi giờ phải quyết: có phát cảnh báo hạn chế hoạt động ngoài trời ở trạm này cho GIỜ TỚI không? "
                  "Cách làm hiện tại là chờ số đo vượt ngưỡng rồi mới báo — luôn trễ 1 giờ. Quy mô: PM2.5 trung "
                  "bình 79,8 µg/m³, khoảng 1/4 số giờ ≥ 115. Chuyển: để trả lời tin được, trước hết phải hiểu dữ liệu.",
                  title_size=34)
    d.card(s, MARGIN, Inches(1.5), Inches(6.6), Inches(1.35), SOFT_ORANGE)
    d.text(s, "“Giờ tới, tại trạm này, PM2.5 có ≥ 115 µg/m³ không?”", MARGIN + Inches(0.3),
           Inches(1.62), Inches(6.0), Inches(1.1), 24, INK, bold=True, italic=True, font=HEAD,
           anchor=MSO_ANCHOR.MIDDLE)
    d.rich(s, [
        [("b", "Ai quyết: "), ("", "trung tâm quan trắc — mỗi giờ, từng trạm")],
        [("b", "Quyết định: "), ("", "phát / không phát cảnh báo hạn chế hoạt động ngoài trời")],
        [("b", "Hiện tại: "), ("", "chờ số đo vượt ngưỡng rồi báo → luôn trễ 1 giờ")],
    ], MARGIN + Inches(0.05), Inches(3.15), Inches(6.5), Inches(2.2), 17, gap=12)
    d.stat(s, "79,8 µg/m³", "PM2.5 trung bình 4 năm · 12 trạm", Inches(7.6), Inches(1.5), Inches(5.13),
           h=Inches(1.5))
    d.stat(s, "≈ 1/4", "số giờ có PM2.5 ≥ 115 µg/m³ — quy mô của vấn đề", Inches(7.6), Inches(3.12),
           Inches(5.13), color=INK, h=Inches(1.5))
    d.stat(s, "Trễ 1 giờ", 'bản chất của cách làm "chờ số đo rồi báo"', Inches(7.6), Inches(4.74),
           Inches(5.13), color=INK, fill=SOFT_ORANGE, h=Inches(1.5))

    # ---- S3 · hai nửa môn học ---------------------------------------------------
    s = d.content("Hai nửa của môn học trong một hệ thống", "A · VÌ SAO & DỮ LIỆU",
                  "1'. Đề tài may mắn ôm trọn hai nửa học phần: kho dữ liệu (OLAP — phân tích quá khứ mọi góc "
                  "nhìn) và khai phá (học từ quá khứ để dự đoán giờ tới). Sơ đồ dưới là toàn luồng: nguồn → lake → "
                  "staging → kho → mart → khai phá → dashboard. Chuyển: đi từ nguồn dữ liệu.")
    d.card(s, MARGIN, Inches(1.5), Inches(5.9), Inches(3.1), SOFT_TEAL)
    d.text(s, "NỬA KHO DỮ LIỆU", MARGIN + Inches(0.3), Inches(1.72), Inches(5.3), Inches(0.5), 20,
           TEAL, bold=True)
    d.text(s, ["Tổ chức 4 năm dữ liệu — phân tích quá khứ", "mọi góc nhìn (OLAP: slice · dice · drill · "
               "roll · pivot)."], MARGIN + Inches(0.3), Inches(2.35), Inches(5.3), Inches(1.9), 18, INK,
           spacing=8)
    d.card(s, Inches(6.83), Inches(1.5), Inches(5.9), Inches(3.1), SOFT_ORANGE)
    d.text(s, "NỬA KHAI PHÁ DỮ LIỆU", Inches(7.13), Inches(1.72), Inches(5.3), Inches(0.5), 20,
           ORANGE, bold=True)
    d.text(s, ["Học từ quá khứ để dự đoán giờ tới —", "phân loại cảnh báo sớm PM2.5 giờ sau."],
           Inches(7.13), Inches(2.35), Inches(5.3), Inches(1.9), 18, INK, spacing=8)
    labels = ["Nguồn", "Lake", "Staging", "Kho", "Mart", "Khai phá", "Dashboard"]
    x = MARGIN
    for i, lab in enumerate(labels):
        wch = Inches(1.38)
        d.chip(s, lab, x, Inches(5.15), wch, fill=INK if i in (0, 6) else TEAL if i in (1, 2) else ORANGE)
        if i < len(labels) - 1:
            d.arrow_right(s, x + wch + Inches(0.03), Inches(5.25), w=Inches(0.3))
        x += wch + Inches(0.36)
    d.text(s, "Sơ đồ hai luồng vào — slide 7.", MARGIN, Inches(5.95), Inches(9), Inches(0.4),
           14, MUTED, italic=True)

    # ---- S4 · hai nguồn ---------------------------------------------------------
    s = d.content("Dữ liệu: hai nguồn mở, có giấy phép", "A · VÌ SAO & DỮ LIỆU",
                  "1'. Nguồn chính UCI #501: 420.768 bản ghi giờ, 12 trạm, 03/2013–02/2017, 12 thuộc tính "
                  "(6 chất ô nhiễm + 5 khí tượng + hướng gió), CC BY 4.0, có DOI. Nguồn phụ Open-Meteo (tái phân "
                  "tích ERA5, theo giờ, cùng 12 toạ độ): lấp chỗ thiếu thời tiết + thêm độ che phủ mây. Không có "
                  "dữ liệu cá nhân. MỒI → LÁ CHẮN: bị hỏi «hai nguồn ghép thế nào?» → A2 (slide 34).")
    for i, (v, lab) in enumerate(((vn_int(420768), "bản ghi giờ"), ("12", "trạm quan trắc"),
                                  ("4 năm", "03/2013 – 02/2017"), ("12", "thuộc tính đo"))):
        d.stat(s, v, lab, MARGIN + i * Inches(3.08), Inches(1.5), Inches(2.85),
               color=ORANGE if i == 0 else INK, h=Inches(1.4))
    d.card(s, MARGIN, Inches(3.2), Inches(5.9), Inches(2.9))
    d.text(s, "UCI #501 — Beijing Multi-Site Air Quality", MARGIN + Inches(0.3), Inches(3.38),
           Inches(5.4), Inches(0.5), 18, TEAL, bold=True)
    d.rich(s, [
        [("b", "6 chất ô nhiễm: "), ("", "PM2.5, PM10, SO2, NO2, CO, O3")],
        [("b", "5 biến khí tượng "), ("", "+ hướng gió (16 hướng)")],
        [("b", "CC BY 4.0 · DOI 10.24432/C5RK5G")],
    ], MARGIN + Inches(0.25), Inches(3.95), Inches(5.45), Inches(2.0), 16, gap=10)
    d.card(s, Inches(6.83), Inches(3.2), Inches(5.9), Inches(2.9))
    d.text(s, "Open-Meteo — Historical Weather (ERA5)", Inches(7.13), Inches(3.38), Inches(5.4),
           Inches(0.5), 18, TEAL, bold=True)
    d.rich(s, [
        [("b", "Theo giờ, cùng 12 toạ độ "), ("", "với UCI")],
        [("b", "Vai trò: "), ("", "lấp chỗ thiếu thời tiết + thêm độ che phủ mây")],
        [("b", "CC BY 4.0 ")],
    ], Inches(7.08), Inches(3.95), Inches(5.45), Inches(2.0), 16, gap=10)

    # ---- S5 · chất lượng dữ liệu -------------------------------------------------
    s = d.content("Chất lượng dữ liệu: quyết định từng cột, không làm đẹp", "A · VÌ SAO & DỮ LIỆU",
                  "1'30\". Làm sạch là một chuỗi quyết định, không phải làm bảng đẹp. CO thiếu nhiều nhất 4,9%, "
                  "chuỗi thiếu liên tục dài tới 1.517 giờ → KHÔNG nội suy (nội suy chuỗi dài = bịa dữ liệu), giữ "
                  "NULL + cờ chất lượng. Thời tiết thiếu ~0,1% → lấp từ Open-Meteo, ghi nguồn từng ô, còn 0 ô "
                  "thiếu. 60 dòng chạm trần thiết bị → giữ + cờ SUSPECT_CAP. ML: median điền thiếu chỉ tính trên "
                  "tập train. MỒI → LÁ CHẮN: «vì sao không nội suy hết cho đẹp?» → A3 (slide 35).",
                  title_size=30)
    cd = CategoryChartData()
    miss = f.missing.sort_values("missing_pct")
    cd.categories = list(miss["variable"])
    cd.add_series("% thiếu", [round(v, 2) for v in miss["missing_pct"]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MARGIN, Inches(1.42), Inches(6.3),
                            Inches(5.35), cd)
    ch = gf.chart
    d.style_chart(ch, [ORANGE], legend=False)
    ch.has_title = True
    ch.chart_title.text_frame.text = "% giá trị thiếu theo cột (gộp 12 trạm)"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    pl = d.labels_on(ch, '0.0"%"', 11)
    pl.gap_width = 60
    d.rich(s, [
        [("b", "CO thiếu nhiều nhất 4,9% "), ("", "— chuỗi thiếu liên tục 343–1.517 giờ")],
        [("o", "Không nội suy: "), ("", "giữ NULL + cờ chất lượng, COUNT(cột) ≠ COUNT(*)")],
        [("b", "Thời tiết thiếu ~0,1% "), ("", "→ lấp từ Open-Meteo, ghi nguồn từng ô → còn 0 ô thiếu")],
        [("b", "60 dòng chạm trần "), ("", "(999 / 10.000) → giữ lại + cờ SUSPECT_CAP")],
        [("b", "Nguyên tắc ML: "), ("", "median điền thiếu chỉ tính trên tập train")],
    ], Inches(7.15), Inches(1.55), Inches(5.55), Inches(5.1), 16, gap=13)

    # ---- S6 · rời rạc hoá ----------------------------------------------------------
    s = d.content("Rời rạc hoá cho phân tích", "A · VÌ SAO & DỮ LIỆU",
                  "30\". Giờ → buổi, ngày → mùa, 16 hướng gió → 4 phần tư → bắc/nam. Đây là các phân cấp sẽ dùng "
                  "lại toàn bộ ở phần kho dữ liệu (dim_date, dim_time, dim_wind).")
    rows = [("Giờ", "→", "Buổi", "sáng · trưa · chiều · tối", "dim_time"),
            ("Ngày", "→", "Mùa", "Xuân · Hạ · Thu · Đông", "dim_date"),
            ("16 hướng gió", "→", "4 phần tư → bắc/nam", "NE · SE · SW · NW → bắc / nam", "dim_wind")]
    for i, (a, ar, b, desc, dim) in enumerate(rows):
        y = Inches(1.65) + i * Inches(1.45)
        d.card(s, MARGIN, y, Inches(12.13), Inches(1.15))
        d.chip(s, a, MARGIN + Inches(0.25), y + Inches(0.36), Inches(1.9), fill=MUTED)
        d.arrow_right(s, MARGIN + Inches(2.3), y + Inches(0.46))
        d.chip(s, b, MARGIN + Inches(2.8), y + Inches(0.36), Inches(3.1), fill=ORANGE)
        d.text(s, desc, MARGIN + Inches(6.2), y + Inches(0.12), Inches(4.3), Inches(0.95), 16, INK,
               anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, dim, MARGIN + Inches(10.6), y + Inches(0.12), Inches(1.4), Inches(0.95), 14, TEAL,
               bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    d.text(s, "Các phân cấp này dùng lại toàn bộ ở phần kho dữ liệu (đúng các bảng chiều).", MARGIN,
           Inches(6.25), Inches(11), Inches(0.4), 15, MUTED, italic=True)

    # ---- S7 · kiến trúc tổng ---------------------------------------------------------
    s = d.content("Kiến trúc tổng: hai luồng vào, một kho", "A · VÌ SAO & DỮ LIỆU",
                  "1'30\". Chỉ theo sơ đồ. Luồng 1 (batch): CSV UCI + Open-Meteo → MinIO data lake (raw → "
                  "processed Parquet). Luồng 2 (tác nghiệp): SQLite mô phỏng nguồn giao dịch → CDC theo watermark "
                  "→ Airflow. Sau đó PostgreSQL staging → dbt dựng kho (fact + dimension) và các mart → đặc trưng "
                  "ML → dashboard FastAPI + React. Airflow điều phối, Docker đóng gói hạ tầng. MỒI → LÁ CHẮN: "
                  "«lake/staging/mart khác nhau thế nào?», «4 công cụ kia đều là database à?» → A4 (slide 36).",
                  title_size=30)
    # Sơ đồ hai luồng — vẽ bằng shapes cho khớp tông deck (không dùng ảnh matplotlib)
    d.chip(s, "LUỒNG 1 · BATCH", MARGIN, Inches(1.34), Inches(2.2), h=Inches(0.3), fill=INK, size=11)
    lane1 = [
        ("Nguồn batch", ["UCI CSV × 12", "Open-Meteo (ERA5)"], MUTED),
        ("MinIO — data lake", ["raw giữ nguyên", "processed Parquet", "cờ chất lượng"], TEAL),
        ("Staging — PostgreSQL", ["staging.*", "đối soát nguồn"], TEAL),
        ("Kho — dbt", ["sao 4 chiều", "130 dbt test"], ORANGE),
        ("Mart & khối", ["trạm×ngày · ×tháng", "CUBE 32 khối"], ORANGE),
        ("ML & Dashboard", ["FastAPI + React", "/olap · /canh-bao"], INK),
    ]
    for i, (name, sub, color) in enumerate(lane1):
        bx = MARGIN + Inches(i * 2.07)
        d.card(s, bx, Inches(1.74), Inches(1.78), Inches(1.48), CARD)
        d.chip(s, name, bx + Inches(0.08), Inches(1.84), Inches(1.62), h=Inches(0.4), fill=color,
               size=11.5)
        d.text(s, sub, bx + Inches(0.06), Inches(2.32), Inches(1.66), Inches(0.85), 10.5, INK,
               align=PP_ALIGN.CENTER)
        if i < len(lane1) - 1:
            d.arrow_right(s, bx + Inches(1.795), Inches(2.38), w=Inches(0.26))
    d.chip(s, "LUỒNG 2 · TÁC NGHIỆP", MARGIN, Inches(4.04), Inches(2.6), h=Inches(0.3), fill=TEAL,
           size=11)
    d.card(s, MARGIN, Inches(4.42), Inches(2.85), Inches(1.12), CARD)
    d.chip(s, "SQLite — tác nghiệp", MARGIN + Inches(0.08), Inches(4.52), Inches(2.69),
           h=Inches(0.38), fill=MUTED, size=11.5)
    d.text(s, ["mô phỏng nguồn giao dịch"], MARGIN + Inches(0.06), Inches(4.98), Inches(2.73),
           Inches(0.45), 10.5, INK, align=PP_ALIGN.CENTER)
    d.arrow_right(s, MARGIN + Inches(2.95), Inches(4.86), w=Inches(0.3))
    d.card(s, MARGIN + Inches(3.4), Inches(4.42), Inches(3.0), Inches(1.12), CARD)
    d.chip(s, "CDC theo watermark", MARGIN + Inches(3.48), Inches(4.52), Inches(2.84),
           h=Inches(0.38), fill=INK, size=11.5)
    d.text(s, ["Airflow DAG · chỉ lấy bản ghi", "sau mốc event_time"], MARGIN + Inches(3.46),
           Inches(4.96), Inches(2.88), Inches(0.52), 10.5, INK, align=PP_ALIGN.CENTER)
    up = s.shapes.add_shape(MSO_SHAPE.UP_ARROW, MARGIN + Inches(4.9), Inches(3.30), Inches(0.3),
                            Inches(1.06))
    up.fill.solid()
    up.fill.fore_color.rgb = TEAL
    up.line.fill.background()
    up.shadow.inherit = False
    d.text(s, ["CDC nạp vào lake → staging bằng upsert theo khóa grain;",
               "load_registry (object key + sha256 + row_count) — chạy lại không nhân đôi."],
           MARGIN + Inches(6.7), Inches(4.5), Inches(5.4), Inches(1.0), 12.5, MUTED, spacing=6)
    d.text(s, "Airflow điều phối cả hai luồng · Docker đóng gói hạ tầng.", MARGIN, Inches(5.9),
           Inches(12.0), Inches(0.4), 14, MUTED, italic=True)

    # ---- S8 · tái lập ------------------------------------------------------------------
    s = d.content("Tái lập được bằng một lệnh", "A · VÌ SAO & DỮ LIỆU",
                  "30\". python scripts/run_all.py — 14 bước, chạy lại từ trạng thái trống mất 11,2 phút, mọi bước "
                  "idempotent (chạy lại không nhân đôi dữ liệu). Lời cam kết: mọi con số hôm nay ai cũng tạo lại "
                  "được.")
    d.card(s, MARGIN, Inches(1.7), Inches(12.13), Inches(1.5), INK)
    d.text(s, "python scripts/run_all.py", MARGIN + Inches(0.4), Inches(1.85), Inches(11.3),
           Inches(1.2), 34, WHITE, bold=True, font="Courier New", anchor=MSO_ANCHOR.MIDDLE)
    for i, (v, lab) in enumerate((("14", "bước — thu thập tới dashboard"), (f"{vn(FROM_SCRATCH_MINUTES, 1)} phút",
                                 "chạy lại từ trạng thái trống"), (f"{f.dbt_pass}/{f.dbt_total}", "dbt test pass"),
                                 ("Idempotent", "chạy lại không nhân đôi dữ liệu"))):
        d.stat(s, v, lab, MARGIN + i * Inches(3.08), Inches(3.6), Inches(2.85), h=Inches(1.6),
               color=ORANGE if i in (0, 1) else INK, value_size=34 if i != 3 else 28)
    d.text(s, "Cam kết: mọi con số trong bài thuyết trình này đều tạo lại được bằng một lệnh.", MARGIN,
           Inches(5.6), Inches(11.5), Inches(0.5), 18, INK, italic=True)

    # ---- S9 · chuyển phần B (dark) ------------------------------------------------------
    s = d.dark('10". NÓI: "Dữ liệu đã nằm đúng chỗ. Câu hỏi tiếp: tổ chức thế nào để phân tích được mọi góc '
               'nhìn mà vẫn kiểm chứng được? Đó là phần kho dữ liệu." → trao cho người B.')
    d.text(s, "PHẦN B · KHO DỮ LIỆU & OLAP", MARGIN, Inches(2.2), Inches(11), Inches(0.4), 15, AMBER,
           bold=True)
    d.text(s, ["Dữ liệu đã nằm đúng chỗ.", "Tổ chức thế nào để phân tích được mọi góc nhìn",
               "mà vẫn kiểm chứng được?"], MARGIN, Inches(2.75), Inches(11.8), Inches(2.6), 30, WHITE,
           bold=True, font=HEAD, spacing=6)

    # ---- S10 · grain ---------------------------------------------------------------------
    s = d.content("Grain: một dòng fact = một trạm × một giờ", "B · KHO DỮ LIỆU & OLAP",
                  "1'. Quyết định quan trọng nhất khi thiết kế kho là phát biểu grain. Chọn độ mịn gốc vì: cuộn "
                  "từ giờ lên ngày/tháng được, ngược lại không; và bài toán cảnh báo giờ tới cần đúng mức giờ. Mọi "
                  "độ đo, mọi khóa chiều phải khớp grain này. MỒI → LÁ CHẮN: «vì sao không lưu mức ngày cho nhẹ?» "
                  "→ A5 (slide 37).")
    d.card(s, MARGIN, Inches(1.55), Inches(12.13), Inches(1.5), INK)
    d.text(s, "1 dòng fact  =  1 trạm  ×  1 giờ", MARGIN, Inches(1.7), Inches(12.13), Inches(1.2), 36,
           WHITE, bold=True, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    d.card(s, MARGIN, Inches(3.4), Inches(5.9), Inches(2.55), SOFT_TEAL)
    d.text(s, "Cuộn được một chiều", MARGIN + Inches(0.3), Inches(3.58), Inches(5.3), Inches(0.5), 18,
           TEAL, bold=True)
    d.text(s, ["giờ → ngày → tháng → năm: tổng hợp được;", "ngược lại (ngày → giờ) mất thông tin vĩnh "
               "viễn."], MARGIN + Inches(0.3), Inches(4.12), Inches(5.3), Inches(1.6), 16, INK, spacing=6)
    d.card(s, Inches(6.83), Inches(3.4), Inches(5.9), Inches(2.55), SOFT_TEAL)
    d.text(s, "Bài toán cần mức giờ", Inches(7.13), Inches(3.58), Inches(5.3), Inches(0.5), 18, TEAL,
           bold=True)
    d.text(s, ["cảnh báo phát cho giờ tới (t+1h) —", "độ mịn gốc phải là trạm × giờ."],
           Inches(7.13), Inches(4.12), Inches(5.3), Inches(1.6), 16, INK, spacing=6)
    d.text(s, "Mọi độ đo và mọi khóa chiều của kho đều phải khớp grain này.", MARGIN, Inches(6.2),
           Inches(11), Inches(0.4), 15, MUTED, italic=True)

    # ---- S11 · sao 4 chiều ------------------------------------------------------------------
    s = d.content("Lược đồ hình sao: fact trung tâm + 4 chiều", "B · KHO DỮ LIỆU & OLAP",
                  "1'. Fact_air_quality_hourly ở giữa + 4 chiều: ngày, giờ, trạm, hướng gió. Hướng gió là chiều "
                  "đặc biệt — trả lời câu hỏi «VÌ SAO ô nhiễm đến» (điều kiện vận tải khí). Thời tiết không làm "
                  "chiều riêng mà là độ đo trên fact. MỒI → LÁ CHẮN: «sao không có chiều thời tiết?» → A6 "
                  "(slide 38).")
    # Sơ đồ hình sao vẽ bằng shapes — fact giữa, 2 chiều mỗi bên, đúng palette deck
    def dim_box(slide, title, rows, x, y, w=Inches(2.6)):
        h = Inches(0.44) + Inches(0.32 * len(rows))
        d.card(slide, x, y, w, h, WHITE, line=RGBColor(0xE2, 0xE8, 0xF0))
        d.chip(slide, title, x + Inches(0.07), y + Inches(0.07), w - Inches(0.14), h=Inches(0.38),
               fill=TEAL, size=12)
        d.text(slide, rows, x + Inches(0.18), y + Inches(0.54), w - Inches(0.32),
               Inches(0.32 * len(rows)), 12, INK, spacing=2)

    d.card(s, Inches(4.1), Inches(2.15), Inches(5.13), Inches(2.9), SOFT_ORANGE)
    d.chip(s, "fact_air_quality_hourly", Inches(4.2), Inches(2.27), Inches(4.93), h=Inches(0.42),
           fill=ORANGE, size=12)
    d.text(s, ["KHÓA: station_key · date_key", "time_key · wind_key",
               "ĐỘ ĐO: pm25 · pm10 · so2 · no2", "co · o3 · temp · pres · dewp",
               "rain · wspm · cờ chất lượng"],
           Inches(4.32), Inches(2.82), Inches(4.69), Inches(2.0), 12.5, INK, spacing=6)
    dim_box(s, "dim_date", ["date_key (PK)", "năm · tháng · quý · mùa", "cuối tuần"], MARGIN,
            Inches(1.7))
    dim_box(s, "dim_time", ["time_key (PK)", "giờ · buổi"], MARGIN, Inches(4.1))
    dim_box(s, "dim_station", ["station_key (PK)", "trạm · quận"], Inches(10.13), Inches(1.7))
    dim_box(s, "dim_wind", ["wind_key (PK)", "16 hướng · phần tư", "chế độ bắc / nam"],
            Inches(10.13), Inches(4.1))
    d.arrow_right(s, Inches(3.28), Inches(2.3))
    d.arrow_right(s, Inches(3.28), Inches(4.45))
    d.arrow_right(s, Inches(9.3), Inches(2.3))
    d.arrow_right(s, Inches(9.3), Inches(4.45))
    d.rich(s, [
        [("o", "Hướng gió = chiều «vì sao»"), ("", " — điều kiện vận tải khí")],
        [("b", "Thời tiết: "), ("", "độ đo trên fact, không làm chiều riêng")],
    ], MARGIN, Inches(5.75), Inches(12.0), Inches(0.95), 15, gap=6)

    # ---- S12 · 180 tổ hợp -------------------------------------------------------------------
    s = d.content("Phân cấp chiều → 180 tổ hợp mức", "B · KHO DỮ LIỆU & OLAP",
                  "1'. 4 phân cấp: ngày→tháng→quý→năm; giờ→buổi; trạm→quận; hướng→phần tư→bắc/nam. Tổng tổ hợp "
                  "mức 5·3·3·4 = 180 (mỗi chiều tính cả mức ALL) — không thể, cũng không cần, tính trước hết; "
                  "chọn theo truy vấn thật.")
    dims = [("dim_date", "ngày → tháng → quý → năm", "5"),
            ("dim_time", "giờ → buổi", "3"),
            ("dim_station", "trạm → quận", "3"),
            ("dim_wind", "hướng → phần tư → bắc/nam", "4")]
    for i, (name, chain, cnt) in enumerate(dims):
        y = Inches(1.55) + i * Inches(1.06)
        d.card(s, MARGIN, y, Inches(7.9), Inches(0.88))
        d.text(s, name, MARGIN + Inches(0.25), y + Inches(0.08), Inches(1.9), Inches(0.7), 15, TEAL,
               bold=True, font="Courier New", anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, chain, MARGIN + Inches(2.2), y + Inches(0.08), Inches(4.4), Inches(0.7), 16, INK,
               anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, f"+ ALL = {cnt} mức", MARGIN + Inches(6.5), y + Inches(0.08), Inches(1.3), Inches(0.7),
               15, ORANGE, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    d.stat(s, "5·3·3·4 = 180", "tổ hợp mức phân cấp (mỗi chiều tính cả mức ALL)", Inches(8.85),
           Inches(1.55), Inches(3.88), h=Inches(2.0), value_size=32)
    d.text(s, ["Không cần tính trước hết 180 tổ hợp —", "chỉ tính trước cái dashboard thật đang dùng "
              "(slide 13)."], Inches(8.85), Inches(3.8), Inches(3.88), Inches(1.6), 15, MUTED,
           spacing=6)

    # ---- S13 · ba khối tính trước ------------------------------------------------------------
    s = d.content("Ba khối tính trước theo truy vấn thật", "B · KHO DỮ LIỆU & OLAP",
                  "1'. mart_station_daily (trạm, ngày) 17.532 dòng — nhỏ hơn fact 24 lần, cho xu hướng; "
                  "mart_station_monthly 576 dòng — báo cáo tháng/mùa; mart_olap_cube 6.075 dòng — trang OLAP đọc "
                  "vài ms. Tiêu chí chọn: đúng các truy vấn dashboard thật đang phát ra. MỒI → LÁ CHẮN: «tính "
                  "trước thì có trả lời được mọi truy vấn không?» → A8 cuối (slide 40).")
    rows = [("mart_station_daily", "trạm × ngày", "17.532", "xu hướng (nhỏ hơn fact 24×)"),
            ("mart_station_monthly", "trạm × tháng", "576", "báo cáo tháng / mùa"),
            ("mart_olap_cube", "CUBE 5 thuộc tính = 32 khối", "6.075", "trang OLAP đọc vài ms")]
    d.dtable(s, ("Khối tính trước", "Khóa / khối con", "Số dòng", "Phục vụ"), rows, MARGIN,
             Inches(1.6), [2.3, 3.1, 1.0, 3.0], row_h=Inches(0.62), size=14)
    d.stat(s, "24×", f"fact {vn_int(420768)} dòng → mart daily 17.532 dòng", Inches(10.2), Inches(1.6),
           Inches(2.53), color=TEAL, fill=SOFT_TEAL, h=Inches(1.6), value_size=30, label_size=12)
    d.card(s, MARGIN, Inches(5.2), Inches(12.13), Inches(1.3), SOFT_ORANGE)
    d.text(s, "Tiêu chí chọn khối: đúng các truy vấn mà dashboard thật đang phát ra — không tính trước "
              "theo cảm tính.", MARGIN + Inches(0.3), Inches(5.32), Inches(11.5), Inches(1.05), 18, INK,
           bold=True, anchor=MSO_ANCHOR.MIDDLE)

    # ---- S14 · CUBE 32 -------------------------------------------------------------------------
    s = d.content("CUBE 5 thuộc tính = 2⁵ = 32 grouping sets", "B · KHO DỮ LIỆU & OLAP",
                  "1'. Cube tổng hợp theo CUBE(năm, mùa, quận, buổi, phần tư gió): 5 thuộc tính độc lập → 2⁵ = 32 "
                  "grouping sets. Nhấn mạnh: số thuộc tính trong CUBE KHÔNG BẰNG số bảng chiều — 2 thuộc tính có "
                  "thể cùng nằm trong một chiều (năm + mùa cùng dim_date). MỒI → LÁ CHẮN: «4 chiều sao ra 32 "
                  "khối?» → A7 (slide 39).")
    attrs = ["năm", "mùa", "quận", "buổi", "phần tư gió"]
    x = MARGIN
    for i, a in enumerate(attrs):
        d.chip(s, a, x, Inches(1.65), Inches(1.75), fill=INK, size=15)
        if i < 4:
            d.text(s, "×2", x + Inches(1.78), Inches(1.68), Inches(0.62), Inches(0.38), 14, MUTED,
                   bold=True, align=PP_ALIGN.CENTER)
        x += Inches(2.4)
    d.text(s, "mỗi thuộc tính: giữ nguyên giá trị  hoặc  cuộn lên ALL → 2 lựa chọn", MARGIN, Inches(2.35),
           Inches(11.5), Inches(0.45), 17, INK)
    d.stat(s, "2⁵ = 32", "grouping sets tính trước trong mart_olap_cube", MARGIN, Inches(3.05),
           Inches(3.9), h=Inches(1.9), value_size=54)
    d.card(s, Inches(4.9), Inches(3.05), Inches(7.83), Inches(1.9), SOFT_ORANGE)
    d.text(s, ["Lưu ý hay bị hỏi:", "số thuộc tính trong CUBE ≠ số bảng chiều —", "năm + mùa cùng nằm "
              "trong dim_date."], Inches(5.2), Inches(3.25), Inches(7.2), Inches(1.55), 17, INK,
           spacing=6)
    d.text(s, "GROUPING SETS: mỗi grouping set = một cách chọn (giữ/ALL) cho 5 thuộc tính → 2×2×2×2×2 = 32 "
              "khối con, từ tổng thể tới từng tổ hợp chi tiết.", MARGIN, Inches(5.25), Inches(12.0),
           Inches(0.9), 15, MUTED)

    # ---- S15 · kiểm chứng -------------------------------------------------------------------------
    s = d.content("Kiểm chứng 100%: tính lại từng ô từ fact", "B · KHO DỮ LIỆU & OLAP",
                  "1'. Phần tự hào nhất của kho: script tính lại TỪNG Ô của các khối trực tiếp từ fact rồi so khớp "
                  "— 24.759/24.759 ô đúng = 100%. Cộng 130 dbt test (khóa, not-null, khóa ngoại, cube phân hoạch "
                  "fact). Truy vấn viết lại: thuộc tính nằm trong cube → đọc mart 2–6 ms; thiếu → quét fact "
                  "130–400 ms. MỒI → LÁ CHẮN: «kiểm chứng cụ thể thế nào?» → A8 (slide 40).")
    s_ = f.olap["summary"]
    d.stat(s, "100%", f"{vn_int(s_['cells_matched'])}/{vn_int(s_['cells_checked'])} ô của khối tính "
           "trước khớp khi tính lại từ fact", MARGIN, Inches(1.55), Inches(4.0), color=TEAL,
           fill=SOFT_TEAL, h=Inches(2.0), value_size=54)
    d.stat(s, f"{f.dbt_pass}/{f.dbt_total}", "dbt test: khóa duy nhất · not-null · khóa ngoại · cube "
           "phân hoạch fact", Inches(4.8), Inches(1.55), Inches(4.0), h=Inches(2.0), value_size=44)
    d.text(s, "Truy vấn được viết lại theo khối:", Inches(9.1), Inches(1.7), Inches(3.6), Inches(0.5),
           16, INK, bold=True)
    d.card(s, Inches(9.1), Inches(2.2), Inches(3.6), Inches(0.75), SOFT_TEAL)
    d.text(s, "đủ thuộc tính trong cube → 2–6 ms", Inches(9.3), Inches(2.28), Inches(3.3), Inches(0.6),
           15, TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    d.card(s, Inches(9.1), Inches(3.05), Inches(3.6), Inches(0.75), SOFT_ORANGE)
    d.text(s, "thiếu 1 thuộc tính → quét fact 130–400 ms", Inches(9.3), Inches(3.13), Inches(3.3),
           Inches(0.6), 15, ORANGE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    d.hline(s, MARGIN, Inches(4.05), Inches(12.13))
    d.rich(s, [
        [("b", "Cách kiểm chứng: "), ("", "chạy lại GROUPING SETS từng khối trực tiếp từ fact, so từng ô "
                                        "với mart đã tính trước")],
        [("b", "Phân hoạch: "), ("", "mỗi dòng fact đóng góp đúng 1 lần vào mỗi grouping set — dbt test "
                                 "bắt lỗi trùng/thiếu")],
        [("b", "Trung bình là độ đo đại số: "), ("", "khối lưu SUM + COUNT rồi chia (slide 18)")],
    ], MARGIN, Inches(4.35), Inches(12.0), Inches(2.3), 17, gap=14)

    # ---- S16 · DEMO OLAP ----------------------------------------------------------------------------
    s = d.content("DEMO OLAP trực tiếp: /olap", "B · KHO DỮ LIỆU & OLAP",
                  "3'. CỬA CẮT: thiếu giờ thì bỏ bước pivot. Kịch bản bấm: mở /olap đang xem Năm × Mùa (đọc từ "
                  "cube) → bấm 2015 drill-down sang quý → roll-up về năm → lọc quận Dongcheng (slice) → thêm "
                  "Xicheng + 2016 (dice) → pivot đổi hàng/cột. Vừa bấm vừa gọi tên thao tác. DỰ PHÒNG: mất "
                  "mạng/Docker lỗi → mở notebooks/olap/olap_demo.ipynb (output đã chạy sẵn), nói như thường.",
                  title_size=32)
    d.picture(s, SHOTS / "olap_nam_x_mua.png", MARGIN, Inches(1.45), h=Inches(5.3))
    ops = [("1", "Roll-up", "quý → năm · tháng → quý (cuộn lên)"),
           ("2", "Drill-down", "bấm 2015 → khoan xuống quý"),
           ("3", "Slice", "một giá trị: quận = Dongcheng"),
           ("4", "Dice", "nhiều chiều: + Xicheng, + 2016"),
           ("5", "Pivot", "đổi hàng ↔ cột")]
    for i, (num, name, desc) in enumerate(ops):
        y = Inches(1.55) + i * Inches(1.0)
        d.badge(s, num, Inches(6.85), y)
        d.text(s, name, Inches(7.45), y - Inches(0.06), Inches(3.5), Inches(0.42), 18, INK, bold=True)
        d.text(s, desc, Inches(7.45), y + Inches(0.34), Inches(4.6), Inches(0.5), 14, MUTED)
    d.card(s, Inches(6.85), Inches(6.46), Inches(5.88), Inches(0.56), CARD)
    d.text(s, "Kịch bản: Năm×Mùa → 2015 → roll-up → slice → dice → pivot", Inches(7.0),
           Inches(6.48), Inches(5.6), Inches(0.5), 12, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    # ---- S17 · insight gió nam -------------------------------------------------------------------------
    s = d.content("Insight từ dice: gió nam mùa đông +60% PM2.5", "B · KHO DỮ LIỆU & OLAP",
                  "1'. Kết quả dice thật: cùng quận lõi đô thị, mùa đông, gió bắc ~85 vs gió nam ~133 µg/m³ "
                  "(+60%) — nguồn khí phía nam (Hà Bắc); tỷ lệ giờ vượt ngưỡng tăng 28% → 43%. Đây là giá trị "
                  "nghiệp vụ của chiều hướng gió. Nếu bị hỏi «sao biết khí từ Hà Bắc?»: nói thẳng đây là TƯƠNG "
                  "QUAN THEO HƯỚNG GIÓ, dự án không mô phỏng đường vận chuyển khí.")
    south = {(x["district"], x["season"]): x for x in f.op("dice_winter_south_core")["preview"]}
    north = {(x["district"], x["season"]): x for x in f.op("dice_compare_north")["preview"]}
    keys = sorted(south, key=lambda k: (k[0], k[1] != "autumn"))
    cd = CategoryChartData()
    cd.categories = [f"{k[0][:9]} · {'Thu' if k[1] == 'autumn' else 'Đông'}" for k in keys]
    cd.add_series("Gió bắc", [round(north[k]["pm25_avg"], 1) for k in keys])
    cd.add_series("Gió nam", [round(south[k]["pm25_avg"], 1) for k in keys])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.45), Inches(7.6),
                            Inches(5.3), cd)
    ch = gf.chart
    d.style_chart(ch, [TEAL, ORANGE])
    pl = d.labels_on(ch, "0", 11)
    pl.gap_width = 80
    ch.value_axis.has_title = True
    ch.value_axis.axis_title.text_frame.text = "PM2.5 trung bình (µg/m³)"
    ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    dc = ("Dongcheng", "winter")
    d.stat(s, f"{vn(north[dc]['pm25_avg'], 0)} → {vn(south[dc]['pm25_avg'], 0)}",
           "µg/m³ · Dongcheng mùa đông, gió bắc → gió nam (+60%)", Inches(8.5), Inches(1.55),
           Inches(4.23), value_size=32, h=Inches(1.7))
    d.stat(s, f"{vn(north[dc]['pct_exceed'], 0)}% → {vn(south[dc]['pct_exceed'], 0)}%",
           "tỷ lệ giờ vượt ngưỡng 115 cùng điều kiện", Inches(8.5), Inches(3.4), Inches(4.23),
           color=INK, h=Inches(1.6), value_size=32)
    d.text(s, "Nguồn khí phía nam (Hà Bắc) — khẳng định ở mức tương quan theo hướng gió, không mô phỏng "
              "đường vận chuyển khí.", Inches(8.5), Inches(5.2), Inches(4.23), Inches(1.4), 13, MUTED,
           italic=True, spacing=4)

    # ---- S18 · tổng hợp đúng ----------------------------------------------------------------------------
    s = d.content("Tổng hợp đúng: khối lưu SUM + COUNT, không lưu trung bình", "B · KHO DỮ LIỆU & OLAP",
                  "1'. Trung bình là độ đo KHÔNG CỘNG. Mart và cube luôn lưu SUM + COUNT rồi chia, nên cube gộp "
                  "đúng kể cả khi mức độ thiếu khác nhau giữa các nhóm. MỒI → LÁ CHẮN: «lấy tổng số giờ bằng SUM "
                  "cả bảng cube được không?», «trung bình của các trung bình?» → A9 (slide 41).", title_size=28)
    d.card(s, MARGIN, Inches(1.6), Inches(12.13), Inches(1.1), INK)
    d.text(s, "avg(PM2.5) của mọi cấp  =  SUM(PM2.5) / COUNT(giờ đo hợp lệ)     — không phải trung bình của "
              "trung bình", MARGIN + Inches(0.3), Inches(1.7), Inches(11.6), Inches(0.9), 19, WHITE,
           bold=True, anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, "Ví dụ 30 giờ: 24 giờ × 100 µg/m³ + 6 giờ × 200 µg/m³", MARGIN, Inches(3.05), Inches(12.0),
           Inches(0.5), 18, INK, bold=True, align=PP_ALIGN.CENTER)
    d.card(s, MARGIN + Inches(0.5), Inches(3.7), Inches(5.2), Inches(1.9), SOFT_TEAL)
    d.text(s, "3.600 / 30 = 120 ✓", MARGIN + Inches(0.8), Inches(3.85), Inches(4.6), Inches(0.7), 26,
           TEAL, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
    d.text(s, "tổng giá trị / tổng số giờ đo", MARGIN + Inches(0.8), Inches(4.7), Inches(4.6),
           Inches(0.6), 15, MUTED, align=PP_ALIGN.CENTER)
    d.card(s, Inches(6.93), Inches(3.7), Inches(5.2), Inches(1.9), SOFT_ORANGE)
    d.text(s, "(100 + 200) / 2 = 150 ✗", Inches(7.23), Inches(3.85), Inches(4.6), Inches(0.7), 26,
           ORANGE, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
    d.text(s, "trung bình của các trung bình — sai", Inches(7.23), Inches(4.7), Inches(4.6),
           Inches(0.6), 15, MUTED, align=PP_ALIGN.CENTER)
    d.text(s, "Cube gộp đúng kể cả khi mức độ thiếu khác nhau giữa các nhóm (COUNT phản ánh số giờ hợp lệ "
              "thật).", MARGIN, Inches(5.85), Inches(12.0), Inches(0.5), 15, MUTED, italic=True)

    # ---- S19 · chuyển phần C (dark) ------------------------------------------------------------------------
    s = d.dark('10". NÓI: "Kho trả lời được «đã xảy ra gì, vì sao». Còn câu hỏi đầu bài — «giờ tới có vượt '
               'ngưỡng không» — là việc của khai phá dữ liệu." → trao cho người C.')
    d.text(s, "PHẦN C · KHAI PHÁ & VẬN HÀNH", MARGIN, Inches(2.2), Inches(11), Inches(0.4), 15, AMBER,
           bold=True)
    d.text(s, ["Kho trả lời được: đã xảy ra gì, vì sao.", "Còn «giờ tới có vượt ngưỡng không»",
               "là việc của khai phá dữ liệu."], MARGIN, Inches(2.75), Inches(11.8), Inches(2.6), 30,
           WHITE, bold=True, font=HEAD, spacing=6)

    # ---- S20 · bài toán phân loại ---------------------------------------------------------------------------
    s = d.content("Bài toán phân loại: nhãn từ ngưỡng, 49 đặc trưng tại thời điểm t", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. Nhãn y = 1 nếu PM2.5(t+1h) ≥ 115 µg/m³ (ngưỡng đợt ô nhiễm khoá từ EDA). 49 đặc trưng, "
                  "mọi đặc trưng chỉ dùng thông tin BIẾT TẠI THỜI ĐIỂM t. MỒI → LÁ CHẮN: «115 là gì, có phải AQI "
                  "115 không?» → A1 (slide 33).", title_size=28)
    d.card(s, MARGIN, Inches(1.5), Inches(12.13), Inches(1.15), INK)
    d.text(s, "y = 1  ⟺  PM2.5 (t+1h) ≥ 115 µg/m³", MARGIN, Inches(1.6), Inches(12.13), Inches(0.95),
           30, WHITE, bold=True, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, "49 đặc trưng — tất cả chỉ dùng thông tin biết tại thời điểm t:", MARGIN, Inches(2.95),
           Inches(11.5), Inches(0.45), 17, INK, bold=True)
    groups = ["PM2.5 hiện tại", "Độ trễ 1–24 giờ", "Xu hướng (trung bình trượt)", "Chất khác: PM10, SO2, NO2, CO, O3",
              "Khí tượng: nhiệt độ, độ ẩm, gió, mưa", "Giờ · tháng · mùa", "Trạm", "Cụm trạm (phân cụm)"]
    for i, g in enumerate(groups):
        gx = MARGIN + (i % 2) * Inches(6.15)
        gy = Inches(3.5) + (i // 2) * Inches(0.78)
        d.card(s, gx, gy, Inches(5.95), Inches(0.62), CARD)
        d.text(s, g, gx + Inches(0.25), gy + Inches(0.06), Inches(5.5), Inches(0.5), 15, INK,
               anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, "Không dùng bất kỳ giá trị của tương lai — kể cả trong nội suy, scaling, chọn ngưỡng.",
           MARGIN, Inches(6.75), Inches(11.8), Inches(0.4), 14, MUTED, italic=True)

    # ---- S21 · chia tập ----------------------------------------------------------------------------------------
    s = d.content("Chia tập theo thời gian & chống rò rỉ", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. Với dữ liệu thời gian, cách chia tập quyết định kết quả có tin được không. Train "
                  "03/2013–12/2015: 276.380 dòng (25,2% nhãn 1); val 01–06/2016: 47.338 (15,8%); test 07–12/2016: "
                  "48.781 (23,9%). Chọn độ sâu cây và ngưỡng xác suất CHỈ trên val; test dùng đúng một lần để báo "
                  "cáo. Median điền thiếu fit trên train. MỒI → LÁ CHẮN: «vì sao không k-fold?», «rò rỉ ở đâu?» "
                  "→ A12 (slide 44).")
    segs = [("TRAIN", "03/2013 – 12/2015", f.cls["rows"]["train"], f.cls["class_balance"]["train"], TEAL, 7.2),
            ("VAL", "01 – 06/2016", f.cls["rows"]["val"], f.cls["class_balance"]["val"], MUTED, 2.4),
            ("TEST", "07 – 12/2016", f.cls["rows"]["test"], f.cls["class_balance"]["test"], ORANGE, 2.53)]
    x = MARGIN
    for name, period, n, bal, color, wdt in segs:
        shp = d.card(s, x, Inches(1.6), Inches(wdt) - Inches(0.08), Inches(1.3), color)
        d.text(s, f"{name} · {period}", x + Inches(0.18), Inches(1.7), Inches(wdt) - Inches(0.35),
               Inches(0.45), 15, WHITE, bold=True)
        d.text(s, f"{vn_int(n)} dòng · {vn(bal * 100, 1)}% nhãn 1", x + Inches(0.18), Inches(2.22),
               Inches(wdt) - Inches(0.35), Inches(0.45), 14, WHITE)
        x += Inches(wdt)
    d.text(s, "372.499 dòng dùng cho ML (cửa sổ 03/2013 – 12/2016; 01–02/2017 để đối soát CDC — xem A12/A14)",
           MARGIN, Inches(3.05), Inches(12.0), Inches(0.4), 14, MUTED)
    d.rich(s, [
        [("o", "Chia theo thời gian, không random: "), ("", "giờ láng giềng gần như trùng nhau — random "
                                                          "split cho kết quả lạc quan")],
        [("b", "Chọn siêu tham số chỉ trên val: "), ("", "độ sâu cây, ngưỡng xác suất")],
        [("b", "Test dùng đúng MỘT lần "), ("", "để báo cáo kết quả cuối")],
        [("b", "Median điền thiếu fit trên train, "), ("", "không nhìn val/test")],
        [("b", "Cửa sổ rolling không tràn qua biên split "), ("", "— rà và khẳng định bằng kiểm tra tự động")],
    ], MARGIN, Inches(3.7), Inches(12.0), Inches(2.9), 15, gap=9)

    # ---- S22 · baseline ----------------------------------------------------------------------------------------
    s = d.content("Baseline: bẫy accuracy và persistence rất mạnh", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. Majority: accuracy 0,761 nhưng F1 = 0 — không phát được cảnh báo nào (bẫy accuracy với lớp "
                  "thiểu số ~24%). Persistence (PM2.5 hiện tại ≥ 115): F1 0,914 — baseline mạnh vì hai giờ liên "
                  "tiếp tương quan cao. Mô hình phải chứng minh giá trị THÊM trên baseline này.")
    d.card(s, MARGIN, Inches(1.6), Inches(5.9), Inches(3.4), CARD)
    d.text(s, "Majority (luôn báo «không vượt»)", MARGIN + Inches(0.3), Inches(1.78), Inches(5.3),
           Inches(0.5), 18, INK, bold=True)
    d.text(s, f"accuracy {vn(f.test('majority')['accuracy'], 3)}", MARGIN + Inches(0.3), Inches(2.35),
           Inches(5.3), Inches(0.75), 34, MUTED, bold=True, font=HEAD)
    d.text(s, ["F1 = 0,000 — không phát được cảnh báo nào.", "Bẫy accuracy khi lớp thiểu số ~24%."],
           MARGIN + Inches(0.3), Inches(3.25), Inches(5.3), Inches(1.5), 16, INK, spacing=6)
    d.card(s, Inches(6.83), Inches(1.6), Inches(5.9), Inches(3.4), CARD)
    d.text(s, "Persistence — «PM2.5 hiện tại ≥ 115»", Inches(7.13), Inches(1.78),
           Inches(5.4), Inches(0.5), 18, INK, bold=True)
    d.text(s, f"F1 {vn(per['f1'], 3)}", Inches(7.13), Inches(2.35), Inches(5.3), Inches(0.75), 34,
           ORANGE, bold=True, font=HEAD)
    d.text(s, ["Baseline mạnh: hai giờ liên tiếp tương quan cao", "(ô nhiễm không đổi trạng thái đột ngột)."],
           Inches(7.13), Inches(3.25), Inches(5.3), Inches(1.5), 16, INK, spacing=6)
    d.card(s, MARGIN, Inches(5.3), Inches(12.13), Inches(1.15), SOFT_ORANGE)
    d.text(s, "Mô hình phải chứng minh giá trị THÊM trên baseline mạnh này — không chỉ F1 cao hơn.",
           MARGIN + Inches(0.3), Inches(5.42), Inches(11.5), Inches(0.9), 18, INK, bold=True,
           anchor=MSO_ANCHOR.MIDDLE)

    # ---- S23 · kết quả 4 mô hình ---------------------------------------------------------------------------------
    s = d.content("Kết quả 4 mô hình: HistGB khoá — F1 0,927", "C · KHAI PHÁ & VẬN HÀNH",
                  f"1'30\". Đọc bảng rồi ma trận: HistGradientBoosting F1 0,927 (precision 0,926, recall 0,928, "
                  f"ROC-AUC 0,993) — mô hình khoá, ngưỡng xác suất 0,50. Cây quyết định entropy F1 0,921. So "
                  f"persistence: giờ ô nhiễm bị bỏ sót (FN) {vn_int(per['fn'])} → {vn_int(loc['fn'])}, cảnh báo "
                  f"nhầm (FP) {vn_int(per['fp'])} → {vn_int(loc['fp'])}. MỒI → LÁ CHẮN: «vì sao chọn cây & "
                  "boosting?» → A12 cuối; «ngưỡng 115 với ngưỡng 0,50?» → A13.")
    rows = [("Majority", vn(f.test("majority")["f1"], 3), "–", vn(f.test("majority")["recall"], 3), "–"),
            ("Persistence", vn(per["f1"], 3), vn(per["precision"], 3), vn(per["recall"], 3), "–"),
            ("Cây quyết định (entropy, sâu 6)", vn(f.test("decision_tree")["f1"], 3),
             vn(f.test("decision_tree")["precision"], 3), vn(f.test("decision_tree")["recall"], 3),
             vn(f.test("decision_tree")["roc_auc"], 3)),
            ("HistGradientBoosting (cutoff 0,50)", vn(loc["f1"], 3), vn(loc["precision"], 3),
             vn(loc["recall"], 3), vn(loc["roc_auc"], 3))]
    d.dtable(s, ("Mô hình", "F1", "Precision", "Recall", "ROC-AUC"), rows, MARGIN, Inches(1.48),
             [4.3, 1.55, 1.9, 1.9, 1.9], row_h=Inches(0.46), size=14, highlight_row=3,
             header_h=Inches(0.44))
    for i, (v, lab, color) in enumerate((
            (f"FN {vn_int(per['fn'])} → {vn_int(loc['fn'])}", "giờ ô nhiễm bị bỏ sót", INK),
            (f"FP {vn_int(per['fp'])} → {vn_int(loc['fp'])}", "cảnh báo nhầm", INK),
            (f"F1 +{vn(loc['f1'] - per['f1'], 3)}", "so với cách hiện tại", ORANGE))):
        d.stat(s, v, lab, Inches(1.9) + i * Inches(3.6), Inches(4.2), Inches(3.4), color=color,
               h=Inches(0.95), value_size=22)
    d.picture(s, ART / "classification" / "figures" / "confusion_matrices_test.png", Inches(2.92),
              Inches(5.24), w=Inches(7.5))

    # ---- S24 · giá trị thật ----------------------------------------------------------------------------------------
    s = d.content("Giá trị thật: báo trước được lúc đợt ô nhiễm bắt đầu", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. F1 +0,013 nghe nhỏ — giá trị thật nằm ở thứ persistence KHÔNG làm được. Trong 1.000 giờ "
                  "đang sạch nhưng giờ sau vượt ngưỡng (lúc đợt ô nhiễm bắt đầu — thời điểm cần cảnh báo nhất), mô "
                  "hình báo trước 304 lần = 30,4%; persistence 0% theo định nghĩa. Kiểm định McNemar trên từng giờ "
                  "test: b = 594 giờ chỉ mô hình đúng, c = 298 giờ chỉ persistence đúng, p ≈ 5×10⁻²³ — khác biệt "
                  "có ý nghĩa thống kê. MỒI → LÁ CHẮN: «McNemar là gì, đủ để triển khai?» → A13 (slide 45).",
                  title_size=30)
    d.stat(s, f"{vn_int(f.onset['locked']['detected'])} / {vn_int(f.onset['locked']['onsets'])}",
           "giờ «đang sạch, giờ sau vượt ngưỡng» được báo trước = 30,4% — persistence: 0% theo định nghĩa",
           MARGIN, Inches(1.6), Inches(6.0), h=Inches(2.1), value_size=48)
    d.stat(s, sci(f.mcnemar["p_value"]),
           f"McNemar: b = {vn_int(f.mcnemar['b'])} giờ chỉ mô hình đúng · c = {vn_int(f.mcnemar['c'])} giờ "
           "chỉ persistence đúng", Inches(6.85), Inches(1.6), Inches(5.88), color=INK, h=Inches(2.1),
           value_size=40)
    d.card(s, MARGIN, Inches(4.0), Inches(12.13), Inches(2.3), SOFT_ORANGE)
    d.text(s, ["Persistence không bao giờ báo trước được lúc đợt bắt đầu — nó chỉ nhìn số đo hiện tại.",
               "304/1.000 giờ khởi đầu được báo trước: đúng lúc người vận hành cần cảnh báo nhất."],
           MARGIN + Inches(0.35), Inches(4.2), Inches(11.4), Inches(1.9), 18, INK, spacing=10,
           anchor=MSO_ANCHOR.MIDDLE)

    # ---- S25 · tầm dự báo --------------------------------------------------------------------------------------------
    s = d.content("Tầm dự báo: càng sớm, mô hình càng có lợi thế", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. F1 test theo tầm: 1h 0,914 → 0,927 (+0,013); 3h 0,806 → 0,832 (+0,026); 6h 0,702 → 0,745 "
                  "(+0,043). 12/12 trạm đều F1 cao hơn persistence. Đổi lại ở tầm xa: precision giảm nhẹ, recall "
                  "tăng mạnh (0,70 → 0,80 ở 6h).")
    cd = CategoryChartData()
    cd.categories = ["1 giờ", "3 giờ", "6 giờ"]
    cd.add_series("Persistence", [round(f.horizon_f1(h, "persistence"), 3) for h in (1, 3, 6)])
    cd.add_series("HistGradientBoosting", [round(f.horizon_f1(h, "hist_gradient_boosting"), 3)
                                           for h in (1, 3, 6)])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.5), Inches(8.0),
                            Inches(5.2), cd)
    ch = gf.chart
    d.style_chart(ch, [MUTED, ORANGE])
    ch.has_title = True
    ch.chart_title.text_frame.text = "F1 trên test theo tầm cảnh báo"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    ch.value_axis.minimum_scale = 0.6
    ch.value_axis.maximum_scale = 1.0
    pl = d.labels_on(ch)
    pl.gap_width = 70
    deltas = [f.horizon_f1(h, "hist_gradient_boosting") - f.horizon_f1(h, "persistence") for h in (1, 3, 6)]
    for i, hname in enumerate(("tầm 1h", "tầm 3h", "tầm 6h")):
        d.stat(s, f"+{vn(deltas[i], 3)}", f"F1 so persistence · {hname}", Inches(8.85),
               Inches(1.6) + i * Inches(1.32), Inches(3.88), h=Inches(1.2), value_size=28)
    d.stat(s, "12 / 12", "trạm có F1 cao hơn persistence trên test", Inches(8.85), Inches(5.55),
           Inches(3.88), color=TEAL, fill=SOFT_TEAL, h=Inches(1.2), value_size=28)

    # ---- S26 · DEMO canh-bao + phat-lai --------------------------------------------------------------------------------
    s = d.content("DEMO: /canh-bao và /phat-lai — dự đoán trước, đối soát sau", "C · KHAI PHÁ & VẬN HÀNH",
                  "2'. CỬA CẮT: thiếu giờ thì rút còn /canh-bao. /canh-bao: hai ma trận nhầm lẫn cạnh nhau (mô "
                  "hình vs cách hiện tại), dòng thời gian xác suất tại một trạm — xanh = cảnh báo đúng, vàng = "
                  "báo nhầm, đỏ = bỏ sót. /phat-lai: kéo mốc thời gian, hệ thống phát dự đoán cho giờ đích, số đo "
                  "thật chỉ hiện khi mốc dữ liệu đã đi qua — «dự đoán trước, đối soát sau». Nói rõ: F1 cao trong "
                  "cửa sổ demo KHÔNG phải bằng chứng cho độ tốt của mô hình (cửa sổ chọn nhiều giờ ô nhiễm). "
                  "DỰ PHÒNG: ảnh chụp màn hình đã lưu trong slide.", title_size=28)
    d.picture(s, SHOTS / "canh_bao.png", MARGIN, Inches(1.45), h=Inches(5.3))
    d.text(s, "/canh-bao", Inches(6.95), Inches(1.5), Inches(5.7), Inches(0.45), 19, TEAL, bold=True)
    d.rich(s, [
        [("b", "Hai ma trận nhầm lẫn "), ("", "cạnh nhau: mô hình vs cách hiện tại")],
        [("b", "Dòng thời gian xác suất "), ("", "tại một trạm")],
        [("", "chấm xanh = cảnh báo đúng · vàng = báo nhầm · đỏ = bỏ sót")],
    ], Inches(6.95), Inches(2.05), Inches(5.75), Inches(1.9), 15, gap=8)
    d.text(s, "/phat-lai", Inches(6.95), Inches(4.1), Inches(5.7), Inches(0.45), 19, ORANGE, bold=True)
    d.rich(s, [
        [("b", "Kéo mốc thời gian "), ("", "— hệ thống phát dự đoán cho giờ đích")],
        [("b", "Số đo thật chỉ hiện khi mốc đã đi qua "), ("", "— không nhìn trộm tương lai")],
    ], Inches(6.95), Inches(4.65), Inches(5.75), Inches(1.5), 15, gap=8)
    d.text(s, "Lưu ý trung thực: F1 cao trong cửa sổ demo không phải bằng chứng cho độ tốt của mô hình.",
           Inches(6.95), Inches(6.3), Inches(5.75), Inches(0.6), 13, MUTED, italic=True)

    # ---- S27 · CDC --------------------------------------------------------------------------------------------------------
    s = d.content("Luồng CDC: hệ thống «sống» sau khi báo cáo", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. SQLite nguồn tác nghiệp → đọc tăng dần theo watermark (chỉ lấy bản ghi sau mốc) → Airflow "
                  "DAG → MinIO → upsert staging. Idempotency bằng load_registry (đối chiếu object key + sha256 + "
                  "row_count). Mỗi batch nạp xong được chấm MAE đối soát với số đo thật, vượt 12 µg/m³ thì bật cờ "
                  "giám sát. Nói thẳng giới hạn: watermark bắt được insert mới; bản ghi cũ bị sửa/xoá sau khi mốc "
                  "đã đi qua thì không bắt được — cần đối soát lại theo partition. MỒI → LÁ CHẮN: «bản ghi cũ bị "
                  "sửa thì sao?», «chạy lại có nhân đôi không?», «vượt ngưỡng MAE rồi làm gì?» → A14 (slide 46).",
                  title_size=30)
    steps = [("SQLite", "nguồn tác nghiệp\n(mô phỏng giao dịch)", MUTED),
             ("Watermark", "đọc tăng dần:\nchỉ bản ghi sau mốc", INK),
             ("Airflow DAG", "điều phối\ntheo lịch", INK),
             ("MinIO", "lưu batch\nthô + đối soát", TEAL),
             ("Staging", "upsert theo\nkhóa grain", TEAL),
             ("Chấm MAE", f"đối soát số đo thật\nngưỡng {vn(12, 0)} µg/m³", ORANGE)]
    x = MARGIN
    for i, (name, desc, color) in enumerate(steps):
        wch = Inches(1.78)
        d.card(s, x, Inches(1.55), wch, Inches(1.55), CARD)
        d.chip(s, name, x + Inches(0.08), Inches(1.67), wch - Inches(0.16), fill=color, size=13)
        d.text(s, desc, x + Inches(0.08), Inches(2.22), wch - Inches(0.16), Inches(0.85), 11, INK,
               align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            d.arrow_right(s, x + wch + Inches(0.02), Inches(2.22), w=Inches(0.24))
        x += wch + Inches(0.28)
    d.stat(s, f"{vn(load_json('artifacts/cdc_eval/mae_by_batch.json')['overall']['mae'], 2)} µg/m³",
           "MAE đối soát liên tục · "
           f"{vn_int(load_json('artifacts/cdc_eval/mae_by_batch.json')['overall']['scored_n'])} giờ · 01/2016–02/2017",
           MARGIN, Inches(3.45), Inches(3.93), h=Inches(1.7), value_size=34)
    d.stat(s, "12 µg/m³", "ngưỡng cờ giám sát từng batch — vượt ngưỡng chỉ BẬT CỜ (monitor_only), mô hình "
           "khoá, không tự retrain", Inches(4.73), Inches(3.45), Inches(3.93), color=INK, h=Inches(1.7),
           value_size=34)
    d.stat(s, "load_registry", "object key + sha256 + schema_version + row_count khớp → skip: chạy lại "
           "không nhân đôi", Inches(8.86), Inches(3.45), Inches(3.87), color=TEAL, fill=SOFT_TEAL,
           h=Inches(1.7), value_size=26)
    d.card(s, MARGIN, Inches(5.45), Inches(12.13), Inches(1.15), SOFT_ORANGE)
    d.text(s, "Giới hạn nói thật: watermark bắt INSERT mới. Bản ghi cũ bị SỬA/XOÁ sau khi mốc đã đi qua "
              "thì không bắt được → cần đối soát lại theo partition.", MARGIN + Inches(0.3), Inches(5.57),
           Inches(11.5), Inches(0.9), 16, INK, anchor=MSO_ANCHOR.MIDDLE)

    # ---- S28 · thí nghiệm 6h (cửa cắt) -------------------------------------------------------------------------------------
    s = d.content("Thí nghiệm cảnh báo 6 giờ — riêng, chưa thay mô hình t+1", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. CỬA CẮT (mặc định bỏ khi 30'). Câu hỏi mở rộng: đang sạch, có đợt (≥3 giờ liên tục vượt "
                  "ngưỡng) bắt đầu trong t+1→t+6 không. Mô hình có đặc trưng trung vị các trạm lân cận: bắt "
                  "506/730 đợt = 69,3%; luật đơn giản «hiện tại ≥ 105»: 468/730 = 64,1%; trung vị thời gian báo "
                  "trước 3 giờ (so với 1 giờ); trần cảnh báo nhầm 0,5/ngày-trạm chọn trên validation. Ghi rõ: thí "
                  "nghiệm riêng, chưa thay mô hình t+1. MỒI → LÁ CHẮN: «định nghĩa đợt?», «vì sao kiểm tra tới "
                  "t+8?» → A15 (slide 47).", title_size=28)
    d.card(s, MARGIN, Inches(1.5), Inches(12.13), Inches(1.05), INK)
    d.text(s, "Đang sạch (PM2.5 < 115)  →  có đợt (≥ 3 giờ liên tục ≥ 115) bắt đầu trong t+1 … t+6 ?",
           MARGIN + Inches(0.2), Inches(1.58), Inches(11.7), Inches(0.9), 19, WHITE, bold=True,
           anchor=MSO_ANCHOR.MIDDLE)
    onset6 = pd.read_csv(ART / "classification" / "onset_6h" / "onset_metrics.csv")
    sel = onset6[(onset6.fold == "2016h1") & (onset6.split == "test") & (onset6.budget == "primary")]
    row_m = sel[sel.variant == "histgb_spatial"].iloc[0]
    row_l = sel[sel.variant == "level"].iloc[0]
    n_ep = int(row_m.n_episodes)
    d.stat(s, f"{vn_int(row_m.n_detected)} / {vn_int(n_ep)}",
           f"đợt được báo trước = {vn(row_m.episode_recall * 100, 1)}% — mô hình có đặc trưng trung vị trạm "
           "lân cận", MARGIN, Inches(2.85), Inches(4.0), h=Inches(1.9), value_size=36)
    d.stat(s, f"{vn_int(row_l.n_detected)} / {vn_int(n_ep)}",
           f"= {vn(row_l.episode_recall * 100, 1)}% — luật đơn giản «hiện tại ≥ 105»", Inches(4.73),
           Inches(2.85), Inches(3.93), color=INK, h=Inches(1.9), value_size=36)
    d.stat(s, f"{vn(row_m.median_lead_h, 0)} giờ", "trung vị thời gian báo trước (luật 105: 1 giờ) · trần "
           "cảnh báo nhầm 0,5/ngày-trạm chọn trên validation", Inches(8.86), Inches(2.85), Inches(3.87),
           color=TEAL, fill=SOFT_TEAL, h=Inches(1.9), value_size=36)
    d.rich(s, [
        [("b", "Đợt bắt đầu muộn nhất t+6 "), ("", "cần xem tới t+8 mới đủ 3 giờ liên tục")],
        [("b", "Nhãn chỉ giữ "), ("", "khi toàn bộ cửa sổ nằm trong một split (không tràn qua biên)")],
        [("o", "Thí nghiệm riêng — "), ("", "chưa thay mô hình t+1 đang khoá")],
    ], MARGIN, Inches(5.05), Inches(12.0), Inches(1.6), 15, gap=9)

    # ---- S29 · bổ trợ (cửa cắt) ---------------------------------------------------------------------------------------------
    s = d.content("Phân tích bổ trợ: hồi quy dự báo số & phân cụm trạm", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. CỬA CẮT. Hồi quy PM2.5(t+1): MAE 8,87 µg/m³ — sai số tuyệt đối trung bình, KHÔNG phải "
                  "trần sai số từng dự báo. Phân cụm: mỗi trạm thành một vector hồ sơ → K-Means, thử K=2,3,4 → "
                  "K=2, silhouette 0,664 cao nhất, cụm 9 trạm + 3 trạm; địa lý chỉ dùng đối chiếu, không đưa vào "
                  "đặc trưng. MỒI → LÁ CHẮN: «MAE nghĩa gì? Phân cụm cái gì?» → A16 (slide 48).", title_size=28)
    d.card(s, MARGIN, Inches(1.55), Inches(5.9), Inches(4.9), CARD)
    d.text(s, "Hồi quy — dự báo số PM2.5(t+1)", MARGIN + Inches(0.3), Inches(1.75), Inches(5.3),
           Inches(0.5), 18, TEAL, bold=True)
    d.text(s, f"MAE {vn(f.reg_mae('hist_gradient_boosting'), 2)} µg/m³", MARGIN + Inches(0.3),
           Inches(2.35), Inches(5.3), Inches(0.8), 38, ORANGE, bold=True, font=HEAD)
    d.rich(s, [
        [("b", "MAE = "), ("", "sai số tuyệt đối TRUNG BÌNH (cùng đơn vị µg/m³)")],
        [("", "không phải trần sai số — từng dự báo có thể sai nhiều hơn")],
        [("b", "Dùng cho: "), ("", "trang /du-bao, kèm phân vùng sai số")],
    ], MARGIN + Inches(0.25), Inches(3.35), Inches(5.4), Inches(2.9), 15, gap=10)
    d.card(s, Inches(6.83), Inches(1.55), Inches(5.9), Inches(4.9), CARD)
    d.text(s, "Phân cụm — 12 trạm thành hồ sơ", Inches(7.13), Inches(1.75), Inches(5.3), Inches(0.5),
           18, TEAL, bold=True)
    d.text(s, "K = 2 · silhouette 0,664", Inches(7.13), Inches(2.35), Inches(5.3), Inches(0.8), 27,
           ORANGE, bold=True, font=HEAD)
    d.rich(s, [
        [("b", "Đối tượng: "), ("", "12 trạm (mỗi trạm 1 vector hồ sơ), không phải 420k giờ")],
        [("b", "Chuẩn hoá trước K-Means; "), ("", "thử K = 2, 3, 4 → K=2 tách cụm rõ nhất (9 + 3)")],
        [("b", "Địa lý: "), ("", "chỉ dùng đối chiếu, không đưa vào đặc trưng")],
    ], Inches(7.08), Inches(3.35), Inches(5.4), Inches(2.9), 15, gap=10)

    # ---- S30 · khuyến nghị & hạn chế -----------------------------------------------------------------------------------------
    s = d.content("Khuyến nghị & hạn chế", "C · KHAI PHÁ & VẬN HÀNH",
                  "1'. Ba khuyến nghị gắn trực tiếp kết quả. Hạn chế nói thật: test chỉ 6 tháng cuối 2016, dữ liệu "
                  "lịch sử; +0,013 F1 ở 1 giờ khiêm tốn; ngưỡng 115 vốn là ngưỡng 24 giờ áp cho từng giờ.")
    d.card(s, MARGIN, Inches(1.55), Inches(5.9), Inches(4.0), SOFT_TEAL)
    d.text(s, "Nên làm", MARGIN + Inches(0.3), Inches(1.72), Inches(5.3), Inches(0.5), 20, TEAL, bold=True)
    d.rich(s, [
        [("b", "Cảnh báo theo xác suất mô hình, "), ("", "không chờ số đo vượt")],
        [("b", "Ưu tiên mùa đông, buổi tối, "), ("", "khi có gió nam")],
        [("b", "Cần sớm 3–6 giờ → "), ("", "chấp nhận precision thấp hơn để recall cao")],
    ], MARGIN + Inches(0.25), Inches(2.35), Inches(5.4), Inches(3.0), 17, gap=14, lead_color=TEAL)
    d.card(s, Inches(6.83), Inches(1.55), Inches(5.9), Inches(4.0), CARD)
    d.text(s, "Hạn chế — nói thật", Inches(7.13), Inches(1.72), Inches(5.3), Inches(0.5), 20, ORANGE,
           bold=True)
    d.rich(s, [
        [("b", "Test chỉ 6 tháng cuối 2016, "), ("", "dữ liệu lịch sử")],
        [("b", "+0,013 F1 ở 1 giờ "), ("", "là khiêm tốn (có ý nghĩa thống kê)")],
        [("b", "Ngưỡng 115 "), ("", "vốn là ngưỡng 24 giờ, áp cho từng giờ")],
    ], Inches(7.08), Inches(2.35), Inches(5.4), Inches(3.0), 17, gap=14)
    d.text(s, "Mọi khuyến nghị đều truy được về một số đo cụ thể đã trình bày ở phần C.", MARGIN,
           Inches(5.85), Inches(12.0), Inches(0.5), 15, MUTED, italic=True)

    # ---- S31 · kết (dark) ------------------------------------------------------------------------------------------------------
    s = d.dark('10". Cảm ơn thầy/cô và các bạn đã lắng nghe, mời hỏi đáp. (Phụ lục A0–A17 chỉ để nhóm dùng '
               'khi bị hỏi — không giới thiệu trên slide. Trong chế độ trình chiếu: gõ số slide + Enter để '
               'nhảy tới phụ lục.)')
    d.text(s, "CẢM ƠN THẦY/CÔ VÀ CÁC BẠN", MARGIN, Inches(1.9), Inches(11.8), Inches(0.5), 15, AMBER,
           bold=True)
    d.text(s, "Hỏi đáp", MARGIN, Inches(2.45), Inches(11.5), Inches(1.2), 48, WHITE, bold=True, font=HEAD)
    d.hline(s, MARGIN, Inches(3.9), Inches(5.2), RGBColor(0x3A, 0x48, 0x5C))
    d.text(s, "Rất mong nhận câu hỏi từ thầy/cô và các bạn.", MARGIN, Inches(4.15), Inches(11.5),
           Inches(0.5), 18, SLATE_L)
    d.text(s, "NHÓM THỰC HIỆN", MARGIN, Inches(5.28), Inches(5), Inches(0.3), 11, AMBER, bold=True)
    for i, (name, mssv) in enumerate(members):
        mx = MARGIN + i * Inches(3.6)
        d.text(s, name, mx, Inches(5.64), Inches(3.5), Inches(0.35), 16, WHITE, bold=True)
        d.text(s, mssv, mx, Inches(6.02), Inches(3.5), Inches(0.3), 12, SLATE_L)

    # ---- S32 = A0 · mục lục phụ lục (dark) --------------------------------------------------------------------------------------
    s = d.dark('A0. Mở ngay khi bắt đầu hỏi đáp để định vị lá chắn. Gõ số slide + Enter để nhảy. Cụm: dữ liệu '
               'A1–A3 · kiến trúc A4 · kho A5–A9 · OLAP A10 · modelling A11 · ML A12–A13 · vận hành A14–A15 · '
               'bổ trợ A16 · demo A17.')
    d.text(s, "PHỤ LỤC — 17 LÁ CHẮN CHO HỎI ĐÁP", MARGIN, Inches(0.65), Inches(11), Inches(0.4), 14,
           AMBER, bold=True)
    d.text(s, "Chỉ mở khi bị hỏi · gõ số slide + Enter để nhảy", MARGIN, Inches(1.1), Inches(11),
           Inches(0.55), 28, WHITE, bold=True, font=HEAD)
    shields = [
        ("A1", "Ngưỡng 115 từ đâu?", 33), ("A2", "Ghép hai nguồn thế nào?", 34),
        ("A3", "Vì sao không nội suy?", 35), ("A4", "Lake/staging/mart · đều là DB à?", 36),
        ("A5", "Vì sao grain giờ?", 37), ("A6", "Sao không có chiều thời tiết?", 38),
        ("A7", "4 chiều sao ra 32 khối?", 39), ("A8", "Kiểm chứng kho thế nào?", 40),
        ("A9", "SUM cube? Trung bình của trung bình?", 41), ("A10", "Kiến trúc OLAP gì? (ROLAP)", 42),
        ("A11", "Khóa & SCD?", 43), ("A12", "K-fold, rò rỉ, 372.499?", 44),
        ("A13", "McNemar & hai ngưỡng?", 45), ("A14", "CDC chi tiết?", 46),
        ("A15", "Đợt & thí nghiệm 6h?", 47), ("A16", "Hồi quy & phân cụm?", 48),
        ("A17", "Demo 3' & tái lập", 49),
    ]
    col_x = [MARGIN, Inches(6.85)]
    for i, (code, topic, num) in enumerate(shields):
        cx = col_x[i // 9]
        cy = Inches(2.0) + (i % 9) * Inches(0.55)
        d.text(s, code, cx, cy, Inches(0.55), Inches(0.4), 15, AMBER, bold=True, font=HEAD)
        d.text(s, topic, cx + Inches(0.6), cy, Inches(4.7), Inches(0.4), 15, WHITE)
        d.text(s, str(num), cx + Inches(5.35), cy, Inches(0.5), Inches(0.4), 15, SLATE_L, bold=True,
               align=PP_ALIGN.RIGHT)

    # ======================================================================= appendix shields
    def blocks(s, items, y0=Inches(1.85), size=15, gap=12, w=Inches(12.13), x=MARGIN):
        """items: list (lead, body) — lead bold màu, body thường, ngăn bằng hairline."""
        y = y0
        for lead, body in items:
            d.hline(s, x, y, w)
            y += Inches(0.12)
            d.text(s, lead, x, y, Inches(2.35), Inches(0.4), size + 1, ORANGE, bold=True)
            d.text(s, body, x + Inches(2.5), y - Inches(0.04), w - Inches(2.5), Inches(1.15), size, INK,
                   spacing=4)
            y += Inches(0.52) + Inches(0.17 * max(0, math.ceil(len(body) / 90) - 1))
        return y

    # A1 (33)
    s = d.shield("A1", "Ngưỡng 115 µg/m³ từ đâu?", "«Vì sao 115? Có phải AQI 115 không?»",
                 "Mở khi bị hỏi về ngưỡng. Trả lời khung: 115 = ranh giới dải «ô nhiễm trung bình» theo quy "
                 "chuẩn AQI Trung Quốc cho PM2.5 — khác với việc đọc «AQI = 115».")
    blocks(s, [
        ("Dự án dùng", "115 µg/m³ làm ngưỡng vượt của PM2.5 — dùng thống nhất cho: nhãn ML, thống kê vượt "
                       "ngưỡng trong cube, khuyến nghị cảnh báo."),
        ("Nghĩa là gì", "Ranh giới dải «ô nhiễm trung bình» theo quy chuẩn AQI Trung Quốc (HJ 633) cho "
                        "PM2.5 — KHÁC với việc đọc «AQI = 115»: đây là nồng độ PM2.5, không phải điểm AQI."),
        ("Khoá từ đâu", "Phân tích EDA đợt ô nhiễm (Phase 8c): đợt tại trạm = ≥ 3 giờ liên tiếp ≥ 115. "
                        "EDA khoá cả quy tắc đa trạm: ≥ 6 trạm cùng lúc."),
        ("Nói thật", "Đây là ngưỡng trung bình 24 giờ được áp cho từng giờ — một giả định khai thác, nêu "
                     "trong hạn chế (slide 30)."),
    ])

    # A2 (34)
    s = d.shield("A2", "Ghép hai nguồn: UCI + Open-Meteo", "«Open-Meteo ghép với UCI thế nào?»",
                 "Mở khi hỏi về ghép nguồn. Khung: cùng 12 toạ độ, cùng mốc giờ; UCI ưu tiên, OM chỉ lấp "
                 "khi UCI null; ghi nguồn từng ô.")
    blocks(s, [
        ("Cách ghép", "Cùng 12 toạ độ trạm, cùng mốc giờ → join theo (trạm, giờ). Quy tắc resolve: UCI "
                      "ƯU TIÊN, Open-Meteo chỉ lấp khi UCI bị NULL — mỗi ô sau lấp đều GHI NGUỒN."),
        ("OM thêm gì", "Độ che phủ mây (chỉ có ở OM); RH không lấy của OM mà suy từ TEMP + DEWP của UCI "
                       "để giữ nhất quán thiết bị đo."),
        ("Đối soát", "Script fill-rate chấm tỷ lệ lấp thật của từng cột — con số in trên slide 5 lấy từ "
                     "đây (thời tiết còn 0 ô thiếu sau lấp)."),
    ])

    # A3 (35)
    s = d.shield("A3", "Vì sao không nội suy hết cho đẹp?", "«Điền hết đi, bảng sạch hơn?»",
                 "Mở khi hỏi về giữ NULL. Khung: chuỗi thiếu dài 343–1.517 giờ — nội suy chuỗi dài = bịa "
                 'dữ liệu, tạo "đợt giả".')
    blocks(s, [
        ("Sự thật dữ liệu", "Chuỗi thiếu CO dài 343 – 1.517 giờ liên tục (trạm worst-case). Nội suy qua "
                            "chuỗi dài ấy là BỊA DỮ LIỆU — có thể tạo ra đợt ô nhiễm giả trong phân tích."),
        ("Dự án làm gì", "Giữ NULL + cờ chất lượng; COUNT(cột) ≠ COUNT(*) là có chủ đích; NULL ≠ 0 "
                         "(không tự biến thiếu thành «sạch»)."),
        ("ML xử lý thế nào", "Median điền thiếu — tính CHỈ trên tập train, áp cho val/test (không nhìn "
                             "tương lai, không rò rỉ phân phối)."),
    ])

    # A4 (36)
    s = d.shield("A4", "Vai trò từng lớp & công cụ — không phải tất cả là database",
                 "«Lake/staging/mart khác nhau thế nào? 4 công cụ kia đều là database à?»",
                 "Mở khi hỏi vai trò lớp/không phải tất cả là database, cộng nửa dòng bảo mật: .env gitignored, "
                 "mạng Docker nội bộ, không có dữ liệu cá nhân.")
    blocks(s, [
        ("Bốn lớp", "Lake (MinIO) = object storage, giữ raw/processed nguyên vẹn · Staging (PostgreSQL) = "
                    "vùng tiếp nhận có cấu trúc, đối soát nguồn · Kho/DWH (PostgreSQL + dbt) = fact/dimension "
                    "phục vụ phân tích · Mart = tập con phục vụ một nhu cầu xác định."),
        ("Công cụ", "MinIO = object storage, PostgreSQL = RDBMS, dbt = khung biến đổi SQL (ELT), Airflow = "
                    "điều phối — KHÔNG phải tất cả là database."),
        ("Bảo mật", ".env bị gitignored; hạ tầng chạy trong mạng Docker nội bộ (air_quality_network); nguồn "
                    "dữ liệu công khai — không có dữ liệu cá nhân."),
    ])

    # A5 (37)
    s = d.shield("A5", "Vì sao grain trạm × giờ?", "«Lưu mức ngày cho nhẹ được không?»",
                 'Mở khi hỏi grain. Khung: giờ→ngày cuộn được, ngược lại mất vĩnh viễn; đổi sang ngày mất '
                 'biến động trong ngày.')
    blocks(s, [
        ("Một chiều", "Giờ → ngày → tháng cuộn (roll-up) được bất cứ lúc nào; ngày → giờ KHÔNG thể — mất "
                      "thông tin vĩnh viễn ngay từ lúc nạp."),
        ("Bài toán", "Cảnh báo phát cho giờ tới (t+1h) — cần đúng mức giờ; độ đo và khóa chiều khớp grain."),
        ("Nếu lưu ngày", "Mất biến động trong ngày (giờ cao điểm buổi tối — 18–24h chính là thứ mô hình "
                         "dùng làm đặc trưng xu hướng)."),
    ])

    # A6 (38)
    s = d.shield("A6", "Vì sao thời tiết không thành chiều?", "«Thiếu dim_weather chứ?»",
                 'Mở khi hỏi thiếu chiều thời tiết. Khung: thời tiết là điều kiện đo CÙNG sự kiện tại grain '
                 'trạm×giờ → độ đo trên fact = một nguồn sự thật.')
    blocks(s, [
        ("Nguyên tắc", "Thời tiết được đo CÙNG sự kiện, tại cùng grain trạm × giờ → là ĐỘ ĐO trên fact "
                       "(cột: TEMP, PRES, DEWP, RAIN, WSPM…), không phải thực thể riêng."),
        ("Lợi ích", "Một nguồn sự thật: không đồng bộ hai nơi; JOIN ít hơn; phân tích tương quan chất — "
                    "thời tiết chỉ cần GROUP BY theo chiều có sẵn."),
        ("Nếu cần riêng", "Muốn phân tích HỒ SƠ thời tiết theo thời gian → dựng fact riêng (grain khác), "
                          "không nhét thêm chiều vào sao này."),
    ])

    # A7 (39)
    s = d.shield("A7", "4 chiều mà 32 khối — hai con số khác nhau", "«4 chiều sao ra 32 khối?»",
                 "Mở khi hỏi 32 vs 180. Khung: CUBE theo 5 THUỘC TÍNH (năm, mùa, quận, buổi, phần tư gió) — "
                 "mỗi thuộc tính giữ-lại-hoặc-cuộn-lên-ALL → 2⁵ = 32 grouping sets.")
    rows = [("dim_date: năm · mùa", "năm / tháng / quý / ngày / ALL — 2 thuộc tính CUBE (năm, mùa) cùng nằm ở đây", "5"),
            ("dim_time: giờ · buổi", "giờ / buổi / ALL — 1 thuộc tính CUBE (buổi)", "3"),
            ("dim_station: trạm · quận", "trạm / quận / ALL — 1 thuộc tính CUBE (quận)", "3"),
            ("dim_wind: hướng · chế độ", "16 hướng / 4 phần tư / bắc-nam / ALL — 1 thuộc tính CUBE (phần tư)", "4")]
    d.dtable(s, ("Chiều", "Các mức (tính cả ALL)", "Số mức"), rows, MARGIN, Inches(1.85), [3.1, 6.9, 1.4],
             row_h=Inches(0.66), size=13, header_h=Inches(0.46))
    d.text(s, ["180 = tổ hợp MỨC PHÂN CẤP của 4 chiều (5·3·3·4) — nói về độ sâu phân tích.",
               "32 = GROUPING SETS của CUBE 5 thuộc tính (2⁵) — nói về khối đã tính trước.",
               "Hai con số đo hai thứ khác nhau — đừng trộn."],
           MARGIN, Inches(5.55), Inches(12.0), Inches(1.4), 15, INK, spacing=6)

    # A8 (40)
    s = d.shield("A8", "Kiểm chứng kho: tính lại từng ô từ fact", "«Làm sao biết khối tính trước đúng?»",
                 "Mở khi hỏi kiểm chứng. Khung: script chạy lại GROUPING SETS từng khối TRỰC TIẾP từ fact, "
                 "so từng ô — 24.759/24.759 = 100%.")
    blocks(s, [
        ("Kiểm chứng ô", "Script chạy lại từng khối (GROUPING SETS) trực tiếp từ fact, so khớp TỪNG Ô: "
                         "24.759 / 24.759 ô đúng = 100% — chạy bằng lệnh, không phải kiểm tra thủ công."),
        ("dbt test", f"{f.dbt_pass}/{f.dbt_total} test: khóa duy nhất · not-null · khóa ngoại · cube PHÂN "
                     "HOẠCH fact (mỗi dòng fact đóng góp đúng 1 lần vào mỗi grouping set — không trùng, "
                     "không thiếu)."),
        ("Giới hạn thật", "Không phải mọi truy vấn dùng được khối: thiếu 1 thuộc tính là phải quét fact "
                          "(130–400 ms) — vì vậy trang OLAP được thiết kế theo thuộc tính của cube."),
    ])

    # A9 (41)
    s = d.shield("A9", "SUM cube & trung bình của trung bình", "«Lấy tổng số giờ bằng SUM cả bảng cube?»",
                 'Mở khi hỏi tổng hợp trên cube. Khung: mỗi dòng fact nằm trong NHIỀU grouping set — SUM '
                 'cả bảng đếm trùng.')
    blocks(s, [
        ("Vì sao sai", "Mỗi dòng fact nằm trong NHIỀU grouping set (từng tổ hợp cuộn-up chứa nó) — SUM cả "
                       "bảng cube đếm cùng một quan trắc nhiều lần."),
        ("Đúng thế nào", "Gộp trung bình = TỔNG giá trị / TỔNG số đo hợp lệ — cube lưu sẵn SUM + COUNT "
                         "cho từng ô nên gộp đúng mọi cấp."),
        ("Ví dụ", "24 giờ × 100 + 6 giờ × 200 → 3.600 / 30 = 120 — KHÔNG PHẢI (100 + 200) / 2 = 150."),
    ])

    # A10 (42)
    s = d.shield("A10", "Kiến trúc OLAP của dự án: ROLAP", "«Nhóm dùng kiến trúc lưu trữ OLAP gì?»",
                 "Mở khi hỏi MOLAP/ROLAP. Khung: ROLAP — dữ liệu quan hệ trong PostgreSQL, tổng hợp SQL + "
                 "khối tính trước trong mart.")
    blocks(s, [
        ("Dự án dùng", "ROLAP: dữ liệu quan hệ trong PostgreSQL; tổng hợp bằng SQL; khối tính trước nằm "
                       "trong mart (bảng quan hệ) — không có engine ma trận đa chiều riêng."),
        ("Vì sao đủ", f"{vn_int(420768)} dòng × 5 thuộc tính cube — PostgreSQL trả lời trong vài ms nhờ "
                      "khối tính trước; chưa cần MOLAP."),
        ("Khi nào đổi", "Nếu 100 trạm × 10 năm (dữ liệu lớn hơn ~2 bậc) → cân nhắc columnar / MPP / Spark "
                        "cho phễu tổng hợp."),
    ])

    # A11 (43)
    s = d.shield("A11", "Khóa & SCD — chiều thay đổi thì sao?", "«Dùng khóa gì? Trạm đổi quận thì sao?»",
                 "Mở khi hỏi khóa/SCD. Khung: khóa tự nhiên/biện minh — hợp lý vì chiều tĩnh, nguồn duy nhất; "
                 "nếu thuộc tính đổi → surrogate key + SCD type 2.")
    blocks(s, [
        ("Hiện dùng", "Khóa tự nhiên/biện minh: station_key = station_id · date_key = YYYYMMDD · time_key "
                      "= 0–23 · wind_key: 0 = không rõ hướng."),
        ("Vì sao ổn", "Chiều TĨNH (trạm không đổi quận trong dữ liệu), nguồn DUY NHẤT, không có thành viên "
                      "thay đổi → khóa tự nhiên đơn giản, đọc được, không cần indirection."),
        ("Nếu thay đổi", "Trạm đổi quận → surrogate key + SCD TYPE 2: bản ghi chiều MỚI + valid_from / "
                         "valid_to; fact cũ vẫn trỏ bản ghi cũ → lịch sử KHÔNG bị viết lại."),
    ])

    # A12 (44)
    s = d.shield("A12", "Chia tập, k-fold & rò rỉ — 420.768 sao còn 372.499?",
                 "«Sao không k-fold? Rò rỉ ở đâu?»",
                 "Mở khi hỏi chia tập/rò rỉ/số dòng. 4 khối: vì sao theo thời gian · 372.499 ≠ 420.768 · rà "
                 "rò rỉ · vì sao cây & HistGB.")
    blocks(s, [
        ("Không k-fold", "Dữ liệu giờ LIÊN TỤC: giờ láng giềng gần như trùng nhau — random split (k-fold) "
                         "cho giờ đó vào cả train lẫn test → kết quả lạc quan. Chia theo thời gian: train → "
                         "val → test; siêu tham số chọn trên val; test dùng đúng một lần."),
        ("372.499 dòng", "Cửa sổ ML chốt 03/2013 – 12/2016 (01–02/2017 nằm NGOÀI ML — dành cho đối soát "
                         "CDC, xem A14). Trong cửa sổ: dòng thiếu nhãn PM2.5(t+1) hoặc không đủ chuỗi 24 "
                         "giờ trước (lag 1–24h + trung bình trượt) bị loại — mỗi giờ thiếu làm gãy chuỗi, "
                         "kéo theo loại cả những giờ ngay sau nó."),
        ("Rò rỉ đã rà", "Cửa sổ rolling và nhãn KHÔNG tràn qua biên split — có kiểm tra tự động; từng phát "
                        "hiện 1 trường hợp tràn và refit: F1 0,926 so với 0,927 — nói thật, không khẳng "
                        "định hoàn hảo."),
        ("Vì sao cây & HistGB", "Cây quyết định: dễ giải thích (vẽ ra trình bày được) + mạnh trên bảng "
                                "đặc trưng đã dựng. HistGB: boosting — cây yếu tuần tự sửa lỗi của cây "
                                "trước, histogram bin cho nhanh."),
    ], size=14, gap=8)

    # A13 (45)
    s = d.shield("A13", "McNemar & hai ngưỡng khác nhau", "«McNemar là gì? 115 với 0,50 là gì?»",
                 "Mở khi hỏi McNemar/ngưỡng. Khung: so 2 mô hình trên cùng tập test, chỉ xét cặp disagree.")
    blocks(s, [
        ("McNemar", "So hai mô hình trên CÙNG tập test, chỉ xét các cặp không thống nhất (disagree): "
                    f"b = {vn_int(f.mcnemar['b'])} giờ chỉ mô hình đúng · c = {vn_int(f.mcnemar['c'])} giờ "
                    f"chỉ persistence đúng → {sci(f.mcnemar['p_value'])}: khác biệt thật, không phải may."),
        ("Ý nghĩa", "Thống kê có ý nghĩa ≠ đủ để triển khai: test chỉ 6 tháng, dữ liệu lịch sử — nói thẳng "
                    "trong hạn chế."),
        ("Hai ngưỡng", "115 µg/m³ tạo NHÃN từ nồng độ (định nghĩa bài toán); 0,50 là NGƯỠNG XÁC SUẤT ra "
                       "quyết định — DỊCH ĐƯỢC (cây dùng 0,45) để đổi recall/precision theo giá của bỏ sót "
                       "(FN) vs báo nhầm (FP)."),
    ])

    # A14 (46)
    cdc = load_json("artifacts/cdc_eval/mae_by_batch.json")
    s = d.shield("A14", "CDC chi tiết: sửa/xoá, nhân đôi, vượt ngưỡng MAE",
                 "«Bản ghi cũ bị sửa thì sao? Chạy lại có nhân đôi? Vượt MAE rồi làm gì?»",
                 "Mở khi hỏi sâu CDC. Khung: watermark bắt insert; sửa/xoá sau mốc không bắt → đối soát "
                 "sha256/load lại partition; load_registry chống nhân đôi; MAE từng batch chỉ bật cờ.")
    blocks(s, [
        ("Watermark", "Mốc event_time đã nạp; mỗi lần chỉ lấy bản ghi SAU mốc → bắt được INSERT mới."),
        ("Sửa / xoá cũ", "Sau khi mốc đã đi qua, bản ghi cũ bị SỬA/XOÁ không xuất hiện lại → KHÔNG tự bắt "
                         "được. Phương án: đối soát sha256 từng batch, hoặc load lại cả partition."),
        ("Không nhân đôi", "Upsert theo khóa grain + load_registry: object key + sha256 + schema_version + "
                           "row_count khớp → skip batch cũ."),
        ("Giám sát MAE", f"Mỗi batch CDC chấm MAE với số đo thật: tổng {vn(cdc['overall']['mae'], 2)} "
                         f"µg/m³ trên {vn_int(cdc['overall']['scored_n'])} giờ, chấm liên tục 01/2016 – "
                         f"28/02/2017; ngưỡng {vn(cdc['policy']['mae_alert_threshold'], 0)} µg/m³ — vượt chỉ "
                         "BẬT CỜ (monitor_only): mô hình khoá, không tự retrain. MAE cao hơn 8,87 của test "
                         "vì đối soát chạy trên dòng thời gian vận hành, không phải một tập test cố định."),
    ], size=14, gap=8)

    # A15 (47)
    s = d.shield("A15", "Đợt ô nhiễm & thí nghiệm 6 giờ", "«Đợt là gì? Vì sao kiểm tra tới t+8?»",
                 "Mở khi hỏi định nghĩa đợt/thí nghiệm 6h. Khung: đợt trạm ≥ 3h liên tục ≥ 115; đa trạm ≥ 6 "
                 "trạm; đợt muộn nhất t+6 cần xem tới t+8.")
    blocks(s, [
        ("Đợt — một trạm", "≥ 3 giờ LIÊN TIẾP ≥ 115 µg/m³ tại trạm đó; thiếu số đo làm gián đoạn chuỗi "
                           "(không nối qua khoảng thiếu)."),
        ("Đợt — đa trạm", "≥ 6 trạm có đợt chồng lấn thời gian → «sự kiện» vùng: 487 sự kiện đa trạm "
                          "trong 4 năm, thời lượng trung vị 10 giờ."),
        ("Thí nghiệm 6h", "Hỏi: đang sạch, có đợt BẮT ĐẦU trong t+1 → t+6 không. Đợt bắt đầu muộn nhất t+6 "
                          "cần theo dõi tới t+8 mới đủ 3 giờ liên tục → nhãn chỉ giữ khi toàn bộ cửa sổ "
                          "nằm trong một split."),
        ("Trần cảnh báo nhầm", "0,5 lần/ngày-trạm — chọn trên validation, không chỉnh trên test."),
    ], size=14, gap=8)

    # A16 (48)
    s = d.shield("A16", "Hồi quy & phân cụm — đúng đối tượng", "«MAE nghĩa gì? Phân cụm cái gì?»",
                 "Mở khi hỏi bổ trợ. Khung: MAE = trung bình |sai số| cùng đơn vị; phân cụm đối tượng là 12 "
                 'TRẠM, không phải 420k giờ.')
    blocks(s, [
        ("MAE", f"{vn(f.reg_mae('hist_gradient_boosting'), 2)} µg/m³ = trung bình trị tuyệt đối sai số "
                "(cùng đơn vị với PM2.5) — KHÔNG phải trần sai số từng dự báo."),
        ("Phân cụm cái gì", "Đối tượng là 12 TRẠM — mỗi trạm thành một vector hồ sơ (mùa, chất, đợt…), "
                            "chuẩn hoá trước K-Means; không phải phân cụm 420 nghìn giờ."),
        ("Chọn K", "Thử K = 2, 3, 4: silhouette 0,664 / 0,345 / 0,277 → K = 2 tách cụm rõ nhất; cụm "
                   "9 trạm + 3 trạm. Silhouette đo tách cụm theo đặc trưng — không phải accuracy."),
        ("Địa lý", "Chỉ dùng để ĐỐI CHIẾU kết quả cụm (cụm có liên hệ vị trí không), không đưa vào đặc "
                   "trưng phân cụm."),
    ], size=14, gap=8)

    # A17 (49)
    s = d.shield("A17", "Demo 3 phút & tái lập", "«Chạy lại được không? Mất mạng thì sao?»",
                 "Mở khi hỏi demo/tái lập. Khung: luồng demo 4 trang + fallback notebook + run_all 11,2 phút.")
    blocks(s, [
        ("Luồng demo 3'", "/ (quy mô) → /olap (1 roll-up + 1 slice) → /canh-bao (nhãn, baseline, kết quả) "
                          "→ /phat-lai (dự đoán trước — đối soát sau)."),
        ("Mất mạng", "notebooks/olap/olap_demo.ipynb đã chạy sẵn output; ảnh chụp màn hình các trang đã "
                     "lưu trong slide."),
        ("Tái lập", f"python scripts/run_all.py — {f.dbt_pass}/{f.dbt_total} dbt test pass · "
                    f"{vn(FROM_SCRATCH_MINUTES, 1)} phút từ trạng thái trống · idempotent (load_registry "
                    "chống nạp đôi)."),
    ])

    assert d.n == N_TOTAL, f"kỳ vọng {N_TOTAL} slide, có {d.n}"
    return d.prs


# --------------------------------------------------------------------------- QA
def qa(prs: Presentation) -> list[str]:
    """Ước lượng overflow text và hộp vượt khung slide — heuristic, sửa các lỗi lớn."""
    issues = []
    sw, sh = prs.slide_width, prs.slide_height
    for idx, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not shp.has_text_frame or not shp.text_frame.text.strip():
                continue
            if shp.left is None:
                continue
            if shp.left < 0 or shp.top < 0 or shp.left + shp.width > sw + Inches(0.02) \
                    or shp.top + shp.height > sh + Inches(0.02):
                issues.append(f"S{idx}: hộp vượt khung — «{shp.text_frame.text[:38]}»")
            w_in = shp.width.inches - 0.1
            est_h = 0.0
            for p in shp.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    est_h += 0.12
                    continue
                size = max((r.font.size.pt if r.font.size else 16) for r in p.runs)
                cpl = max(6, w_in / (size * 0.0071))
                lines = max(1, math.ceil(len(txt) / cpl))
                est_h += lines * size * 1.25 / 72 + (p.space_after.pt / 72 if p.space_after else 0)
            if est_h > shp.height.inches + 0.12:
                issues.append(f"S{idx}: nghi overflow ({est_h:.2f}\" vs {shp.height.inches:.2f}\") — "
                              f"«{shp.text_frame.text[:44]}»")
    return issues


def main() -> int:
    prs = build(Facts())
    prs.save(str(OUT))
    print(OUT.name, len(prs.slides), "slides")
    for msg in qa(prs):
        print("QA:", msg)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
