"""python-docx helpers: A4 Vietnamese academic layout, captions, tables, TOC and page numbers."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


def _set_font(style_or_run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    font = style_or_run.font
    font.name = FONT
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    # Built-in heading styles point at theme fonts, which win over w:ascii → drop them.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]


def _field(paragraph, instr: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    run._r.append(text)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self.chapter = 0
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
        styles = self.doc.styles
        normal = styles["Normal"]
        _set_font(normal, 13)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.25
        for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
            st = styles[name]
            _set_font(st, size, bold=True, color=RGBColor(0, 0, 0))
            st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number", "Caption"):
            _set_font(styles[name], 13 if name != "Caption" else 11)
        styles["Caption"].font.italic = True
        styles["Caption"].font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # --- text --------------------------------------------------------------
    def p(self, text: str, *, bold_prefix: str | None = None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          italic: bool = False, size: float | None = None):
        para = self.doc.add_paragraph()
        para.alignment = align
        if bold_prefix:
            r = para.add_run(bold_prefix)
            r.bold = True
            if size:
                r.font.size = Pt(size)
        r = para.add_run(text)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        return para

    def rich(self, parts: list[tuple[str, str]], align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        """parts = [("b", "bold text"), ("", "plain"), ("i", "italic"), ("c", "code")]"""
        para = self.doc.add_paragraph()
        para.alignment = align
        for kind, text in parts:
            r = para.add_run(text)
            r.bold = "b" in kind
            r.italic = "i" in kind
            if "c" in kind:
                r.font.name = "Consolas"
                r.font.size = Pt(11)
        return para

    def bullets(self, items: list[str | list[tuple[str, str]]], style: str = "List Bullet") -> None:
        for item in items:
            para = self.doc.add_paragraph(style=style)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.space_after = Pt(2)
            if isinstance(item, str):
                para.add_run(item)
            else:
                for kind, text in item:
                    r = para.add_run(text)
                    r.bold = "b" in kind
                    r.italic = "i" in kind
                    if "c" in kind:
                        r.font.name = "Consolas"
                        r.font.size = Pt(11)

    def code(self, text: str) -> None:
        for line in text.strip("\n").splitlines():
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.left_indent = Cm(0.5)
            r = para.add_run(line if line else " ")
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F1F5F9")
            para._p.get_or_add_pPr().append(shd)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # --- structure ---------------------------------------------------------
    def h1(self, text: str) -> None:
        self.chapter += 1
        self.fig_no = 0
        self.tab_no = 0
        self.doc.add_heading(f"Chương {self.chapter}. {text}", level=1)

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)

    def h3(self, text: str) -> None:
        self.doc.add_heading(text, level=3)

    def unnumbered_h1(self, text: str) -> None:
        self.doc.add_heading(text, level=1)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- figures & tables --------------------------------------------------
    def figure(self, path: Path, caption: str, width_cm: float = 15.0) -> None:
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        self.fig_no += 1
        cap = self.doc.add_paragraph(f"Hình {self.chapter}.{self.fig_no}. {caption}", style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, headers: list[str], rows: list[list[str]], widths_cm: list[float], caption: str | None,
              font_size: float = 10.5, align_right: set[int] | None = None) -> None:
        if caption:
            self.tab_no += 1
            cap = self.doc.add_paragraph(f"Bảng {self.chapter}.{self.tab_no}. {caption}", style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = True
        align_right = align_right or set()
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for r_i, values in enumerate([headers, *rows]):
            for c_i, value in enumerate(values):
                cell = t.cell(r_i, c_i)
                cell.width = Cm(widths_cm[c_i])
                cell.text = ""
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                if c_i in align_right and r_i > 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = para.add_run(str(value))
                run.font.size = Pt(font_size)
                if r_i == 0:
                    run.bold = True
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "DCE6F1")
                    cell._tc.get_or_add_tcPr().append(shd)
        # repeat header row on page breaks
        trpr = t.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trpr.append(hdr)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- fields ------------------------------------------------------------
    def toc(self) -> None:
        para = self.doc.add_paragraph()
        _field(para, 'TOC \\o "1-2" \\h \\z \\u', "Nhấn chuột phải → Update Field (hoặc F9) để cập nhật mục lục.")

    def start_numbered_section(self) -> None:
        """New section (new page) whose footer shows the page number."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.footer.is_linked_to_previous = False
        para = sec.footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(para, "PAGE", "1")
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        sec._sectPr.append(pg)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
