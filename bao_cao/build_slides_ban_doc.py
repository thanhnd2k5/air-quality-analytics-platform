#!/usr/bin/env python3
"""Generate Slide_BTL_BAN_DOC.pptx — TÀI LIỆU ĐỌC HIỂU, bản thiết kế lại (~49 slide).

Bản để ĐỌC (không chiếu): văn xuôi giải thích bên trái + một "bảng hình" minh hoạ bên phải
(số to, thẻ so sánh, sơ đồ, timeline, thanh) cho MỖI slide; cuối slide có dải GHI NHỚ.
Nội dung + số liệu bám KE_HOACH_NOI_DUNG_THUYET_TRINH.md và ON_TAP_BAO_VE_HOI_DAP.md
(đã đối chiếu artifacts 24/10/2026). Bộ chiếu phòng thủ vẫn là
Slide_BTL_THUYET_TRINH_30phut_v3.pptx.

    py -3 bao_cao/build_slides_ban_doc.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

OUT = HERE / "Slide_BTL_BAN_DOC.pptx"

# Palette — đồng bộ bộ v3: charcoal chủ đạo, cam cảnh báo điểm nhấn, teal hỗ trợ.
INK = RGBColor(0x1F, 0x2A, 0x37)
INK_SOFT = RGBColor(0x33, 0x41, 0x55)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
MUTED = RGBColor(0x64, 0x74, 0x8B)
SLATE_L = RGBColor(0xC5, 0xD0, 0xDE)
AMBER = RGBColor(0xFD, 0xBA, 0x74)
SOFT_TEAL = RGBColor(0xD6, 0xF5, 0xEE)
SOFT_ORANGE = RGBColor(0xFF, 0xEC, 0xDC)
CARD = RGBColor(0xF2, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

SERIF = "Cambria"
SANS = "Calibri"
MONO = "Courier New"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CW = W - 2 * MARGIN
BODY_Y = Inches(1.5)
BAND_Y, BAND_H = Inches(6.10), Inches(1.40)
PROSE_W = Inches(7.15)
PANEL_X, PANEL_W = Inches(8.0), Inches(4.73)

TOKEN = re.compile(r"(\*\*.+?\*\*|##.+?##)")


def parse_runs(s: str):
    out = []
    for part in TOKEN.split(s):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((True, False, part[2:-2]))
        elif part.startswith("##") and part.endswith("##"):
            out.append((True, True, part[2:-2]))
        else:
            out.append((False, False, part))
    return out


class Doc:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.pages: list[tuple[object, bool]] = []

    # --- khung -----------------------------------------------------------------
    def light(self, kicker: str, title: str, title_size: int = 27):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = WHITE
        self.text(s, kicker, MARGIN, Inches(0.16), Inches(10.5), Inches(0.3), 11, MUTED,
                  bold=True, font=SANS)
        self.text(s, title, MARGIN, Inches(0.5), CW, Inches(0.95), title_size, INK,
                  bold=True, font=SERIF)
        self.pages.append((s, False))
        return s

    def dark(self, kicker: str, title: str, lines: list[str], big: str | None = None,
             title_y: float = 1.95):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = INK
        if big:
            self.text(s, big, Inches(9.3), Inches(1.2), Inches(3.6), Inches(4.2), 220,
                      INK_SOFT, bold=True, font=SERIF, align=PP_ALIGN.RIGHT)
        self.text(s, kicker, MARGIN, Inches(0.85), Inches(10), Inches(0.35), 14, AMBER,
                  bold=True, font=SANS)
        self.text(s, title, MARGIN, Inches(title_y), Inches(8.6), Inches(0.85), 36, WHITE,
                  bold=True, font=SERIF)
        y = title_y + 1.25
        for ln in lines:
            box = self.text(s, "", MARGIN, Inches(y), Inches(8.3), Inches(0.4), 15.5,
                            SLATE_L, font=SANS)
            p = box.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "▸  "
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = SANS, Pt(15.5), True, ORANGE
            for bold, orange, txt in parse_runs(ln):
                r = p.add_run()
                r.text = txt
                r.font.name, r.font.size = SANS, Pt(15.5)
                r.font.bold = bold
                r.font.color.rgb = AMBER if orange else SLATE_L
            y += 0.52
        self.pages.append((s, True))
        return s

    # --- nguyên thuỷ -------------------------------------------------------------
    def text(self, slide, text, x, y, w, h, size=15, color=INK, bold=False, italic=False,
             font=SERIF, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.03)
        tf.margin_top = tf.margin_bottom = 0
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if spacing:
                p.space_after = Pt(spacing)
            r = p.add_run()
            r.text = line
            r.font.name, r.font.size, r.font.bold, r.font.italic = font, Pt(size), bold, italic
            r.font.color.rgb = color
        return box

    def prose(self, slide, paras: list[str], y=None, h=None, size=14.5, gap=9,
              x=MARGIN, w=PROSE_W):
        box = slide.shapes.add_textbox(x, y or BODY_Y, w, h or Inches(4.4))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.03)
        tf.margin_top = tf.margin_bottom = 0
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.JUSTIFY
            p.space_after = Pt(gap)
            p.line_spacing = 1.13
            for bold, orange, txt in parse_runs(para):
                r = p.add_run()
                r.text = txt
                r.font.name, r.font.size, r.font.bold = SERIF, Pt(size), bold
                r.font.color.rgb = ORANGE if orange else INK
        return box

    def remember(self, slide, sentence: str):
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, BAND_Y, W, BAND_H)
        band.fill.solid()
        band.fill.fore_color.rgb = SOFT_TEAL
        band.line.fill.background()
        band.shadow.inherit = False
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN,
                                      BAND_Y + Inches(0.17), Inches(1.35), Inches(0.4))
        chip.adjustments[0] = 0.5
        chip.fill.solid()
        chip.fill.fore_color.rgb = TEAL
        chip.line.fill.background()
        chip.shadow.inherit = False
        tf = chip.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "GHI NHỚ"
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = SANS, Pt(12), True, WHITE
        self.text(slide, sentence, MARGIN + Inches(1.55), BAND_Y + Inches(0.15),
                  CW - Inches(1.55), Inches(1.1), 15, INK, bold=True, font=SERIF,
                  anchor=MSO_ANCHOR.MIDDLE, spacing=2)

    def card(self, slide, x, y, w, h, fill=CARD, line=None, rounded=True):
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shp = slide.shapes.add_shape(shape, x, y, w, h)
        if rounded:
            shp.adjustments[0] = 0.06
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if line:
            shp.line.color.rgb = line
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def chip(self, slide, text, x, y, w, h=Inches(0.42), fill=INK, color=WHITE, size=12,
             font=SANS, rounded=True):
        w_in = Emu(int(w)).inches
        while size > 8 and len(text) * size * 0.0078 > w_in - 0.14:
            size -= 0.5
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shp = slide.shapes.add_shape(shape, x, y, w, h)
        if rounded:
            shp.adjustments[0] = 0.5
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = font, Pt(size), True, color

    def badge(self, slide, num: str, x, y, d=Inches(0.44), fill=ORANGE):
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
        c.fill.solid()
        c.fill.fore_color.rgb = fill
        c.line.fill.background()
        c.shadow.inherit = False
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = num
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = SANS, Pt(14), True, WHITE

    # --- bảng hình bên phải ---------------------------------------------------------
    def panel(self, slide, title: str):
        """Thẻ nội dung bên phải; trả về (x, y, w, đáy)."""
        self.card(slide, PANEL_X, BODY_Y, PANEL_W, Inches(4.45), CARD)
        self.text(slide, title, PANEL_X + Inches(0.25), BODY_Y + Inches(0.16),
                  PANEL_W - Inches(0.5), Inches(0.3), 11, MUTED, bold=True, font=SANS)
        return PANEL_X + Inches(0.25), BODY_Y + Inches(0.56), PANEL_W - Inches(0.5), 5.82

    def stat(self, slide, x, y, w, value, label, vsize=26, color=ORANGE, vfont=SERIF,
             hvalue: float | None = None, label_lines: int = 2):
        hv = Inches(hvalue) if hvalue else Inches(vsize * 1.5 / 72 + 0.06)
        self.text(slide, value, x, y, w, hv, vsize, color, bold=True, font=vfont)
        self.text(slide, label, x, y + hv + Inches(0.02), w,
                  Inches(0.2 * label_lines + 0.05), 11.5, MUTED, font=SANS, spacing=1)
        return y + hv + Inches(0.28)

    def hline(self, slide, x, y, w, color=LINE):
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(0.75))
        ln.fill.solid()
        ln.fill.fore_color.rgb = color
        ln.line.fill.background()
        ln.shadow.inherit = False

    def kv(self, slide, x, y, w, term, meaning, size=12.5, h=0.66, term_color=TEAL):
        box = slide.shapes.add_textbox(x, y, w, Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.line_spacing = 1.05
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = term + "  "
        r.font.name, r.font.size, r.font.bold = SANS, Pt(size), True
        r.font.color.rgb = term_color
        for bold, orange, txt in parse_runs(meaning):
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = SANS, Pt(size), bold
            r.font.color.rgb = ORANGE if orange else INK
        return y + Inches(h)

    def mini_card(self, slide, x, y, w, height, tag, lines, fill=SOFT_TEAL, tag_color=TEAL,
                  size=12, **kw):
        h = height
        self.card(slide, x, y, w, h, fill)
        self.text(slide, tag, x + Inches(0.18), y + Inches(0.13), w - Inches(0.36),
                  Inches(0.3), 13, tag_color, bold=True, font=SANS)
        box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.47),
                                       w - Inches(0.36), h - Inches(0.55))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = 0
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.08
            p.space_after = Pt(3)
            for bold, orange, txt in parse_runs(ln):
                r = p.add_run()
                r.text = txt
                r.font.name, r.font.size, r.font.bold = SANS, Pt(size), bold
                r.font.color.rgb = ORANGE if orange else INK

    def flow_rows(self, slide, x, y, w, rows, size=11, chip_h=Inches(0.4), gap=Inches(0.12)):
        """rows = list[list[str]]; mỗi row một hàng chip nối tiếp bằng →."""
        for row in rows:
            n = len(row)
            arrow_w = Inches(0.24)
            cw = Emu(int((w - (n - 1) * arrow_w) / n))
            cx = x
            for i, t in enumerate(row):
                self.chip(slide, t, cx, y, cw, chip_h, fill=INK_SOFT, size=size)
                cx += cw
                if i < n - 1:
                    self.text(slide, "→", cx - Inches(0.01), y + Inches(0.03), arrow_w,
                              chip_h, size + 2, ORANGE, bold=True, font=SANS,
                              align=PP_ALIGN.CENTER)
                    cx += arrow_w
            y += chip_h + gap
        return y

    def star_diagram(self, slide, x, y, w):
        """Sơ đồ sao nhỏ: fact giữa + 4 chiều quanh."""
        cx = x + Emu(int(w / 2)) - Inches(1.45)
        self.card(slide, cx, y + Inches(1.32), Inches(2.9), Inches(0.86), SOFT_ORANGE)
        self.text(slide, "fact_air_quality_hourly", cx + Inches(0.1), y + Inches(1.4),
                  Inches(2.7), Inches(0.35), 11.5, ORANGE, bold=True, font=MONO,
                  align=PP_ALIGN.CENTER)
        self.text(slide, "độ đo + 4 khóa chiều", cx + Inches(0.1), y + Inches(1.75),
                  Inches(2.7), Inches(0.3), 10.5, INK, font=SANS, align=PP_ALIGN.CENTER)
        dims = [("dim_date", "ngày→năm · mùa", 0.0, 0.0),
                ("dim_time", "giờ · buổi", 0.0, 2.5),
                ("dim_station", "trạm · quận", 2.25, 0.0),
                ("dim_wind", "hướng · bắc/nam", 2.25, 2.5)]
        for name, sub, dx, dy in dims:
            bx = x + Inches(dx)
            by = y + Inches(dy)
            self.card(slide, bx, by, Inches(1.95), Inches(0.78), WHITE, line=LINE)
            self.text(slide, name, bx + Inches(0.1), by + Inches(0.07), Inches(1.75),
                      Inches(0.3), 11, TEAL, bold=True, font=MONO, align=PP_ALIGN.CENTER)
            self.text(slide, sub, bx + Inches(0.08), by + Inches(0.38), Inches(1.79),
                      Inches(0.3), 10, MUTED, font=SANS, align=PP_ALIGN.CENTER)
            # nối tới fact
            fx1 = bx + Inches(1.95) if dx == 0 else bx
            mx = (x + Emu(int(w / 2)))
            self.hline(slide, min(fx1, mx), y + Inches(1.7) + Inches(0.02),
                       abs(mx - min(fx1, mx)), LINE)

    def bar_rows(self, slide, x, y, w, rows, max_val, unit=""):
        """rows = (label, value, text); vẽ thanh ngang tỷ lệ theo max_val."""
        for label, val, text in rows:
            self.text(slide, label, x, y, Inches(0.85), Inches(0.3), 12, INK, bold=True,
                      font=SANS)
            bw = Emu(int(Inches(2.0) * val / max_val))
            bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.95), y,
                                         bw, Inches(0.3))
            bar.adjustments[0] = 0.5
            bar.fill.solid()
            bar.fill.fore_color.rgb = ORANGE
            bar.line.fill.background()
            bar.shadow.inherit = False
            self.text(slide, f"+{text}{unit}", x + Inches(1.0) + bw + Inches(0.08), y,
                      Inches(1.1), Inches(0.3), 13, ORANGE, bold=True, font=SERIF)
            y += Inches(0.52)
        return y

    def numbered_rows(self, slide, rows: list[tuple[str, str, str]], y=None, size=15.5,
                      gap: float = 0.9):
        y = y or 1.66
        for num, lead, rest in rows:
            self.badge(slide, num, MARGIN, Inches(y + 0.02))
            box = slide.shapes.add_textbox(MARGIN + Inches(0.68), Inches(y),
                                           CW - Inches(0.68), Inches(0.8))
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.02)
            tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.line_spacing = 1.1
            r = p.add_run()
            r.text = lead + " "
            r.font.name, r.font.size, r.font.bold = SERIF, Pt(size), True
            r.font.color.rgb = INK
            for bold, orange, txt in parse_runs(rest):
                r = p.add_run()
                r.text = txt
                r.font.name, r.font.size, r.font.bold = SERIF, Pt(size), bold
                r.font.color.rgb = ORANGE if orange else INK
            y += gap

    def finish(self):
        n = len(self.pages)
        for i, (s, dark) in enumerate(self.pages, 1):
            self.text(s, f"{i} / {n}", W - MARGIN - Inches(1.4), Inches(0.16), Inches(1.4),
                      Inches(0.3), 11, SLATE_L if dark else MUTED, font=SANS,
                      align=PP_ALIGN.RIGHT)
        self.prs.core_properties.title = "Air Quality Analytics Platform — Tài liệu đọc hiểu"
        self.prs.core_properties.author = "Nhóm INT1422"
        self.prs.save(OUT)
        print(f"OK {OUT.name}: {n} slides")


# --------------------------------------------------------------------------- phần chung
def cover_parts(d: Doc):
    s = d.prs.slides.add_slide(d.blank)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK
    d.text(s, "INT1422 · KHO DỮ LIỆU VÀ KHAI PHÁ DỮ LIỆU · PTIT", MARGIN, Inches(0.7),
           Inches(11), Inches(0.35), 13, AMBER, bold=True, font=SANS)
    d.text(s, "TÀI LIỆU ĐỌC HIỂU", MARGIN, Inches(1.2), Inches(6.6), Inches(0.95), 42,
           WHITE, bold=True, font=SERIF)
    d.text(s, "Air Quality Analytics Platform — kho dữ liệu chất lượng không khí Bắc Kinh "
              "& cảnh báo sớm PM2.5 vượt ngưỡng", MARGIN, Inches(2.25), Inches(6.3),
           Inches(1.3), 17, SLATE_L, font=SERIF, spacing=5)
    d.chip(s, "49 slide · đọc hết ≈ 45 phút", MARGIN, Inches(3.75), Inches(3.3), h=Inches(0.44),
           fill=ORANGE, size=13)
    d.chip(s, "số cam = số phải thuộc", MARGIN + Inches(3.5), Inches(3.75), Inches(3.0),
           h=Inches(0.44), fill=INK_SOFT, color=AMBER, size=13)
    d.text(s, "Bộ CHIẾU phòng thủ: Slide_BTL_THUYET_TRINH_30phut_v3.pptx — bộ này để ĐỌC "
              "HIỂU và ÔN, không dùng để trình chiếu.", MARGIN, Inches(4.45), Inches(6.4),
           Inches(0.75), 12.5, MUTED, italic=True, font=SERIF, spacing=3)
    d.hline(s, MARGIN, Inches(5.5), Inches(6.3), INK_SOFT)
    d.text(s, "Phạm Nhật Khánh · Phạm Xuân Kỳ · Nguyễn Đức Thành", MARGIN, Inches(5.68),
           Inches(6.3), Inches(0.35), 13.5, WHITE, bold=True, font=SANS)
    d.text(s, "GVHD: TS. Phan Lý Huỳnh — Học viện Công nghệ Bưu chính Viễn thông", MARGIN,
           Inches(6.08), Inches(6.3), Inches(0.35), 12, SLATE_L, font=SANS)
    parts = [("A", "Bài toán & dữ liệu", "câu hỏi nghiệp vụ · 2 nguồn · làm sạch · kiến trúc",
              "slide 4–14"),
             ("B", "Kho dữ liệu & OLAP", "grain · sao 4 chiều · 180 vs 32 · kiểm chứng",
              "slide 16–28"),
             ("C", "Khai phá dữ liệu", "nhãn · chia tập · baseline · McNemar · tầm dự báo",
              "slide 30–42"),
             ("D", "Vận hành & tổng ôn", "CDC · dashboard · khuyến nghị · bộ số phải thuộc",
              "slide 44–48")]
    y = 1.2
    for code, name, desc, rng in parts:
        d.card(s, Inches(7.35), Inches(y), Inches(5.38), Inches(1.06), INK_SOFT, rounded=True)
        d.badge(s, code, Inches(7.55), Inches(y + 0.3), d=Inches(0.46))
        d.text(s, name, Inches(8.2), Inches(y + 0.14), Inches(3.4), Inches(0.35), 15,
               WHITE, bold=True, font=SANS)
        d.text(s, desc, Inches(8.2), Inches(y + 0.52), Inches(4.35), Inches(0.42), 11,
               SLATE_L, font=SANS, spacing=1)
        d.text(s, rng, Inches(11.55), Inches(y + 0.14), Inches(1.05), Inches(0.35), 11,
               AMBER, bold=True, font=SANS, align=PP_ALIGN.RIGHT)
        y += 1.24
    d.pages.append((s, True))
    return s


# --------------------------------------------------------------------------- build
def build() -> None:
    d = Doc()

    # ---- 1 · bìa ------------------------------------------------------------------
    cover_parts(d)

    # ---- 2 · cách đọc ----------------------------------------------------------------
    s = d.light("PHẦN 0 · CÁCH DÙNG TÀI LIỆU NÀY", "Đọc tài liệu này thế nào cho đỡ tốn công")
    d.prose(s, [
        "Tài liệu này viết theo kiểu **sách đọc**: mỗi slide là một khái niệm giải thích trọn "
        "vẹn, thuật ngữ được **định nghĩa ngay trước khi dùng** — đọc lần lượt từ đầu là hiểu "
        "toàn bộ đề tài mà không phải quay lại tài liệu gốc. Bên trái mỗi slide là phần giải "
        "thích; bên phải là **bảng hình** tóm nội dung — dùng nó để ôn nhanh ở lượt đọc sau.",
        "Ba quy ước: **(1)** ##số in đậm màu cam là số phải thuộc##; **(2)** cuối mỗi slide có "
        "dải xanh **GHI NHỚ** đúng một câu; **(3)** các đoạn hay bị hỏi trong bảo vệ (vì sao "
        "115, vì sao không nội suy, SCD là gì…) được giải thích ngay trong văn.",
        "Sau khi đọc xong, chuyển sang luyện đáp với ##ON_TAP_BAO_VE_HOI_DAP.md## (tự trả lời "
        "không nhìn), cuối cùng tập nói theo ##KE_HOACH_NOI_DUNG_THUYET_TRINH.md## kèm bộ "
        "chiếu v3.",
    ])
    x, y, w, bot = d.panel(s, "BA CẤP TỐC ĐỘ ÔN")
    y = d.stat(s, x, y, w, "45 phút", "đọc hết một lượt từ đầu đến cuối — vừa làm ở lần đầu",
               vsize=24)
    d.hline(s, x, y + Inches(0.02), w)
    y = d.stat(s, x, y + Inches(0.14), w, "15 phút",
               "đọc lại các dải GHI NHỚ + 4 slide CHỐT cuối mỗi phần", vsize=24)
    d.hline(s, x, y + Inches(0.02), w)
    y = d.stat(s, x, y + Inches(0.14), w, "+30 phút",
               "ON_TAP_BAO_VE_HOI_DAP.md — tự trả lời, yếu chỗ nào quay lại đúng slide đó",
               vsize=24)
    d.remember(s, "Đọc 1 lượt hiểu mạch → luyện Q&A → tập nói theo kịch bản — ba bước, mỗi "
                  "bước một tài liệu riêng.")

    # ---- PHẦN A -------------------------------------------------------------------
    d.dark("PHẦN A · SLIDE 4–14", "Bài toán & dữ liệu",
           ["Câu hỏi nghiệp vụ: vì sao cần cảnh báo sớm",
            "Hai nguồn dữ liệu mở và quy tắc ghép chúng",
            "Chỗ thiếu: vì sao KHÔNG điền hết cho đẹp",
            "Kiến trúc tổng: hai luồng vào, mỗi lớp một việc"], big="A")

    s = d.light("PHẦN A · BÀI TOÁN", "Đề tài bắt đầu từ một câu hỏi của người vận hành")
    d.prose(s, [
        "Hãy tưởng tượng mình là cán bộ trung tâm quan trắc: **mỗi giờ, với từng trạm**, bạn "
        "phải quyết định có phát cảnh báo hạn chế hoạt động ngoài trời hay không. Cách làm "
        "hiện tại là chờ **số đo vượt ngưỡng rồi mới báo** — cảnh báo luôn **trễ 1 giờ**, phát "
        "ra sau khi người dân đã ở ngoài trời trong giờ ô nhiễm.",
        "Đề tài đổi câu hỏi lại: **“Giờ tới, tại trạm này, PM2.5 có ≥ 115 µg/m³ không?”** — "
        "trả lời TRƯỚC khi giờ đó đến. Đây là bài toán **phân loại nhị phân** (vượt / không "
        "vượt), không phải dự báo con số tuyệt đối. Và vấn đề không hiếm: trung bình 4 năm "
        "đã ##79,8 µg/m³##, khoảng ##1/4 số giờ## nằm trên ngưỡng 115.",
    ])
    x, y, w, bot = d.panel(s, "QUY MÔ VẤN ĐỀ")
    y = d.stat(s, x, y, w, "79,8 µg/m³", "PM2.5 trung bình — 12 trạm, 4 năm")
    d.hline(s, x, y, w)
    y = d.stat(s, x, y + Inches(0.12), w, "≈ 1/4 số giờ", "có PM2.5 ≥ 115 µg/m³")
    d.hline(s, x, y, w)
    d.mini_card(s, x, y + Inches(0.14), w, Inches(1.1), "CÁCH HIỆN TẠI",
                ["“chờ số đo vượt rồi báo”", "→ luôn trễ 1 giờ"], fill=SOFT_ORANGE,
                tag_color=ORANGE)
    d.remember(s, "Đề tài = cảnh báo sớm cho GIỜ TỚI, thay cho cách “chờ số đo vượt rồi "
                  "báo” luôn trễ 1 giờ.")

    s = d.light("PHẦN A · BÀI TOÁN", "Hai nửa của môn học trong một hệ thống")
    d.prose(s, [
        "Đề tài ôm trọn hai nửa học phần. **Nửa kho dữ liệu (OLAP):** tổ chức 4 năm dữ liệu "
        "để trả lời “đã xảy ra gì, ở đâu, lúc nào, kèm điều kiện gì” ở mọi góc nhìn. **Nửa "
        "khai phá:** học từ quá khứ để dự đoán giờ tới — bài toán phân loại cảnh báo.",
        "Quan trọng: đây là **một hệ thống, không phải hai bài riêng** — cùng một luồng dữ "
        "liệu đi qua ##nguồn → lake → staging → kho → mart → đặc trưng → mô hình → "
        "dashboard##. Kho dựng đúng thì đặc trưng cho mô hình mới sạch; mô hình tốt mà kho "
        "sai số liệu thì kết quả cũng sai theo.",
    ])
    x, y, w, bot = d.panel(s, "MỘT LUỒNG, HAI NỬA")
    d.mini_card(s, x, y, w, Inches(1.28), "NỬA KHO — OLAP",
                ["tổ chức quá khứ để phân tích", "mọi góc nhìn: mùa · quận · gió"])
    d.mini_card(s, x, y + Inches(1.42), w, Inches(1.28), "NỬA KHAI PHÁ",
                ["học từ quá khứ", "→ dự đoán cảnh báo giờ tới"], fill=SOFT_ORANGE,
                tag_color=ORANGE)
    d.flow_rows(s, x, y + Inches(2.92), w, [["nguồn", "lake", "staging"],
                                            ["kho", "mart", "ML", "dashboard"]], size=10.5)
    d.remember(s, "Kho trả lời QUÁ KHỨ (OLAP), mô hình trả lời TƯƠNG LAI (phân loại) — cùng "
                  "một luồng dữ liệu.")

    s = d.light("PHẦN A · DỮ LIỆU", "Dữ liệu chính: UCI #501 — 420.768 giờ đo, 12 trạm, 4 năm")
    d.prose(s, [
        "Nguồn chính là bộ dữ liệu mở ##UCI #501 — Beijing Multi-Site Air Quality##: mỗi bản "
        "ghi là một lần đo **trạm × giờ**, gom lại thành ##420.768 dòng## từ ##12 trạm## "
        "trong ##4 năm## (03/2013 – 02/2017). Đọc số cho đúng: 420.768 là **số dòng "
        "trạm×giờ** (12 × 35.064), không phải 420.768 mốc thời gian, và cũng không có nghĩa "
        "mọi ô đều có số đo.",
        "Mỗi bản ghi ##12 thuộc tính##: **6 chất ô nhiễm** (PM2.5, PM10, SO2, NO2, CO, O3), "
        "**5 biến khí tượng** (nhiệt độ, áp suất, điểm sương, mưa, tốc độ gió) và **hướng "
        "gió** 16 hướng. Bộ dữ liệu mở ##CC BY 4.0## kèm DOI — trích dẫn được, và **không "
        "chứa dữ liệu cá nhân**.",
    ])
    x, y, w, bot = d.panel(s, "BỘ SỐ DỮ LIỆU")
    y = d.stat(s, x, y, w, "420.768", "dòng giờ-trạm = 12 trạm × 35.064 giờ")
    d.hline(s, x, y, w)
    y = d.stat(s, x, y + Inches(0.12), w, "12 + 12", "trạm quan trắc · thuộc tính mỗi bản ghi")
    d.hline(s, x, y, w)
    y = d.kv(s, x, y + Inches(0.16), w, "6 chất ô nhiễm",
             "PM2.5 · PM10 · SO2 · NO2 · CO · O3")
    y = d.kv(s, x, y, w, "5 khí tượng + gió",
             "nhiệt độ · áp suất · điểm sương · mưa · tốc độ + 16 hướng gió", h=0.66)
    d.chip(s, "CC BY 4.0 · DOI 10.24432/C5RK5G", x, y + Inches(0.06), w, h=Inches(0.4),
           fill=TEAL, size=11.5)
    d.remember(s, "420.768 = dòng trạm×giờ (12 trạm × 4 năm) · 12 thuộc tính = 6 ô nhiễm + 5 "
                  "khí tượng + hướng gió.")

    s = d.light("PHẦN A · DỮ LIỆU", "Nguồn bổ trợ Open-Meteo và quy tắc ghép hai nguồn")
    d.prose(s, [
        "Nguồn phụ là ##Open-Meteo Historical Weather##: dữ liệu thời tiết **tái phân tích "
        "ERA5** (chạy lại mô hình khí tượng trên dữ liệu quá khứ), theo giờ, lấy đúng **12 "
        "toạ độ trạm** của UCI. Nó làm hai việc: **lấp chỗ thiếu thời tiết** của UCI và "
        "**thêm độ che phủ mây** — thứ UCI không đo.",
        "Quy tắc ghép là đoạn hay bị hỏi nhất: UCI có sẵn thì dùng UCI; Open-Meteo chỉ bước "
        "vào khi ô UCI rỗng, và **mỗi ô lấp được ghi rõ nguồn** để truy vết — vì khí tượng "
        "tái phân tích không đồng nhất tuyệt đối với phép đo tại trạm, biết ô nào đến từ "
        "đâu mới kiểm tra được khi nghi ngờ. Độ ẩm tương đối không lấy của Open-Meteo mà "
        "**suy từ TEMP + DEWP** của chính UCI.",
    ])
    x, y, w, bot = d.panel(s, "QUY TẮC GHÉP")
    y = d.kv(s, x, y, w, "Khóa ghép", "cùng trạm + cùng mốc giờ (thống nhất múi giờ)")
    y = d.kv(s, x, y, w, "Ưu tiên", "##UCI có → dùng UCI##; rỗng → mới lấy Open-Meteo")
    y = d.kv(s, x, y, w, "Truy vết", "mỗi ô lấp đều ghi rõ nguồn từng ô")
    y = d.kv(s, x, y, w, "Độ ẩm RH", "suy từ TEMP + DEWP (UCI), không lấy của OM")
    y = d.kv(s, x, y, w, "Thêm mới", "độ che phủ mây — chỉ có ở Open-Meteo")
    d.mini_card(s, x, y + Inches(0.08), w, Inches(0.88), "SAU KHI LẤP",
                ["thời tiết còn ##0 ô thiếu##"])
    d.remember(s, "UCI ưu tiên tuyệt đối — Open-Meteo chỉ lấp ô trống, và ô lấp nào cũng ghi "
                  "rõ nguồn để truy vết.")

    s = d.light("PHẦN A · DỮ LIỆU", "Chỗ thiếu: quyết định từng cột, không “làm đẹp”")
    d.prose(s, [
        "Thiếu dữ liệu không đều: ##CO thiếu nhiều nhất 4,9%## với các **chuỗi thiếu liên "
        "tục dài 343–1.517 giờ**; thời tiết thiếu rất ít (~0,1%). Mỗi cột một cách xử lý — "
        "đây là chuỗi **quyết định**, không phải một thao tác “làm sạch” duy nhất.",
        "Với chuỗi thiếu **dài**: ##không nội suy##. Nội suy chuỗi hàng trăm giờ là **bịa "
        "dữ liệu** — có thể tạo “đợt ô nhiễm giả” hoặc xóa mờ đợt thật. Quyết định: **giữ "
        "NULL + cờ chất lượng**; phân biệt kỹ: ##NULL ≠ 0## (0 là giá trị đo được, NULL là "
        "chưa đo). Với chỗ thiếu **ngắn**: lấp từ Open-Meteo theo quy tắc slide trước, sau "
        "lấp còn ##0 ô thiếu##.",
    ])
    x, y, w, bot = d.panel(s, "MỖI CỘT MỘT QUYẾT ĐỊNH")
    d.mini_card(s, x, y, w, Inches(1.5), "CHUỖI THIẾU DÀI",
                ["CO: chuỗi liên tục 343–1.517 giờ", "→ ##GIỮ NULL## + cờ chất lượng",
                 "(nội suy chuỗi dài = bịa dữ liệu)"])
    d.mini_card(s, x, y + Inches(1.64), w, Inches(1.24), "THIẾU NGẮN — THỜI TIẾT",
                ["thiếu ~0,1%", "→ lấp từ Open-Meteo, còn ##0 ô##"], fill=SOFT_ORANGE,
                tag_color=ORANGE)
    d.chip(s, "CO thiếu 4,9% — nhiều nhất", x, y + Inches(3.04), w, h=Inches(0.42),
           fill=INK_SOFT, color=AMBER, size=12)
    d.remember(s, "Chuỗi thiếu dài → giữ NULL + cờ (nội suy = bịa dữ liệu); thiếu ngắn → lấp "
                  "từ nguồn phụ và ghi nguồn từng ô.")

    s = d.light("PHẦN A · DỮ LIỆU", "Trần thiết bị và nguyên tắc điền thiếu phía học máy")
    d.prose(s, [
        "Có ##60 dòng chạm trần thiết bị## (giá 999 / 10.000 tùy cột) — thiết bị đo tới "
        "giới hạn rồi đứng lại. Quyết định: **giữ nguyên + cờ SUSPECT_CAP**, vì đó là số đo "
        "thật của thiết bị đang quá tải; xóa đi là mất bằng chứng về những giờ nặng nhất.",
        "Phía học máy, một số biến đầu vào thiếu được điền bằng **median** (trung vị — giá "
        "trị đứng giữa sau khi sắp xếp). Nguyên tắc nghiêm ngặt: ##median chỉ được tính trên "
        "tập train## — nếu tính trên toàn bộ, thông tin của val/test đã “rò” vào khâu chuẩn "
        "bị, một dạng rò rỉ dữ liệu làm điểm đánh giá lạc quan hơn thực tế.",
    ])
    x, y, w, bot = d.panel(s, "NGUYÊN TẮC")
    y = d.stat(s, x, y, w, "60 dòng", "chạm trần thiết bị → giữ + cờ SUSPECT_CAP")
    d.chip(s, "NULL ≠ 0", x, y + Inches(0.04), Inches(1.3), h=Inches(0.38), fill=TEAL,
           size=11.5)
    d.text(s, "0 là giá trị đo — NULL là chưa đo", x + Inches(1.45), y + Inches(0.07),
           w - Inches(1.45), Inches(0.3), 11.5, INK, font=SANS)
    y += Inches(0.62)
    d.hline(s, x, y, w)
    d.text(s, "Mọi phép biến đổi “học” từ dữ liệu đều fit trên train:", x,
           y + Inches(0.14), w, Inches(0.32), 12, MUTED, bold=True, font=SANS)
    d.flow_rows(s, x, y + Inches(0.52), w, [["train", "median", "val / test"]], size=11)
    d.remember(s, "Trần thiết bị: giữ + cờ. Mọi phép biến đổi học từ dữ liệu (median, "
                  "scaler) chỉ fit trên tập train.")

    s = d.light("PHẦN A · DỮ LIỆU", "Rời rạc hoá — gom mức chi tiết thành mức phân tích")
    d.prose(s, [
        "Dữ liệu thô quá mịn cho câu hỏi nghiệp vụ: không ai hỏi “quận Dongcheng giờ thứ "
        "17 thế nào”, họ hỏi “**buổi tối mùa đông** thế nào”. Vì vậy một số cột được **rời "
        "rạc hoá** (gom thành nhóm): ##giờ → buổi##, ##ngày → mùa##, ##16 hướng gió → 4 "
        "phần tư → bắc / nam##.",
        "Giới hạn của phép gom phải biết: gom giờ thành buổi **mất khả năng so sánh từng "
        "giờ riêng** — nên rời rạc hoá chỉ tạo **cột bổ sung**, grain gốc mức giờ vẫn giữ "
        "nguyên (phần B). Các mức này chính là **phân cấp của bảng chiều** sau này (buổi "
        "trong dim_time, mùa trong dim_date, phần tư gió trong dim_wind) — làm một lần, "
        "dùng lại toàn hệ thống.",
    ])
    x, y, w, bot = d.panel(s, "BA PHÉP GOM")
    for src, dst in (("giờ", "buổi"), ("ngày", "mùa"), ("16 hướng gió", "bắc / nam")):
        d.chip(s, src, x, y, Inches(1.85), h=Inches(0.44), fill=INK_SOFT, size=12)
        d.text(s, "→", x + Inches(1.92), y + Inches(0.05), Inches(0.3), Inches(0.34), 15,
               ORANGE, bold=True, font=SANS, align=PP_ALIGN.CENTER)
        d.chip(s, dst, x + Inches(2.28), y, w - Inches(2.28), h=Inches(0.44), fill=ORANGE,
               size=12)
        y += Inches(0.62)
    d.mini_card(s, x, y + Inches(0.08), w, Inches(0.96), "LƯU Ý",
                ["cột BỔ SUNG — grain gốc", "mức giờ vẫn giữ nguyên"])
    d.remember(s, "Rời rạc hoá tạo CỘT bổ sung (buổi, mùa, phần tư gió) — grain gốc mức giờ "
                  "vẫn được giữ nguyên.")

    s = d.light("PHẦN A · KIẾN TRÚC", "Kiến trúc tổng: hai luồng vào, một kho")
    d.prose(s, [
        "**Luồng 1 — batch (nạp lịch sử):** CSV UCI + Open-Meteo → ##MinIO data lake## → "
        "##PostgreSQL staging## → ##dbt## dựng kho fact + dimension và các mart → đặc trưng "
        "ML → dashboard ##FastAPI + React##.",
        "**Luồng 2 — tác nghiệp (mô phỏng vận hành):** ##SQLite## đóng vai nguồn giao dịch → "
        "CDC theo watermark (phần D) → góp vào cùng lake và staging. ##Airflow## điều phối "
        "cả hai luồng theo lịch, ##Docker## đóng gói hạ tầng để chạy lại được y hệt.",
    ], y=Inches(1.42), h=Inches(1.15), w=CW)
    d.chip(s, "LUỒNG 1 · BATCH", MARGIN, Inches(2.62), Inches(2.3), h=Inches(0.32), fill=INK,
           size=11)
    lane1 = [("Nguồn batch", "UCI CSV · Open-Meteo", MUTED),
             ("Data lake", "MinIO · raw + Parquet", TEAL),
             ("Staging", "PostgreSQL · đối soát", TEAL),
             ("Kho — dbt", "fact + 4 chiều", ORANGE),
             ("Mart & khối", "OLAP · báo cáo", ORANGE),
             ("ML & Dashboard", "FastAPI + React", INK)]
    for i, (name, sub, color) in enumerate(lane1):
        bx = MARGIN + Inches(i * 2.02)
        d.chip(s, name, bx, Inches(3.06), Inches(1.74), h=Inches(0.44), fill=color, size=12)
        d.text(s, sub, bx, Inches(3.56), Inches(1.74), Inches(0.55), 11, MUTED, font=SANS,
               align=PP_ALIGN.CENTER)
        if i < len(lane1) - 1:
            d.text(s, "→", bx + Inches(1.74), Inches(3.08), Inches(0.28), Inches(0.4), 15,
                   MUTED, bold=True, font=SANS, align=PP_ALIGN.CENTER)
    d.chip(s, "LUỒNG 2 · TÁC NGHIỆP", MARGIN, Inches(4.32), Inches(2.6), h=Inches(0.32),
           fill=TEAL, size=11)
    lane2 = [("SQLite — tác nghiệp", MUTED), ("CDC theo watermark", INK),
             ("lake → staging (upsert)", TEAL)]
    for i, (name, color) in enumerate(lane2):
        bx = MARGIN + Inches(i * 2.5)
        d.chip(s, name, bx, Inches(4.76), Inches(2.24), h=Inches(0.44), fill=color, size=12)
        if i < len(lane2) - 1:
            d.text(s, "→", bx + Inches(2.24), Inches(4.78), Inches(0.26), Inches(0.4), 15,
                   MUTED, bold=True, font=SANS, align=PP_ALIGN.CENTER)
    d.text(s, "Airflow điều phối cả hai luồng · Docker đóng gói hạ tầng · mọi bước idempotent.",
           MARGIN, Inches(5.5), Inches(12.0), Inches(0.4), 13.5, MUTED, italic=True,
           font=SERIF)
    d.remember(s, "Mọi dữ liệu đều đi qua lake → staging → kho → mart; không có đường tắt "
                  "thẳng vào kho.")

    s = d.light("PHẦN A · KIẾN TRÚC", "Mỗi lớp một việc — không phải tất cả đều là database")
    d.prose(s, [
        "Câu hỏi bẫy kinh điển: “4 công cụ kia đều là database à?”. Không — mỗi công cụ một "
        "loại: ##MinIO## là **object storage** (giữ file thô nguyên vẹn, không phải RDBMS); "
        "##PostgreSQL## mới là RDBMS, giữ staging lẫn kho; ##dbt## là khung viết SQL biến "
        "đổi dữ liệu **sau khi nạp** (triết lý ELT); ##Airflow## là bộ điều phối — chạy việc "
        "gì, theo lịch nào.",
        "Bốn “vùng” cũng khác nhau: **lake** giữ bản gốc bất biến (khi phát hiện lỗi, quay "
        "về gốc nạp lại — kho dựng lại từ lake bất cứ lúc nào); **staging** tiếp nhận có "
        "cấu trúc để đối soát nguồn trước khi tin; **kho (DWH)** tổ chức fact + dimension "
        "cho phân tích; **mart** là tập con phục vụ đúng một nhu cầu.",
    ])
    x, y, w, bot = d.panel(s, "X LÀ GÌ?")
    y = d.kv(s, x, y, w, "MinIO", "object storage — kho file thô, ##không phải RDBMS##")
    y = d.kv(s, x, y, w, "PostgreSQL", "RDBMS — giữ staging và kho (fact + chiều)")
    y = d.kv(s, x, y, w, "dbt", "khung SQL biến đổi SAU khi nạp (ELT)")
    y = d.kv(s, x, y, w, "Airflow", "điều phối — việc gì chạy theo lịch nào")
    y = d.kv(s, x, y, w, "FastAPI + React", "backend + frontend dashboard")
    d.mini_card(s, x, y + Inches(0.1), w, Inches(1.34), "BỐN VÙNG",
                ["lake = giữ gốc · staging = đối soát", "kho = phân tích · mart = phục vụ "
                 "một nhu cầu"], h=1.34)
    d.remember(s, "Lake = giữ gốc · staging = tiếp nhận & đối soát · kho = phân tích "
                  "(fact/dim) · mart = phục vụ một nhu cầu.")

    s = d.light("PHẦN A · KIẾN TRÚC", "Tái lập bằng một lệnh — vì sao đó là lời cam kết")
    d.prose(s, [
        "Cả hệ thống chạy lại từ đầu bằng đúng một lệnh: ##python scripts/run_all.py## — "
        "##14 bước##, từ trạng thái trống tới dashboard mất ##11,2 phút##. Mọi bước "
        "**idempotent** (chạy lặp lại không đổi kết quả): nạp theo khóa grain + có sổ đăng "
        "ký từng batch (chi tiết phần D), nên chạy lại không nhân đôi dữ liệu.",
        "Ý nghĩa lớn hơn mặt kỹ thuật: mọi kết quả trong bài — F1 của mô hình, 24.759 ô "
        "khối kiểm chứng, các biểu đồ — **ai cũng tạo lại được như hệt**. Nó biến các con "
        "số từ “tôi chạy được đấy” thành **bằng chứng kiểm chứng được**; và cũng là công cụ "
        "ôn tốt nhất: muốn biết số nào đến từ đâu, mở script tương ứng đọc.",
    ])
    x, y, w, bot = d.panel(s, "MỘT LỆNH — TOÀN HỆ THỐNG")
    d.chip(s, "python scripts/run_all.py", x, y, w, h=Inches(0.5), fill=INK, size=13,
           font=MONO)
    y += Inches(0.72)
    y = d.stat(s, x, y, w, "14 bước", "thu thập → làm sạch → kho → ML → dashboard")
    d.hline(s, x, y, w)
    y = d.stat(s, x, y + Inches(0.12), w, "11,2 phút", "chạy lại từ trạng thái trống")
    d.hline(s, x, y, w)
    d.chip(s, "idempotent ✓ chạy lại không nhân đôi", x, y + Inches(0.14), w,
           h=Inches(0.42), fill=TEAL, size=12)
    d.remember(s, "run_all.py: 14 bước · 11,2 phút từ trống · idempotent — mọi con số trong "
                  "bài đều tái lập được.")

    s = d.light("PHẦN A · CHỐT", "Nếu chỉ nhớ 5 điều về phần A")
    d.numbered_rows(s, [
        ("1", "Câu hỏi nghiệp vụ:", "giờ tới trạm này có ≥ 115 µg/m³ không — cảnh báo trước "
         "khi xảy ra, thay cho cách báo trễ 1 giờ."),
        ("2", "Dữ liệu:", "UCI #501 — ##420.768## dòng trạm×giờ, 12 trạm, 03/2013–02/2017; "
         "Open-Meteo lấp thiếu thời tiết, UCI ưu tiên, ghi nguồn từng ô."),
        ("3", "Làm sạch:", "chuỗi thiếu dài 343–1.517 giờ giữ NULL + cờ; 60 dòng trần giữ + "
         "cờ SUSPECT_CAP; median chỉ tính trên train."),
        ("4", "Kiến trúc:", "hai luồng (batch + CDC) → MinIO lake → staging → dbt dựng kho "
         "→ mart → ML → dashboard; Airflow điều phối, Docker đóng gói."),
        ("5", "Tái lập:", "##run_all.py## — 14 bước, ##11,2 phút## từ trống, idempotent."),
    ], y=1.66, gap=0.9)
    d.remember(s, "Phần A trả lời: vì sao làm · dữ liệu gì · sạch thế nào · hệ thống mấy lớp "
                  "· chạy lại thế nào.")

    # ---- PHẦN B -------------------------------------------------------------------
    d.dark("PHẦN B · SLIDE 16–28", "Kho dữ liệu & OLAP",
           ["Grain — quyết định quan trọng nhất khi thiết kế kho",
            "Sao 4 chiều: vì sao gió là chiều mà thời tiết thì không",
            "180 tổ hợp mức vs 32 grouping sets — hai số hay trộn",
            "Kiểm chứng 100% · tổng hợp đúng (SUM + COUNT)",
            "Năm phép OLAP · ROLAP vs MOLAP · khóa & SCD · insight gió nam"], big="B")

    s = d.light("PHẦN B · KHO", "Grain — quyết định quan trọng nhất khi thiết kế kho")
    d.prose(s, [
        "**Grain** là câu trả lời cho “một dòng của bảng fact đại diện cho cái gì”. Ở dự án "
        "này: ##một dòng fact = một trạm tại một giờ##. Grain phải phát biểu đầu tiên: mọi "
        "độ đo, mọi khóa chiều sau này đều phải khớp grain; đổi grain về sau gần như thiết "
        "kế lại kho.",
        "Vì sao mức giờ chứ không mức ngày cho “nhẹ”: **(1)** giờ cuộn lên ngày/tháng/năm "
        "được, ngày về giờ thì **không bao giờ** — mất thông tin vĩnh viễn; **(2)** bài "
        "toán cảnh báo cần đúng mức giờ; **(3)** mức ngày mất biến động trong ngày (đỉnh "
        "buổi tối) — chính thứ mô hình dùng làm đặc trưng. Nguyên tắc chung: **lưu ở độ mịn "
        "gốc nhất mà nghiệp vụ cần**, tổng hợp để sau.",
    ])
    x, y, w, bot = d.panel(s, "GRAIN CỦA DỰ ÁN")
    d.card(s, x, y, w, Inches(0.95), SOFT_ORANGE)
    d.text(s, "1 dòng fact = 1 trạm × 1 giờ", x + Inches(0.15), y + Inches(0.08),
           w - Inches(0.3), Inches(0.8), 19, ORANGE, bold=True, font=SERIF,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.15)
    y = d.kv(s, x, y, w, "Cuộn một chiều", "giờ → ngày → năm được; ngược lại không")
    y = d.kv(s, x, y, w, "Bài toán cần", "cảnh báo t+1h cần đúng mức giờ")
    y = d.kv(s, x, y, w, "Mất gì nếu lưu ngày", "biến động buổi tối — đặc trưng của ML")
    d.remember(s, "Fact = trạm × giờ. Cuộn giờ→ngày được, ngày→giờ không — nên giữ mức mịn "
                  "gốc.")

    s = d.light("PHẦN B · KHO", "Lược đồ hình sao: fact ở giữa, 4 chiều xung quanh")
    d.prose(s, [
        "Bảng trung tâm là ##fact_air_quality_hourly## — mỗi dòng một lần đo, chứa các **độ "
        "đo** (nồng độ 6 chất, 5 biến khí tượng, cờ chất lượng) và khóa nối 4 bảng **chiều** "
        "(dimension): ##dim_date##, ##dim_time##, ##dim_station##, ##dim_wind##.",
        "Phân công rất rõ: **fact lưu số để tổng hợp** (cộng, trung bình); **dimension lưu "
        "ngữ cảnh để lọc và nhóm** (mùa nào? quận nào? gió hướng gì?). Mọi truy vấn phân "
        "tích đều là “chọn subset của fact theo điều kiện chiều, rồi tổng hợp độ đo”. Gọi "
        "là **hình sao** vì fact ở giữa, các chiều tỏa quanh — mỗi chiều nối fact bằng một "
        "liên kết. (Snowflake chuẩn hoá chiều thành nhiều bảng con — ít dư thừa nhưng join "
        "dài hơn, với 4 chiều nhỏ không đáng.)",
    ])
    x, y, w, bot = d.panel(s, "SƠ ĐỒ HÌNH SAO")
    d.star_diagram(s, x, y + Inches(0.05), w)
    d.remember(s, "Fact = sự kiện đo (độ đo + khóa nối) · dimension = ngữ cảnh để lọc/nhóm — "
                  "mỗi chiều một bước tới fact.")

    s = d.light("PHẦN B · KHO", "Vì sao hướng gió là chiều, còn thời tiết thì không")
    d.prose(s, [
        "Một cột nên làm **chiều** khi nó là ngữ cảnh lọc/nhóm với ít giá trị lặp lại; nên "
        "làm **độ đo** khi nó là thuộc tính số của chính sự kiện. **Hướng gió làm chiều** "
        "vì nó trả lời “**vì sao ô nhiễm đến**” — gió là điều kiện vận tải khí, nhóm theo "
        "hướng gió là cách soi nguồn khí (slide 27).",
        "**Thời tiết không làm chiều riêng**: nhiệt độ, mưa… được đo **cùng lúc, cùng grain "
        "trạm×giờ với chính sự kiện** — tức là thuộc tính số của sự kiện, để làm độ đo trên "
        "fact. Phản chứng dễ thấy: nếu nhiệt độ thành chiều, mỗi giá trị đo một dòng chiều "
        "— bảng chiều nổ ra hàng nghìn dòng vô nghĩa phân tích. Nếu sau này cần phân tích "
        "“hồ sơ thời tiết” riêng → dựng **fact riêng**, không nhét thêm chiều.",
    ])
    x, y, w, bot = d.panel(s, "PHÂN CÔNG NÀY")
    d.mini_card(s, x, y, w, Inches(1.5), "HƯỚNG GIÓ → CHIỀU",
                ["ngữ cảnh lọc/nhóm, ít giá trị", "điều kiện vận tải khí —", "soi nguồn "
                 "khí theo hướng gió"])
    d.mini_card(s, x, y + Inches(1.64), w, Inches(1.5), "THỜI TIẾT → ĐỘ ĐO",
                ["đo CÙNG sự kiện, cùng grain", "thuộc tính số của lần đo —", "nếu thành "
                 "chiều: nổ bảng chiều"], fill=SOFT_ORANGE, tag_color=ORANGE)
    d.chip(s, "cần phân tích thời tiết riêng → fact riêng", x, y + Inches(3.3), w,
           h=Inches(0.42), fill=INK_SOFT, color=AMBER, size=11)
    d.remember(s, "Chiều = ngữ cảnh lọc/nhóm (ít giá trị) · độ đo = thuộc tính số của sự "
                  "kiện — thời tiết đo cùng sự kiện nên là độ đo.")

    s = d.light("PHẦN B · KHO", "Phân cấp chiều và con số 180 tổ hợp mức")
    d.prose(s, [
        "Mỗi chiều có **phân cấp** (hierarchy) để cuộn lên — khoan xuống: ngày→tháng→quý→năm; "
        "giờ→buổi; trạm→quận; hướng→phần tư→bắc/nam. Đếm mức mỗi chiều **tính cả mức ALL** "
        "(tổng thể, không lọc): dim_date có 5 mức, dim_time 3, dim_station 3, dim_wind 4 — "
        "nhân lại ##5·3·3·4 = 180## tổ hợp mức.",
        "Đọc 180 đúng nghĩa: đó là **số “góc nhìn” khác nhau** kho có thể trả lời — không "
        "phải số bảng, cũng không phải số gì đã tính trước. ##Không cần — và không nên — "
        "tính trước hết 180 tổ hợp##: chỉ tính trước góc nhìn dashboard thật dùng (slide "
        "21), còn lại tính lúc truy vấn từ fact.",
    ])
    x, y, w, bot = d.panel(s, "ĐẾM MỨC — CÓ ALL")
    y = d.kv(s, x, y, w, "dim_date", "ngày · tháng · quý · năm + ALL = ##5 mức##")
    y = d.kv(s, x, y, w, "dim_time", "giờ · buổi + ALL = ##3 mức##")
    y = d.kv(s, x, y, w, "dim_station", "trạm · quận + ALL = ##3 mức##")
    y = d.kv(s, x, y, w, "dim_wind", "hướng · phần tư · bắc-nam + ALL = ##4 mức##")
    d.card(s, x, y + Inches(0.1), w, Inches(0.9), SOFT_ORANGE)
    d.text(s, "5·3·3·4 = 180", x + Inches(0.15), y + Inches(0.18), w - Inches(0.3),
           Inches(0.6), 26, ORANGE, bold=True, font=SERIF, align=PP_ALIGN.CENTER)
    d.remember(s, "180 = tổ hợp MỨC của 4 phân cấp (5·3·3·4, có ALL) — không phải số đã "
                  "tính trước; chỉ tính trước góc nhìn thật dùng.")

    s = d.light("PHẦN B · KHO", "Khóa chiều: khi nào khóa tự nhiên ổn, SCD là gì")
    d.prose(s, [
        "Dự án đang dùng **khóa tự nhiên** (natural key) — giá trị có sẵn từ nghiệp vụ: "
        "##station_key = station_id##, ##date_key = YYYYMMDD##, ##time_key = 0–23##, "
        "##wind_key## với 0 = “không rõ”. Ổn vì các chiều **tĩnh**, có **một nguồn duy "
        "nhất**, không có thành viên đổi thuộc tính.",
        "Kịch bản buộc phải khác: một trạm **đổi quận** sau khi dữ liệu đã nạp. Nếu chỉ "
        "UPDATE dòng chiều, mọi fact cũ tự động “nằm” ở quận mới — **lịch sử bị viết lại**. "
        "Giải pháp chuẩn: ##surrogate key## (khóa vô nghĩa do hệ thống sinh) + **SCD type "
        "2**: khi thuộc tính đổi, **thêm bản ghi mới** có valid_from / valid_to; fact cũ nối "
        "bản mô tả đúng thời điểm của nó — lịch sử nguyên vẹn.",
    ])
    x, y, w, bot = d.panel(s, "KHÓA HIỆN TẠI → SCD2")
    y = d.kv(s, x, y, w, "station_key", "= station_id (tự nhiên)")
    y = d.kv(s, x, y, w, "date_key", "= YYYYMMDD · time_key = 0–23")
    y = d.kv(s, x, y, w, "Vì sao ổn", "chiều tĩnh · một nguồn · không đổi thuộc tính")
    y += Inches(0.1)
    d.mini_card(s, x, y, w, Inches(1.62), "NẾU TRẠM ĐỔI QUẬN",
                ["UPDATE chiều = viết lại lịch sử ✗", "→ surrogate key + ##SCD type 2##:",
                 "bản ghi MỚI + valid_from / valid_to", "— fact cũ giữ mô tả đúng thời điểm"],
                fill=SOFT_ORANGE, tag_color=ORANGE, size=11.5, h=1.62)
    d.remember(s, "Chiều tĩnh → khóa tự nhiên đủ. Chiều đổi thuộc tính → surrogate key + "
                  "SCD type 2 (bản ghi mới + valid_from/to).")

    s = d.light("PHẦN B · KHO", "Ba khối tính trước (mart) — chọn theo truy vấn thật")
    d.prose(s, [
        "Ba mart chính: ##mart_station_daily## (trạm × ngày) ##17.532 dòng## — fact cuộn "
        "lại nhỏ hơn ##24 lần##, vẽ xu hướng nhanh; ##mart_station_monthly## ##576 dòng## — "
        "báo cáo tháng/mùa; ##mart_olap_cube## ##6.075 dòng## — chứa 32 khối con của CUBE, "
        "trang OLAP đọc chỉ vài ms.",
        "Tiêu chí chọn: **đúng các truy vấn dashboard thật phát ra** — không tính trước "
        "theo cảm tính. Truy vấn đủ thuộc tính nằm trong cube → đọc mart ##2–6 ms##; thiếu "
        "một thuộc tính là quét fact ##130–400 ms##. Đây là bài đánh đổi chung của mọi kho: "
        "tính trước = nhanh nhưng phải giữ đồng bộ; tính lúc chạy = đúng mọi truy vấn "
        "nhưng chậm. Dự án chọn: tính trước phần hay dùng, phần còn lại quét fact vẫn "
        "chịu được.",
    ])
    x, y, w, bot = d.panel(s, "BA MART + TỐC ĐỘ")
    y = d.stat(s, x, y, w, "17.532", "mart_station_daily — trạm × ngày (nhỏ hơn fact 24×)",
               vsize=22)
    y = d.stat(s, x, y + Inches(0.06), w, "576", "mart_station_monthly — báo cáo tháng/mùa",
               vsize=22)
    y = d.stat(s, x, y + Inches(0.06), w, "6.075", "mart_olap_cube — 32 khối cho trang OLAP",
               vsize=22)
    y += Inches(0.08)
    d.mini_card(s, x, y, w, Inches(0.86), "TỐC ĐỘ ĐỌC",
                ["đủ thuộc tính → ##2–6 ms##", "thiếu → quét fact ##130–400 ms##"],
                size=11.5)
    d.remember(s, "3 mart: 17.532 (ngày) · 576 (tháng) · 6.075 (cube) — chọn theo truy vấn "
                  "thật; đủ thuộc tính đọc 2–6 ms, thiếu thì quét fact.")

    s = d.light("PHẦN B · KHO", "CUBE 5 thuộc tính = 2⁵ = 32 grouping sets")
    d.prose(s, [
        "Trang OLAP tổng hợp theo ##CUBE(năm, mùa, quận, buổi, phần tư gió)##. Với CUBE, "
        "mỗi thuộc tính có 2 lựa chọn — **giữ nguyên giá trị** hoặc **cuộn lên ALL** — và "
        "CUBE sinh **mọi tổ hợp**: 2×2×2×2×2 = ##32 grouping sets##. Ví dụ một khối: "
        "(2015, Đông, ALL, ALL, Nam) — tổng cả thành phố mùa đông 2015, chỉ giờ gió nam.",
        "Hai bẫy hay hỏi: **(1)** 5 thuộc tính trong CUBE ##không bằng## 4 bảng chiều — "
        "năm và mùa cùng nằm trong dim_date; **(2)** ##32 ≠ 180## — 32 là grouping sets của "
        "CUBE này; 180 là tổ hợp mức phân cấp của cả 4 chiều (kể cả thuộc tính ngoài CUBE "
        "như trạm cụ thể, ngày cụ thể). GROUPING SETS là cú pháp SQL đằng sau: mỗi grouping "
        "set = một cách GROUP BY một tập cột.",
    ])
    x, y, w, bot = d.panel(s, "2⁵ = 32")
    d.card(s, x, y, w, Inches(1.05), SOFT_TEAL)
    d.text(s, "2⁵ = 32", x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), Inches(0.85),
           34, TEAL, bold=True, font=SERIF, align=PP_ALIGN.CENTER,
           anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.25)
    y = d.kv(s, x, y, w, "Mỗi thuộc tính", "giữ nguyên giá trị hoặc cuộn lên ALL")
    y = d.kv(s, x, y, w, "Ví dụ khối", "(2015, Đông, ALL, ALL, Nam)")
    d.mini_card(s, x, y + Inches(0.08), w, Inches(1.3), "HAI BẪY",
                ["5 thuộc tính ##≠## 4 bảng chiều —", "năm + mùa cùng dim_date",
                 "##32 ≠ 180## (mức phân cấp)"], fill=SOFT_ORANGE, tag_color=ORANGE,
                size=11.5, h=1.3)
    d.remember(s, "2⁵ = 32 grouping sets (giữ/ALL cho từng thuộc tính) · thuộc tính CUBE ≠ "
                  "số bảng chiều · 32 ≠ 180.")

    s = d.light("PHẦN B · KHO", "Kiểm chứng kho: tính lại từng ô — 100% khớp")
    d.prose(s, [
        "Phần tự hào nhất của kho: mọi ô tính trước đều được **kiểm chứng độc lập**. Script "
        "chạy lại GROUPING SETS từng khối **trực tiếp từ bảng fact**, rồi so **từng ô một** "
        "với mart đã tính trước: ##24.759 / 24.759 ô khớp = 100%##.",
        "Lớp thứ hai là ##130 dbt test## chạy tự động mỗi lần dựng kho: khóa duy nhất, "
        "not-null, khóa ngoại fact↔chiều, và test **phân hoạch** — mỗi dòng fact đóng góp "
        "đúng MỘT lần vào mỗi grouping set (nếu một dòng rơi vào 2 khối, tổng các khối sẽ "
        "vượt tổng fact → bị bắt). Nói đủ cả hạn chế: khối tính trước chỉ phục vụ truy "
        "v vấn có đủ thuộc tính trong cube; kiểm chứng không phải “mọi truy vấn đều "
        "nhanh” mà là “**mọi ô đã tính đều đúng, sai thì máy móc phát hiện**”.",
    ])
    x, y, w, bot = d.panel(s, "HAI LỚP KIỂM CHỨNG")
    y = d.stat(s, x, y, w, "100%", "24.759/24.759 ô khớp khi tính lại trực tiếp từ fact",
               vsize=40)
    d.hline(s, x, y + Inches(0.05), w)
    y = d.stat(s, x, y + Inches(0.17), w, "130 dbt test",
               "khóa duy nhất · not-null · khóa ngoại · phân hoạch fact", vsize=22)
    d.mini_card(s, x, y + Inches(0.1), w, Inches(1.0), "PHÂN HOẠCH",
                ["mỗi dòng fact đóng góp đúng 1 lần", "vào MỖI grouping set"],
                size=11.5)
    d.remember(s, "24.759/24.759 ô = 100% khớp khi tính lại từ fact · 130 dbt test — hai "
                  "lớp kiểm chứng độc lập.")

    s = d.light("PHẦN B · KHO", "Tổng hợp đúng: lưu SUM + COUNT, không lưu trung bình")
    d.prose(s, [
        "Trung bình là **độ đo “không cộng”** — không thể cộng hai trung bình thành một "
        "trung bình đúng. Vì vậy mart và cube chỉ lưu ##SUM## và ##COUNT## của từng nhóm, "
        "chia ra khi đọc. Cùng nguyên tắc, ##không được SUM cả bảng cube## để đếm tổng số "
        "giờ: mỗi dòng fact nằm trong NHIỀU grouping set — cộng cả bảng là **đếm trùng** "
        "một quan trắc nhiều lần.",
        "Lưu SUM + COUNT còn có lợi nữa: khi **mức thiếu khác nhau giữa các nhóm** (nhóm "
        "này thiếu nhiều giờ hơn nhóm kia), COUNT phản ánh đúng số giờ đo hợp lệ thật — "
        "gộp lên cấp cao vẫn đúng trọng số.",
    ])
    x, y, w, bot = d.panel(s, "VÍ DỤ 30 GIỜ ĐO")
    d.text(s, "24 giờ × 100 + 6 giờ × 200 µg/m³", x, y, w, Inches(0.3), 12, MUTED,
           font=SANS, align=PP_ALIGN.CENTER)
    y += Inches(0.4)
    d.mini_card(s, x, y, w, Inches(1.14), "ĐÚNG ✓",
                ["(2.400 + 1.200) / 30", "= ##3.600 / 30 = 120##"])
    y += Inches(1.28)
    d.mini_card(s, x, y, w, Inches(1.14), "SAI ✗",
                ["trung bình của trung bình", "=(100 + 200)/2 = ##150##"], fill=SOFT_ORANGE,
                tag_color=ORANGE)
    d.chip(s, "SUM cả bảng cube = đếm trùng", x, y + Inches(1.4), w, h=Inches(0.42),
           fill=INK_SOFT, color=AMBER, size=11.5)
    d.remember(s, "Khối lưu SUM + COUNT rồi chia lúc đọc · không SUM cả cube để đếm (đếm "
                  "trùng) · không lấy trung bình của trung bình.")

    s = d.light("PHẦN B · OLAP", "Năm phép OLAP — từ khóa gọi tên đúng thao tác")
    d.prose(s, [
        "OLAP (xử lý phân tích trực tuyến) có 5 thao tác kinh điển, và demo /olap gọi tên "
        "đúng từng thao tác khi bấm. Kịch bản demo nối đủ 5 phép: mở trang đang xem **Năm × "
        "Mùa** (đọc từ cube) → bấm 2015 **drill-down** quý → **roll-up** về năm → lọc "
        "Dongcheng (**slice**) → thêm Xicheng + 2016 (**dice**) → đổi hàng/cột "
        "(**pivot**).",
        "Nhớ bằng bản chất: slice/dice là **chọn tập con** (1 giá trị hay nhiều giá trị); "
        "drill/roll là **đổi mức mịn-thô** trên phân cấp; pivot chỉ **đổi góc trình bày**, "
        "dữ liệu không đổi.",
    ])
    x, y, w, bot = d.panel(s, "NĂM PHÉP")
    ops = [("Slice", "chọn MỘT giá trị — quận = Dongcheng"),
           ("Dice", "chọn NHIỀU giá trị / miền — + Xicheng, 2016"),
           ("Drill-down", "xuống mức mịn hơn — 2015 → quý"),
           ("Roll-up", "lên mức thô hơn — quý → năm, hoặc về ALL"),
           ("Pivot", "xoay trục — đổi hàng ↔ cột")]
    for name, desc in ops:
        d.chip(s, name, x, y, Inches(1.45), h=Inches(0.4), fill=TEAL, size=12)
        d.text(s, desc, x + Inches(1.6), y + Inches(0.04), w - Inches(1.6), Inches(0.6),
               11.5, INK, font=SANS, spacing=1)
        y += Inches(0.56)
    d.remember(s, "Slice 1 giá trị · Dice nhiều giá trị/miền · Drill-down xuống mịn · "
                  "Roll-up lên thô · Pivot xoay trục.")

    s = d.light("PHẦN B · OLAP", "ROLAP vs MOLAP — dự án dùng gì và vì sao")
    d.prose(s, [
        "Hai kiến trúc lưu trữ OLAP: ##ROLAP## — dữ liệu nằm trong **cơ sở dữ liệu quan "
        "hệ**, tổng hợp bằng SQL + khối tính trước trong mart; ##MOLAP## — engine **ma "
        "trận đa chiều riêng**, tính trước mọi ô của khối. Dự án dùng **ROLAP** "
        "(PostgreSQL + SQL + mart_olap_cube).",
        "Vì sao: ##420 nghìn dòng fact và 5 thuộc tính CUBE là quy mô PostgreSQL xử lý "
        "trong vài ms## — xây engine MOLAP lúc này là phức tạp không mua thêm tốc độ. Câu "
        "“khi nào đổi”: nếu lên ~100 trạm × 10 năm (hàng chục triệu dòng) thì cân nhắc kho "
        "cột (columnar), engine MPP, hoặc Spark — vẫn giữ mô hình fact/dimension, chỉ đổi "
        "lớp lưu trữ bên dưới.",
    ])
    x, y, w, bot = d.panel(s, "DỰ ÁN DÙNG ROLAP")
    d.mini_card(s, x, y, w, Inches(1.5), "ROLAP ✓",
                ["quan hệ (PostgreSQL) + SQL", "+ mart tính trước — ##dự án dùng##", "đủ "
                 "nhanh: vài ms ở 420k dòng"])
    d.mini_card(s, x, y + Inches(1.64), w, Inches(1.3), "MOLAP — khi nào?",
                ["engine ma trận đa chiều riêng", "tính trước mọi ô — đáng khi ~100 trạm × "
                 "10 năm"], fill=SOFT_ORANGE, tag_color=ORANGE, h=1.3)
    d.chip(s, "quy mô lớn → columnar · MPP · Spark", x, y + Inches(3.1), w,
           h=Inches(0.42), fill=INK_SOFT, color=AMBER, size=11.5)
    d.remember(s, "Dự án = ROLAP (quan hệ + SQL + mart tính trước) — đủ nhanh ở 420k dòng; "
                  "quy mô lớn hẳn mới cần MOLAP/columnar/MPP.")

    s = d.light("PHẦN B · INSIGHT", "Đọc một insight đúng cách: gió nam mùa đông +60%")
    d.prose(s, [
        "Kết quả dice thật từ kho: cùng nhóm quận lõi đô thị, mùa đông, những giờ **gió "
        "bắc** có PM2.5 trung bình ##~85 µg/m³## còn giờ **gió nam** lên ##~133 µg/m³ "
        "(+60%)##; tỷ lệ giờ vượt ngưỡng tăng ##28% → 43%##. Đây là giá trị nghiệp vụ của "
        "việc đưa hướng gió thành chiều: nó chỉ ra **khi nào cần cảnh báo gắt hơn**.",
        "Diễn giải phải đúng mức: đây là ##tương quan theo hướng gió## — **gợi ý** nguồn "
        "khí phía nam (vùng công nghiệp Hà Bắc), nhưng dự án **không mô phỏng đường vận "
        "chuyển khí** nên không khẳng định nhân quả. Kho thấy **sự đi kèm**, mô hình tìm "
        "quan hệ để **dự đoán** — chứng minh nhân quả là việc của một thiết kế nghiên cứu "
        "khác.",
    ])
    x, y, w, bot = d.panel(s, "DICE THẬT — MÙA ĐÔNG")
    y = d.stat(s, x, y, w, "85 → 133", "µg/m³ — gió bắc → gió nam (+60%)", vsize=28)
    d.hline(s, x, y, w)
    y = d.stat(s, x, y + Inches(0.12), w, "28% → 43%", "tỷ lệ giờ vượt ngưỡng 115 — cùng "
               "điều kiện", vsize=28)
    d.mini_card(s, x, y + Inches(0.14), w, Inches(1.14), "DIỄN GIẢI",
                ["##tương quan## theo hướng gió", "≠ khẳng định nguồn phát thải"],
                size=11.5)
    d.remember(s, "Gió nam mùa đông: 85 → 133 µg/m³ (+60%), vượt ngưỡng 28% → 43% — khẳng "
                  "định ở mức TƯƠNG QUAN, không phải nhân quả.")

    s = d.light("PHẦN B · CHỐT", "Nếu chỉ nhớ 5 điều về phần B")
    d.numbered_rows(s, [
        ("1", "Grain:", "một dòng fact = trạm × giờ — lưu mức mịn gốc vì cuộn lên được, "
         "ngược lại không."),
        ("2", "Sao 4 chiều:", "date · time · station · wind; gió là chiều (vận tải khí), "
         "thời tiết là độ đo trên fact."),
        ("3", "Hai con số dễ trộn:", "##180## = tổ hợp mức phân cấp (5·3·3·4) ≠ ##32## = "
         "grouping sets của CUBE (2⁵)."),
        ("4", "Đúng & kiểm chứng:", "lưu SUM + COUNT; ##24.759/24.759 = 100%## + ##130## "
         "dbt test; ROLAP."),
        ("5", "Insight:", "gió nam mùa đông ##+60%## PM2.5 — tương quan theo hướng gió, "
         "không phải nhân quả."),
    ], y=1.66, gap=0.9)
    d.remember(s, "Kho = grain → sao 4 chiều → tính trước theo truy vấn thật → kiểm chứng "
                  "100% → đọc insight đúng mức.")

    # ---- PHẦN C -------------------------------------------------------------------
    d.dark("PHẦN C · SLIDE 30–42", "Khai phá dữ liệu",
           ["Nhãn 115 µg/m³ và 49 đặc trưng chỉ nhìn quá khứ",
            "Chia tập theo thời gian, chống rò rỉ — vì sao không k-fold",
            "Baseline: bẫy accuracy của majority, sức mạnh của persistence",
            "Kết quả HistGB 0,927 · hai ngưỡng 115 vs 0,50 · McNemar",
            "Tầm 1/3/6 giờ · thí nghiệm đợt 6 giờ · hồi quy & phân cụm"], big="C")

    s = d.light("PHẦN C · BÀI TOÁN", "Nhãn từ đâu và 49 đặc trưng “chỉ nhìn về quá khứ”")
    d.prose(s, [
        "Bài toán: với mỗi giờ t tại mỗi trạm, dự đoán ##y = 1 khi PM2.5(t+1h) ≥ 115 "
        "µg/m³##. Con số 115 là **ranh giới dải “ô nhiễm trung bình”** theo quy chuẩn AQI "
        "Trung Quốc cho PM2.5 — là ##nồng độ µg/m³##, **không phải điểm AQI = 115**; khóa "
        "từ phân tích EDA các đợt ô nhiễm (đợt tại một trạm = ≥ 3 giờ liên tiếp ≥ 115).",
        "Bên vào là ##49 đặc trưng##, nguyên tắc vàng: **mọi đặc trưng chỉ dùng thông tin "
        "BIẾT TẠI thời điểm t** — không một đặc trưng nào, kể cả gián tiếp qua nội suy hay "
        "chuẩn hoá, được nhìn thấy giá trị của t+1 trở đi. Bài toán giả lập “đứng ở giờ t "
        "dự đoán giờ t+1” — đặc trưng nhìn trộm tương lai thì điểm đẹp chỉ là ảo ảnh, "
        "thực chiến sẽ sụp.",
    ])
    x, y, w, bot = d.panel(s, "NHÃN & ĐẶC TRƯNG")
    d.card(s, x, y, w, Inches(1.0), SOFT_ORANGE)
    d.text(s, "y = 1 ⟺ PM2.5(t+1h) ≥ 115", x + Inches(0.12), y + Inches(0.08),
           w - Inches(0.24), Inches(0.85), 17, ORANGE, bold=True, font=SERIF,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.2)
    y = d.stat(s, x, y, w, "49 đặc trưng", "tất cả chỉ dùng thông tin biết tại thời điểm t",
               vsize=22)
    for i, g in enumerate(("PM2.5 + độ trễ 1–24h", "xu hướng · 5 chất khác · khí tượng",
                           "giờ · tháng · mùa · trạm · cụm trạm")):
        d.chip(s, g, x, y + Inches(0.06) + Inches(i * 0.52), w, h=Inches(0.42),
               fill=INK_SOFT, size=11.5)
    d.remember(s, "y = 1 ⟺ PM2.5(t+1h) ≥ 115 µg/m³ (nồng độ, không phải AQI) · 49 đặc "
                  "trưng chỉ dùng thông tin biết tại t.")

    s = d.light("PHẦN C · PHƯƠNG PHÁP", "Chia tập theo thời gian — vì sao không k-fold")
    d.prose(s, [
        "Với dữ liệu thời gian, **cách chia tập quyết định kết quả có tin được không**. "
        "Hai giờ kề nhau gần như trùng nhau (ô nhiễm không nhảy trạng thái đột ngột) — "
        "nếu chia random, giờ láng giềnh nằm cả hai phía train/test, mô hình gần như "
        "“được xem đề” → kết quả **lạc quan** vô nghĩa thực chiến; k-fold random vì thế "
        "không dùng.",
        "Chia đúng: ##train (03/2013–12/2015) → val (01–06/2016) → test (07–12/2016)##, "
        "thời gian đi một chiều. **Val** để chọn siêu tham số (độ sâu cây, ngưỡng xác "
        "suất); ##test dùng đúng MỘT lần## để báo cáo. Rò rỉ được rà chủ động: cửa sổ "
        "rolling và nhãn không tràn qua biên chia — có script kiểm tra; phát hiện thấy "
        "và refit (F1 0,926 so với 0,927 — nói thật, không khẳng định hoàn hảo).",
    ])
    x, y, w, bot = d.panel(s, "CHIA THEO THỜI GIAN")
    segs = [("TRAIN", "03/2013 – 12/2015", "276.380 dòng · 25,2% nhãn 1", TEAL),
            ("VAL", "01 – 06/2016", "47.338 dòng · 15,8%", MUTED),
            ("TEST", "07 – 12/2016", "48.781 dòng · 23,9%", ORANGE)]
    for name, period, info, color in segs:
        d.chip(s, name, x, y, Inches(1.0), h=Inches(0.42), fill=color, size=12)
        d.text(s, period, x + Inches(1.15), y + Inches(0.02), w - Inches(1.15),
               Inches(0.3), 12.5, INK, bold=True, font=SANS)
        d.text(s, info, x + Inches(1.15), y + Inches(0.28), w - Inches(1.15),
               Inches(0.3), 11, MUTED, font=SANS)
        y += Inches(0.62)
    y += Inches(0.06)
    d.mini_card(s, x, y, w, Inches(1.3), "NGUYÊN TẮC",
                ["chọn siêu tham số ##chỉ trên val##", "##test dùng đúng MỘT lần## — dùng "
                 "lại là tự học vào test"], h=1.3, size=11.5)
    d.remember(s, "Dữ liệu thời gian → chia THEO THỜI GIAN train→val→test; chọn siêu tham "
                  "số chỉ trên val; test dùng đúng một lần.")

    s = d.light("PHẦN C · PHƯƠNG PHÁP", "Ba tập và vì sao 372.499 dòng chứ không 420.768")
    d.prose(s, [
        "Số dòng ba tập: ##train 276.380## (25,2% nhãn 1) · ##val 47.338## (15,8%) · "
        "##test 48.781## (23,9%) — cộng lại 372.499, thiếu gần 50 nghìn so với 420.768. "
        "Đây là chỗ hay bị hỏi “số đâu mất?”",
        "Ba lý do mất dòng: **(1)** cửa sổ ML khóa ##03/2013–12/2016## — 01–02/2017 dành "
        "riêng cho đối soát CDC (phần D); **(2)** dòng thiếu ##nhãn PM2.5(t+1)## bị loại — "
        "không nhãn thì không học được; **(3)** dòng không đủ ##chuỗi 24 giờ trước## (độ "
        "trễ + trung bình trượt) bị loại — và mỗi giờ thiếu làm **gãy chuỗi**, kéo theo "
        "loại cả những giờ ngay sau nó. Tỷ lệ nhãn 1 khác nhau giữa ba tập (nửa đầu 2016 "
        "ô nhiễm ít hơn nửa cuối) — thêm lý do không dùng k-fold random.",
    ])
    x, y, w, bot = d.panel(s, "420.768 → 372.499")
    y = d.stat(s, x, y, w, "−48.269 dòng", "ba lý do dưới đây", vsize=24)
    d.hline(s, x, y + Inches(0.04), w)
    y = d.kv(s, x, y + Inches(0.16), w, "1 · Ngoài cửa sổ",
             "01–02/2017 để đối soát CDC")
    y = d.kv(s, x, y, w, "2 · Thiếu nhãn", "không có PM2.5(t+1)")
    y = d.kv(s, x, y, w, "3 · Gãy chuỗi lag",
             "thiếu 1 giờ → mất cả những giờ ngay sau nó (chuỗi 24h)")
    d.mini_card(s, x, y + Inches(0.1), w, Inches(0.96), "CỘNG LẠI",
                ["##276.380 + 47.338 + 48.781 = 372.499##"], size=12)
    d.remember(s, "372.499 = 420.768 − ngoài cửa sổ ML (01–02/2017 cho CDC) − thiếu nhãn "
                  "t+1 − gãy chuỗi lag 24 giờ.")

    s = d.light("PHẦN C · BASELINE", "Baseline 1 — majority và cái bẫy accuracy")
    d.prose(s, [
        "Baseline (mốc so sánh tối thiểu) thứ nhất: **majority** — luôn đoán lớp đa số "
        "“không vượt”. Kết quả: ##accuracy 0,761 nhưng F1 = 0,000## — không phát được "
        "cảnh báo nào thật sự. Vì accuracy chỉ đếm đúng/sai từng dòng: lớp “không vượt” "
        "~76% nên đoán mù cũng đúng 76%, trong khi **~24% giờ nguy hiểm bị bỏ qua hết**.",
        "Ba chỉ số đọc thay accuracy: ##precision## = trong các cảnh báo phát ra, bao "
        "nhiêu đúng (báo nhầm tốn uy tín); ##recall## = trong các giờ thật sự vượt ngưỡng, "
        "bắt được bao nhiêu (bỏ sót là người dân hít khói độc); ##F1## = trung bình điều "
        "hoà hai cái — chỉ số chính của bài này. Bài học: **lớp mất cân bằng làm accuracy "
        "mất ý nghĩa** — bài toán cảnh báo luôn phải báo precision/recall/F1.",
    ])
    x, y, w, bot = d.panel(s, "BẪY ACCURACY")
    d.mini_card(s, x, y, w, Inches(1.3), "MAJORITY",
                ["accuracy ##0,761## — nghe cao", "nhưng F1 = ##0,000##: không bắt được "
                 "cảnh báo nào"], fill=SOFT_ORANGE, tag_color=ORANGE, h=1.3)
    y += Inches(1.44)
    y = d.kv(s, x, y, w, "precision", "trong các cảnh báo phát ra, bao nhiêu đúng")
    y = d.kv(s, x, y, w, "recall", "giờ thật sự vượt ngưỡng, bắt được bao nhiêu")
    y = d.kv(s, x, y, w, "F1", "điều hoà precision và recall — chỉ số chính")
    d.chip(s, "lớp thiểu số ~24% → accuracy bẫy", x, y + Inches(0.08), w,
           h=Inches(0.42), fill=INK_SOFT, color=AMBER, size=11.5)
    d.remember(s, "Majority: accuracy 0,761 nhưng F1 = 0,000 — với lớp thiểu số ~24%, "
                  "accuracy là chỉ số bẫy; nhìn F1.")

    s = d.light("PHẦN C · BASELINE", "Baseline 2 — persistence mạnh đến mức nào")
    d.prose(s, [
        "Baseline thứ hai: ##persistence## — “giờ tới giữ nguyên hiện tại”: PM2.5 lúc này "
        "≥ 115 thì báo vượt. F1 đạt ##0,914## — rất cao, vì hai giờ liên tiếp tương quan "
        "chặt (ô nhiễm tích tụ và tan dần từ từ, ít đổi trạng thái đột ngột).",
        "Hệ quả phương pháp luận: mô hình học máy phải chứng minh ##giá trị THÊM## so với "
        "baseline mạnh này — chứ không phải so với majority cho đẹp. So với baseline yếu "
        "thì phần thắng bị thổi phồng; ở đây “sàn” đã 0,914. Câu hỏi đúng để đánh giá mô "
        "hình không phải “F1 bao nhiêu?” mà “**làm được gì mà persistence không làm "
        "được?**” — hai slide nữa trả lời đúng câu đó.",
    ])
    x, y, w, bot = d.panel(s, "SÀN PHẢI VƯỢT")
    y = d.stat(s, x, y, w, "F1 0,914", "persistence — “giờ tới giữ nguyên hiện tại”",
               vsize=36)
    d.mini_card(s, x, y + Inches(0.12), w, Inches(1.34), "VÌ SAO MẠNH",
                ["hai giờ liên tiếp tương quan chặt —", "ô nhiễm ít đổi trạng thái đột "
                 "ngột", "→ mô hình phải chứng minh giá trị ##THÊM##"], h=1.34, size=11.5)
    d.chip(s, "majority F1 0,000 — không phải mốc so", x, y + Inches(3.02), w,
           h=Inches(0.42), fill=INK_SOFT, color=AMBER, size=11.5)
    d.remember(s, "Persistence F1 0,914 là “sàn” phải vượt có ý nghĩa — luôn so với "
                  "baseline mạnh nhất, không so với majority.")

    s = d.light("PHẦN C · KẾT QUẢ", "Kết quả 4 mô hình — đọc bảng đúng cách")
    d.prose(s, [
        "Xếp hạng trên tập test: ##majority F1 0,000## · ##persistence 0,914## · ##cây "
        "quyết định (entropy, sâu 6) 0,921## · ##HistGradientBoosting 0,927## (precision "
        "0,926 · recall 0,928 · ROC-AUC 0,993 · ngưỡng 0,50) — **mô hình khóa**.",
        "So persistence theo số giờ cụ thể trên 6 tháng test: giờ ô nhiễm bị bỏ sót "
        "(FN) ##1.000 → 838##, cảnh báo nhầm (FP) ##999 → 865## — cả hai loại lỗi "
        "đều giảm, không đổi cái này lấy cái kia. Vì sao cây & boosting: dễ giải thích + "
        "mạnh trên bảng đặc trưng đã dựng; HistGB = **ensemble** — từng cây yếu nối tiếp "
        "nhau sửa lỗi cây trước, bin histogram cho nhanh.",
    ])
    x, y, w, bot = d.panel(s, "F1 TRÊN TẬP TEST")
    models = [("Majority", "0,000", False), ("Persistence", "0,914", False),
              ("Cây quyết định · sâu 6", "0,921", False),
              ("HistGradientBoosting", "0,927", True)]
    for name, f1, hot in models:
        fill = SOFT_ORANGE if hot else CARD
        d.card(s, x, y, w, Inches(0.62), fill)
        d.text(s, name, x + Inches(0.16), y + Inches(0.06), Inches(2.7), Inches(0.5),
               12, INK, bold=hot, font=SANS, anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, f1, x + w - Inches(1.2), y + Inches(0.06), Inches(1.05), Inches(0.5),
               20 if hot else 17, ORANGE if hot else INK, bold=True, font=SERIF,
               align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.74)
    y += Inches(0.04)
    d.text(s, "FN 1.000 → 838 · FP 999 → 865 — giảm cả hai lỗi so persistence", x, y,
           w, Inches(0.6), 12, TEAL, bold=True, font=SANS, spacing=2)
    d.remember(s, "HistGB 0,927 khóa (AUC 0,993, cutoff 0,50) · so persistence: FN "
                  "1.000→838 và FP 999→865 — giảm cả hai loại lỗi.")

    s = d.light("PHẦN C · KẾT QUẢ", "Hai ngưỡng khác nhau: 115 và 0,50")
    d.prose(s, [
        "Hay bị hỏi lẫn: hệ thống có **hai ngưỡng hoàn toàn khác vai trò**. ##115 µg/m³## "
        "tạo **nhãn** — dịch từ nồng độ sang câu hỏi phân loại, thuộc phần định nghĩa bài "
        "toán. ##0,50## là **ngưỡng xác suất** ra quyết định — mô hình xuất xác suất "
        "0…1, so với ngưỡng này mới thành cảnh báo.",
        "Ngưỡng xác suất **có thể dịch** (cây quyết định dùng 0,45) để đổi cán cân lỗi: "
        "cảnh báo sức khoẻ thường đặt **bỏ sót (FN) đắt hơn báo nhầm (FP)** — hạ ngưỡng "
        "thì recall tăng (bắt gần hết giờ nguy hiểm), precision giảm (nhiều cảnh báo sai "
        "hơn). Chọn ngưỡng nào là quyết định nghiệp vụ, không phải con số kỹ thuật cố "
        "định.",
    ])
    x, y, w, bot = d.panel(s, "HAI NGƯỠNG — HAI VAI")
    d.mini_card(s, x, y, w, Inches(1.4), "115 µg/m³ — NHÃN",
                ["dịch nồng độ → nhãn 0/1", "thuộc ##định nghĩa bài toán##, cố định"])
    y += Inches(1.54)
    d.mini_card(s, x, y, w, Inches(1.4), "0,50 — QUYẾT ĐỊNH",
                ["cắt xác suất của mô hình thành cảnh báo", "##dịch được## — đổi "
                 "recall/precision"], fill=SOFT_ORANGE, tag_color=ORANGE, h=1.4)
    d.chip(s, "hạ 0,50 → recall↑ · precision↓", x, y + Inches(1.56), w, h=Inches(0.42),
           fill=INK_SOFT, color=AMBER, size=11.5)
    d.remember(s, "115 µg/m³ tạo nhãn (định nghĩa bài toán) · 0,50 cắt xác suất ra cảnh "
                  "báo (dịch được để đổi recall/precision).")

    s = d.light("PHẦN C · KẾT QUẢ", "Giá trị thật: báo trước được lúc đợt ô nhiễm bắt đầu")
    d.prose(s, [
        "F1 hơn baseline ##+0,013## nghe rất nhỏ. Giá trị thật nằm ở chỗ persistence "
        "**không bao giờ làm được**: tách riêng ##1.000 giờ “đang sạch nhưng giờ sau vượt "
        "ngưỡng”## — chính là các thời điểm **đợt ô nhiễm bắt đầu**, lúc người vận hành "
        "cần cảnh báo nhất.",
        "Trong nhóm giờ đó, mô hình báo trước ##304 lần = 30,4%##; persistence ##0% theo "
        "đúng định nghĩa## — nó chỉ nhìn số đo hiện tại, mà giờ hiện tại vẫn sạch thì "
        "nhất định không báo. Đây không phải “hơi hơn” mà là **khả năng về bản chất "
        "khác**. Đọc số cho đúng: 30,4% là tỷ lệ bắt được giờ khởi đầu — không phải độ "
        "chính xác chung; ~70% đợt còn lại được bắt ở giờ thứ hai, thứ ba (khi "
        "persistence cũng bắt được).",
    ])
    x, y, w, bot = d.panel(s, "1.000 GIỜ KHỞI ĐẦU ĐỢT")
    y = d.stat(s, x, y, w, "304 / 1.000", "giờ “đang sạch, giờ sau vượt” được báo trước "
               "= 30,4%", vsize=30)
    d.hline(s, x, y, w)
    y = d.stat(s, x, y + Inches(0.12), w, "0%", "persistence — KHÔNG THỂ báo trước theo "
               "định nghĩa", vsize=30)
    d.mini_card(s, x, y + Inches(0.14), w, Inches(1.1), "VÌ SAO",
                ["persistence chỉ nhìn số hiện tại —", "giờ này sạch → nhất định không "
                 "báo"], fill=SOFT_ORANGE, tag_color=ORANGE, h=1.1, size=11.5)
    d.remember(s, "Giờ khởi đầu đợt: mô hình báo trước 304/1.000 = 30,4% — persistence 0% "
                  "theo định nghĩa; giá trị nằm ở chỗ baseline không thể làm.")

    s = d.light("PHẦN C · KẾT QUẢ", "McNemar — kiểm định “hơn thật, không phải may”")
    d.prose(s, [
        "F1 hơn nhau 0,013 — có thể chỉ là nhiễu? ##McNemar## là kiểm định dành cho đúng "
        "câu này: so **hai mô hình trên cùng một tập test**, chỉ xét những giờ hai mô "
        "hình **chênh kết quả** (những giờ hai bên cùng đúng/cùng sai không phân biệt "
        "được ai hơn ai).",
        "Kết quả: ##b = 594## giờ chỉ mô hình đúng, ##c = 298## giờ chỉ persistence đúng "
        "→ ##p ≈ 5×10⁻²³## — xác suất hai mô hình thật ra ngang nhau mà quan sát được "
        "chênh 594 vs 298 gần bằng 0: khác biệt **có ý nghĩa thống kê**. Nhưng nói đủ: ý "
        "nghĩa thống kê ##không đương nhiên## đủ triển khai — test chỉ 6 tháng cuối 2016, "
        "dữ liệu lịch sử; cần theo dõi thêm nhiều mùa.",
    ])
    x, y, w, bot = d.panel(s, "MCNEMAR TRÊN TỪNG GIỜ TEST")
    d.mini_card(s, x, y, w, Inches(0.98), "b = 594", ["giờ CHỈ mô hình đúng"])
    d.mini_card(s, x, y + Inches(1.12), w, Inches(0.98), "c = 298",
                ["giờ CHỈ persistence đúng"], fill=SOFT_ORANGE, tag_color=ORANGE)
    y = d.stat(s, x, y + Inches(2.28), w, "p ≈ 5×10⁻²³",
               "xác suất chênh như vậy do may ≈ 0 → khác biệt thật", vsize=26)
    d.remember(s, "McNemar trên từng giờ test: b = 594 vs c = 298, p ≈ 5×10⁻²³ → khác "
                  "biệt thật — nhưng test chỉ 6 tháng, chưa đủ triển khai.")

    s = d.light("PHẦN C · KẾT QUẢ", "Tầm dự báo 1–3–6 giờ: càng sớm, mô hình càng có lợi")
    d.prose(s, [
        "Chạy cùng hệ thống cho ba tầm, F1 trên tập test: ##1h: 0,914 → 0,927## · ##3h: "
        "0,806 → 0,832## · ##6h: 0,702 → 0,745##.",
        "Mẫu hình rõ: **càng phải báo sớm, persistence càng yếu** (càng xa hiện tại, "
        "“giữ nguyên trạng thái” càng sai) trong khi mô hình đọc được xu hướng trước đó — "
        "nên **lợi thế của mô hình tăng theo tầm**: +0,013 → +0,026 → +0,043. Ở 6h, "
        "recall tăng mạnh ##0,70 → 0,80##, đổi lại precision giảm nhẹ — đúng hướng cảnh "
        "báo sớm cần. Kiểm tra từng trạm: ##12/12 trạm## đều có F1 cao hơn persistence — "
        "lợi thế không nhờ vài trạm đặc biệt.",
    ])
    x, y, w, bot = d.panel(s, "LỢI THẾ F1 THEO TẦM")
    d.text(s, "HistGB trừ persistence, tập test:", x, y, w, Inches(0.3), 11.5, MUTED,
           font=SANS)
    y = d.bar_rows(s, x, y + Inches(0.4), w,
                   [("tầm 1h", 0.013, "0,013"), ("tầm 3h", 0.026, "0,026"),
                    ("tầm 6h", 0.043, "0,043")], max_val=0.043)
    d.hline(s, x, y + Inches(0.1), w)
    y = d.stat(s, x, y + Inches(0.22), w, "12 / 12", "trạm đều F1 cao hơn persistence",
               vsize=26)
    d.chip(s, "recall 6h: 0,70 → 0,80", x, y + Inches(0.1), w, h=Inches(0.42),
           fill=TEAL, size=12)
    d.remember(s, "Lợi thế F1 theo tầm: +0,013 (1h) · +0,026 (3h) · +0,043 (6h) — càng "
                  "báo sớm mô hình càng hơn; 12/12 trạm đều hơn.")

    s = d.light("PHẦN C · MỞ RỘNG", "Thí nghiệm 6 giờ: cảnh báo cả đợt ô nhiễm")
    d.prose(s, [
        "Câu hỏi mở rộng: giờ này đang sạch, liệu có **đợt** ô nhiễm (##≥ 3 giờ liên tiếp "
        "≥ 115## tại một trạm) bắt đầu trong t+1 → t+6 không? Mô hình thêm đặc trưng "
        "**trung vị PM2.5 các trạm lân cận** — nắm thông tin không gian giữa các trạm.",
        "Kết quả: bắt ##506/730 = 69,3%## số đợt; luật đơn giản “hiện tại ≥ 105” chỉ "
        "##64,1%##; ##trung vị thời gian báo trước 3 giờ## (trước đây 1 giờ); trần cảnh "
        "báo nhầm ##0,5 lần/ngày-trạm## chọn trên validation để đỡ “kêu ồn”. Chi tiết "
        "hay hỏi: kiểm tra tới ##t+8## vì đợt bắt đầu muộn nhất t+6 cần tới t+8 mới đủ 3 "
        "giờ liên tiếp. Đây là **thí nghiệm riêng**, chưa thay mô hình t+1 đang khóa.",
    ])
    x, y, w, bot = d.panel(s, "BẮT ĐỢT TRONG 6 GIỜ TỚI")
    d.mini_card(s, x, y, w, Inches(1.06), "MÔ HÌNH",
                ["##506/730 = 69,3%## số đợt được bắt"])
    d.mini_card(s, x, y + Inches(1.2), w, Inches(1.06), "LUẬT “≥ 105”",
                ["468/730 = ##64,1%## — đơn giản nhưng yếu hơn"], fill=SOFT_ORANGE,
                tag_color=ORANGE, h=1.06)
    y = d.stat(s, x, y + Inches(2.42), w, "3 giờ",
               "trung vị thời gian báo trước (trước đây 1 giờ)", vsize=26)
    d.chip(s, "trần 0,5 cảnh báo nhầm/ngày-trạm · kiểm tra tới t+8", x, y + Inches(0.08),
           w, h=Inches(0.42), fill=INK_SOFT, color=AMBER, size=10.5)
    d.remember(s, "Cảnh báo đợt 6h: 69,3% vs luật đơn giản 64,1% · báo trước trung vị 3 "
                  "giờ · trần 0,5 FP/ngày-trạm — thí nghiệm riêng, chưa thay mô hình "
                  "chính.")

    s = d.light("PHẦN C · BỔ TRỢ", "Hai bài bổ trợ: hồi quy số và phân cụm trạm")
    d.prose(s, [
        "**Hồi quy PM2.5 t+1** (dự báo con số thay vì nhãn): sai số ##MAE = 8,87 µg/m³##. "
        "Đọc đúng: MAE là **trung bình** sai số tuyệt đối, cùng đơn vị µg/m³ — không phải "
        "trần sai số từng lần dự báo (có lần sai nhiều hơn 8,87).",
        "**Phân cụm K-Means**: đối tượng là ##12 trạm## — mỗi trạm một vector hồ sơ "
        "(nồng độ, mùa vụ, giờ cao điểm…), chuẩn hoá trước khi phân cụm, thử K = 2, 3, 4. "
        "Kết quả ##K = 2 với silhouette 0,664## cao nhất, chia ##9 trạm + 3 trạm##. "
        "Silhouette đo độ tách cụm theo đặc trưng (−1…1) — **không phải accuracy**; địa "
        "lý chỉ dùng đối chiếu sau, không đưa vào đặc trưng. Hai bài này bổ trợ hiểu dữ "
        "liệu và tạo đặc trưng “nhóm trạm”, không phải đầu ra chính.",
    ])
    x, y, w, bot = d.panel(s, "HAI BÀI BỔ TRỢ")
    y = d.stat(s, x, y, w, "MAE 8,87", "µg/m³ — TRUNG BÌNH sai số tuyệt đối, không phải "
               "trần", vsize=26)
    d.hline(s, x, y, w)
    y = d.stat(s, x, y + Inches(0.12), w, "K = 2 · 0,664", "silhouette cao nhất trong "
               "K = 2, 3, 4 — cụm 9 + 3 trạm", vsize=26)
    y = d.kv(s, x, y + Inches(0.1), w, "Đối tượng cụm", "12 TRẠM (vector hồ sơ), không "
             "phải 420k giờ")
    y = d.kv(s, x, y, w, "Silhouette", "đo độ tách cụm — ##không phải accuracy##")
    d.remember(s, "Hồi quy MAE 8,87 µg/m³ (trung bình, không phải trần) · K-Means cụm 12 "
                  "trạm: K=2, silhouette 0,664, cụm 9+3.")

    s = d.light("PHẦN C · CHỐT", "Nếu chỉ nhớ 5 điều về phần C")
    d.numbered_rows(s, [
        ("1", "Bài toán:", "y = 1 ⟺ PM2.5(t+1h) ≥ 115 µg/m³; 49 đặc trưng chỉ nhìn quá "
         "khứ (no leakage)."),
        ("2", "Chia tập:", "theo thời gian ##276.380 / 47.338 / 48.781##; siêu tham số "
         "chỉ chọn trên val; test dùng một lần."),
        ("3", "Baseline:", "majority acc 0,761 / F1 0,000 (bẫy accuracy) — persistence F1 "
         "##0,914## là sàn phải vượt."),
        ("4", "Kết quả:", "HistGB F1 ##0,927##, AUC 0,993; FN 1.000→838, FP 999→865; khởi "
         "đầu đợt ##30,4%## vs 0%; McNemar ##p ≈ 5×10⁻²³##."),
        ("5", "Tầm & bổ trợ:", "+0,013/+0,026/+0,043 theo 1/3/6h, ##12/12## trạm; đợt 6h "
         "##69,3%##; MAE ##8,87##; K-Means K=2 (0,664)."),
    ], y=1.66, gap=0.9)
    d.remember(s, "Khai phá = nhãn 115 → chia đúng → vượt baseline mạnh → chứng minh "
                  "bằng giờ khởi đầu + McNemar, không chỉ F1.")

    # ---- PHẦN D -------------------------------------------------------------------
    d.dark("PHẦN D · SLIDE 44–48", "Vận hành & tổng ôn",
           ["CDC theo watermark: hệ thống “sống” sau khi báo cáo",
            "Idempotency: chạy lại không nhân đôi · giám sát MAE từng batch",
            "Dashboard 4 trang và cách đọc demo trung thực",
            "Khuyến nghị & hạn chế nói thật — và bộ số phải thuộc"], big="D")

    s = d.light("PHẦN D · VẬN HÀNH", "CDC theo watermark — hệ thống “sống” sau khi báo cáo")
    d.prose(s, [
        "Đến đây kho mới chỉ “chết” trên dữ liệu lịch sử. Luồng ##CDC## (change data "
        "capture — bắt thay đổi) mô phỏng vận hành thật: ##SQLite## đóng vai nguồn giao "
        "dịch (trạm gửi số đo mỗi giờ), Airflow theo lịch kéo dữ liệu mới về lake rồi "
        "upsert vào staging.",
        "Cơ chế ##watermark##: hệ thống ghi mốc ##event_time## đã nạp tới đâu; mỗi lần "
        "chạy chỉ lấy **các bản ghi sau mốc** — không quét lại cả nguồn; đây là cách nạp "
        "tăng dần tiêu biểu khi nguồn không push được sự kiện. Giới hạn phải nói thẳng: "
        "watermark chỉ bắt được ##INSERT mới## — bản ghi cũ bị **sửa hoặc xoá sau khi "
        "mốc đã đi qua thì không ai biết**; phương án là đối soát sha256 từng batch "
        "hoặc nạp lại nguyên partition khi nghi ngờ.",
    ])
    x, y, w, bot = d.panel(s, "LUỒNG CDC")
    y = d.flow_rows(s, x, y, w, [["SQLite", "watermark"], ["Airflow", "lake"],
                                 ["staging", "chấm MAE"]], size=11)
    y += Inches(0.1)
    y = d.kv(s, x, y, w, "Watermark", "mốc event_time đã nạp — chỉ lấy bản ghi SAU mốc")
    y = d.kv(s, x, y, w, "Nạp", "upsert theo khóa grain (không nhân đôi)")
    d.mini_card(s, x, y + Inches(0.1), w, Inches(1.34), "GIỚI HẠN",
                ["bắt được ##INSERT mới##", "sửa / xoá bản ghi cũ sau mốc → ##không "
                 "bắt được## — đối soát lại partition"], fill=SOFT_ORANGE,
                tag_color=ORANGE, h=1.34, size=11.5)
    d.remember(s, "Watermark = mốc event_time, chỉ lấy bản ghi sau mốc → bắt insert mới; "
                  "sửa/xoá bản ghi cũ phải đối soát lại partition.")

    s = d.light("PHẦN D · VẬN HÀNH", "Chạy lại không nhân đôi — idempotency & giám sát MAE")
    d.prose(s, [
        "Vì sao chạy lại không nhân đôi dữ liệu — hai lớp bảo vệ: nạp kiểu ##upsert theo "
        "khóa grain## (trạm × giờ trùng thì ghi đè, không thêm dòng) và ##load_registry## "
        "— sổ ghi từng batch với ##object key + sha256 + schema_version + row_count##; "
        "batch nào đã có “vân tay” khớp thì bỏ qua. Đó chính là nghĩa của **idempotent**.",
        "Mỗi batch nạp xong còn được **chấm MAE đối soát** với số đo thật: batch mới nhất "
        "lệch trung bình ##> 12 µg/m³## → ##bật cờ giám sát (monitor_only)## — mô hình "
        "đang khóa **không tự retrain**, người quyết định xem xét. Tổng quan: MAE đối "
        "soát ##9,41 µg/m³## trên ##111.944 giờ##, chấm liên tục 01/2016–28/02/2017 — "
        "cao hơn 8,87 của tập test vì chạy trên dòng thời gian vận hành, không phải một "
        "tập test cố định.",
    ])
    x, y, w, bot = d.panel(s, "IDEMPOTENT + GIÁM SÁT")
    y = d.kv(s, x, y, w, "Lớp 1", "upsert theo khóa grain trạm × giờ")
    y = d.kv(s, x, y, w, "Lớp 2", "load_registry: key + sha256 + row_count khớp → skip")
    d.hline(s, x, y + Inches(0.06), w)
    y = d.stat(s, x, y + Inches(0.18), w, "9,41 µg/m³",
               "MAE đối soát CDC — 111.944 giờ, 01/2016–02/2017", vsize=26)
    d.chip(s, "vượt 12 µg/m³ → bật cờ · KHÔNG tự retrain", x, y + Inches(0.1), w,
           h=Inches(0.44), fill=TEAL, size=11.5)
    d.remember(s, "Upsert theo khóa + load_registry (sha256) = idempotent · batch nào "
                  "vượt MAE 12 µg/m³ thì bật cờ — không tự retrain.")

    s = d.light("PHẦN D · DASHBOARD", "Dashboard bốn trang — và cách đọc demo trung thực")
    d.prose(s, [
        "Bốn trang chính: ##/## quy mô vấn đề · ##/olap## năm phép thao tác trên cube "
        "(phần B) · ##/canh-bao## nhãn, baseline, kết quả — hai ma trận nhầm lẫn đặt "
        "cạnh nhau thấy ngay mô hình vs cách hiện tại · ##/phat-lai## kéo mốc thời gian, "
        "hệ thống phát dự đoán cho giờ đích, **số đo thật chỉ hiện khi mốc dữ liệu đã đi "
        "qua** — minh hoạ nguyên tắc “dự đoán trước, đối soát sau”.",
        "Một câu trung thực phải nhớ khi demo: ##F1 cao trong cửa sổ demo KHÔNG phải bằng "
        "chứng độ tốt của mô hình## — cửa sổ demo chọn nhiều giờ ô nhiễm cho dễ thấy; "
        "bằng chứng thật nằm ở tập test đã khóa. Dự phòng: mất mạng / Docker lỗi thì mở "
        "notebook OLAP đã chạy sẵn + ảnh chụp màn hình trong slide; toàn hệ thống dựng "
        "lại bằng run_all.py 11,2 phút.",
    ])
    x, y, w, bot = d.panel(s, "BỐN TRANG")
    y = d.kv(s, x, y, w, "/", "quy mô vấn đề")
    y = d.kv(s, x, y, w, "/olap", "năm phép OLAP trên cube")
    y = d.kv(s, x, y, w, "/canh-bao", "2 ma trận nhầm lẫn cạnh nhau + dòng xác suất")
    y = d.kv(s, x, y, w, "/phat-lai", "dự đoán trước — số thật chỉ hiện khi mốc đi qua")
    d.mini_card(s, x, y + Inches(0.12), w, Inches(1.3), "TRUNG THỰC",
                ["F1 trong cửa sổ demo ##không phải bằng chứng##", "— cửa sổ chọn nhiều "
                 "giờ ô nhiễm; bằng chứng ở test"], fill=SOFT_ORANGE, tag_color=ORANGE,
                h=1.3, size=11.5)
    d.remember(s, "/ · /olap · /canh-bao · /phat-lai (dự đoán trước, đối soát sau) — số "
                  "trong cửa sổ demo không dùng làm bằng chứng đánh giá.")

    s = d.light("PHẦN D · TỔNG KẾT", "Khuyến nghị & hạn chế — nói chủ động, đừng chờ bị hỏi")
    d.prose(s, [
        "Ba khuyến nghị rút ra: **(1)** cảnh báo theo ##xác suất mô hình## thay vì chờ số "
        "đo vượt — tận dụng 30,4% giờ khởi đầu; **(2)** ưu tiên nguồn lực vào ##mùa đông, "
        "buổi tối, giờ gió nam## — các điều kiện nguy hiểm cao nhất từ kho; **(3)** khi "
        "cần báo sớm 3–6 giờ, ##chấp nhận precision thấp hơn để recall cao## — bỏ sót "
        "đắt hơn báo nhầm.",
        "Bốn hạn chế nói thật: ##test chỉ 6 tháng cuối 2016## (một mùa, dữ liệu lịch sử); "
        "##+0,013 F1## ở tầm 1 giờ là khiêm tốn; ##ngưỡng 115 vốn là ngưỡng 24 giờ## của "
        "quy chuẩn đang áp cho từng giờ — quy ước thí nghiệm; ##chưa kết nối trạm "
        "thật## — CDC mới mô phỏng. Nêu hạn chế chủ động là điểm cộng: nó cho thấy nhóm "
        "hiểu giới hạn của chính kết quả mình.",
    ])
    x, y, w, bot = d.panel(s, "NÊN / CHƯA")
    d.mini_card(s, x, y, w, Inches(1.78), "3 KHUYẾN NGHỊ",
                ["cảnh báo theo xác suất mô hình", "ưu tiên đông · tối · gió nam",
                 "cần sớm → đổi precision lấy recall"], h=1.78, size=11.5)
    d.mini_card(s, x, y + Inches(1.92), w, Inches(1.78), "4 HẠN CHẾ",
                ["test 6 tháng cuối 2016 · +0,013 khiêm tốn", "115 vốn ngưỡng 24h áp cho "
                 "giờ", "chưa kết nối trạm thật"], fill=SOFT_ORANGE, tag_color=ORANGE,
                h=1.78, size=11.5)
    d.remember(s, "3 khuyến nghị (cảnh báo theo xác suất · ưu tiên đông-tối-gió nam · đổi "
                  "precision lấy recall khi cần sớm) + 4 hạn chế nói thật.")

    s = d.light("PHẦN D · TỔNG KẾT", "Bộ số phải thuộc — đọc được không nhìn slide")
    blocks = [("DỮ LIỆU", TEAL,
               ["##420.768## dòng trạm×giờ · 12 trạm · 03/2013–02/2017",
                "PM2.5 TB ##79,8## · ~1/4 giờ ≥ 115 · CO thiếu ##4,9%##",
                "run_all ##14 bước / 11,2 phút## · 60 dòng cờ trần"]),
              ("KHO", ORANGE,
               ["grain trạm×giờ · 4 chiều · ##180## tổ hợp mức · CUBE = ##32##",
                "marts ##17.532 / 576 / 6.075## · ##24.759/24.759 = 100%## · ##130## test",
                "cube ##2–6 ms## vs fact ##130–400 ms## · gió nam ##+60%## (28→43%)"]),
              ("KHAI PHÁ", TEAL,
               ["split ##276.380 / 47.338 / 48.781## · majority ##0,000##",
                "persistence ##0,914## — cây ##0,921## — HistGB ##0,927## (AUC 0,993)",
                "FN 1.000→838 · FP 999→865 · khởi đầu ##30,4%## · p ≈ ##5×10⁻²³##",
                "tầm ##+0,013 / +0,026 / +0,043## · ##12/12## trạm"]),
              ("VẬN HÀNH", ORANGE,
               ["đợt 6h ##506/730 (69,3%)## vs luật ##64,1%## · trung vị ##3h##",
                "trần ##0,5 FP/ngày-trạm## · MAE hồi quy ##8,87## · CDC ##9,41##",
                "cờ giám sát ở ##12 µg/m³## · K-Means ##K=2 · 0,664 · 9+3##"])]
    gw, gh = Inches(5.98), Inches(2.18)
    for i, (name, color, lines) in enumerate(blocks):
        gx = MARGIN + (i % 2) * (gw + Inches(0.17))
        gy = Inches(1.52) + (i // 2) * (gh + Inches(0.16))
        d.card(s, gx, gy, gw, gh, CARD)
        d.chip(s, name, gx + Inches(0.2), gy + Inches(0.16), Inches(1.7), h=Inches(0.4),
               fill=color, size=12)
        box = s.shapes.add_textbox(gx + Inches(0.2), gy + Inches(0.68),
                                   gw - Inches(0.4), gh - Inches(0.8))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = 0
        for j, ln in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(4)
            p.line_spacing = 1.05
            r = p.add_run()
            r.text = "· "
            r.font.name, r.font.size, r.font.bold = SANS, Pt(12.5), True
            r.font.color.rgb = color
            for bold, orange, txt in parse_runs(ln):
                r = p.add_run()
                r.text = txt
                r.font.name, r.font.size, r.font.bold = SANS, Pt(12.5), bold
                r.font.color.rgb = ORANGE if orange else INK
    d.remember(s, "Thuộc 4 cụm số này là “đủ đạn” cho hầu hết câu hỏi — mỗi số đều giải "
                  "thích được nguồn gốc.")

    # ---- kết -----------------------------------------------------------------------
    s = d.prs.slides.add_slide(d.blank)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK
    d.text(s, "ĐỌC XONG — BA LƯỢT ÔN THEO THỨ TỰ NÀY", MARGIN, Inches(0.9), Inches(11),
           Inches(0.4), 14, AMBER, bold=True, font=SANS)
    d.text(s, "Hiểu rồi — giờ chuyển từ hiểu sang nói được", MARGIN, Inches(1.45),
           Inches(11.6), Inches(0.9), 33, WHITE, bold=True, font=SERIF)
    steps = [("LƯỢT 1", "đọc hiểu toàn bộ (vừa xong) — không cần ghi gì thêm.", SOFT_TEAL,
              TEAL),
             ("LƯỢT 2", "đọc lại các dải GHI NHỚ + 4 slide CHỐT (≈ 15 phút, trước khi ngủ "
              "càng tốt).", SOFT_ORANGE, ORANGE),
             ("LƯỢT 3", "mở ON_TAP_BAO_VE_HOI_DAP.md — tự trả lời không nhìn, yếu chỗ nào "
              "quay lại đúng slide đó.", SOFT_TEAL, TEAL),
             ("CUỐI", "tập nói theo KE_HOACH_NOI_DUNG_THUYET_TRINH.md với bộ chiếu v3 — "
              "mồi nào nằm slide nào, lá chắn nằm phụ lục nào.", SOFT_ORANGE, ORANGE)]
    y = 2.75
    for code, txt, fill, accent in steps:
        d.card(s, MARGIN, Inches(y), Inches(12.13), Inches(0.82), INK_SOFT)
        d.chip(s, code, MARGIN + Inches(0.18), Inches(y + 0.2), Inches(1.25),
               h=Inches(0.42), fill=accent, size=12)
        d.text(s, txt, MARGIN + Inches(1.65), Inches(y + 0.14), Inches(10.2),
               Inches(0.55), 14.5, SLATE_L, font=SERIF, anchor=MSO_ANCHOR.MIDDLE,
               spacing=2)
        y += 0.98
    d.text(s, "Chúc ôn tốt — hẹn gặp ở phần hỏi đáp.", MARGIN, Inches(6.85), Inches(11),
           Inches(0.4), 15, WHITE, italic=True, font=SERIF)
    d.pages.append((s, True))

    d.finish()


if __name__ == "__main__":
    build()
